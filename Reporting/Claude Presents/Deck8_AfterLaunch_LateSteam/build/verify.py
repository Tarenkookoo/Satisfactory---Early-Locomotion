"""Verify tag counts, text contrast, links, notes, transitions, document properties and package integrity.

usage: python3 verify.py [path.pptx]
"""
import os
import re
import sys
import zipfile
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import C

DEFAULT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                          'After_Launch_Late_Steam_Driving_and_Prestige_Explained.pptx'))
EXPECT = {'Agreed': 16, 'Proposed': 4, 'Provisional': 1, 'Undecided': 1, 'Background': 0, None: 3}
EXPECT_SLIDES = {
    'Agreed': [3, 4, 5, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 19, 20, 21],
    'Proposed': [6, 7, 18, 22], 'Provisional': [23], 'Undecided': [24], None: [1, 2, 25]}


def lum(h):
    def f(v):
        v /= 255.0
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def ratio(a, b):
    la, lb = lum(a), lum(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)


def main(path):
    prs = Presentation(path)
    ok = True
    # tag chips
    found = {}
    for n, s in enumerate(prs.slides, 1):
        chip = None
        for sh in s.shapes:
            if sh.name == 'Status chip':
                chip = sh.text_frame.text.strip().capitalize()
        found.setdefault(chip, []).append(n)
    print('Chips by slide:')
    for k in ('Agreed', 'Proposed', 'Provisional', 'Undecided', 'Background', None):
        got = found.get(k, [])
        exp = EXPECT_SLIDES.get(k, [])
        flag = 'OK' if got == exp else 'MISMATCH (expected %s)' % exp
        if got != exp:
            ok = False
        print('  %-12s %2d  %s  %s' % (k or 'no chip', len(got), got, flag))
    # inline chips
    print('Inline chips:')
    for n, s in enumerate(prs.slides, 1):
        for sh in walk(s.shapes):
            if sh.name == 'Inline chip' and sh.has_text_frame:
                print('  slide %2d: %s' % (n, sh.text_frame.text.replace('\v', ' ')))
    # contrast against own fill or the slide background / panel
    print('Contrast (text colour against own fill, else against background and panel):')
    low = 0
    for n, s in enumerate(prs.slides, 1):
        for sh in walk(s.shapes):
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            try:
                fill = sh.fill.fore_color.rgb if sh.fill.type == 1 else None
            except Exception:
                fill = None
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        fg = str(r.font.color.rgb)
                    except Exception:
                        continue
                    bgs = [str(fill)] if fill else [C['bg'], C['panel']]
                    for bg in bgs:
                        rt = ratio(fg, bg)
                        if rt < 4.5 and not (r.hyperlink and r.hyperlink.address):
                            low += 1
                            print('  LOW slide %d "%s" %s on %s = %.2f' % (n, r.text[:25], fg, bg, rt))
    print('  low-contrast runs:', low)
    # notes, transitions, animations
    for n, s in enumerate(prs.slides, 1):
        if s.has_notes_slide:
            print('  slide %d has notes' % n)
            ok = False
        x = etree.tostring(s._element).decode()
        if 'p:transition' in x or 'p:timing' in x or 'transition' in x.lower():
            print('  slide %d has a transition or timing' % n)
            ok = False
    # document properties
    cp = prs.core_properties
    print('Title:', cp.title, '| Author:', cp.author, '| Language:', cp.language)
    # languages
    nolang = 0
    for n, s in enumerate(prs.slides, 1):
        for el in s._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr'):
            if el.get('lang') != 'en-GB':
                nolang += 1
    print('runs without lang en-GB:', nolang)
    # hyperlinks
    print('Hyperlinks:')
    for n, s in enumerate(prs.slides, 1):
        for rel in s.part.rels.values():
            if rel.reltype.endswith('/hyperlink'):
                print('  slide %d -> %s' % (n, rel._target))
    # empty placeholders
    for n, s in enumerate(prs.slides, 1):
        for sh in s.placeholders:
            if sh.has_text_frame and not sh.text_frame.text.strip():
                print('  empty placeholder on slide', n)
                ok = False
    # pictures
    for n, s in enumerate(prs.slides, 1):
        for sh in walk(s.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                print('  picture on slide %d: %s, alt: %s' % (n, sh.name, sh._element.xpath('.//p:cNvPr')[0].get('descr')[:60]))
    # package integrity
    z = zipfile.ZipFile(path)
    bad = z.testzip()
    print('zip test:', 'ok' if bad is None else bad)
    for name in z.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):
            etree.fromstring(z.read(name))
    print('all XML parts parse:', len(z.namelist()), 'parts')
    print('notes slides in package:', [nm for nm in z.namelist() if 'notesSlide' in nm])
    print('file size: %.2f MB' % (os.path.getsize(path) / 1e6))
    print('RESULT:', 'OK' if ok else 'CHECK ABOVE')


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
