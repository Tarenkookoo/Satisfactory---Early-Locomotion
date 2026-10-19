"""Audit the deck: word counts, smallest font, text-fit estimate, alt text, notes, chips.

usage: python3 check.py [path.pptx]
Counted words: caption + diagram labels + inline chips that carry their own text.
Not counted: title, the top-right status chip, legend chips, slide number, credit lines, bare numbers.
"""
import os
import re
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import text_w, wrap_lines

DEFAULT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                          'After_Launch_Late_Steam_Driving_and_Prestige_Explained.pptx'))
SKIP_COUNT = ('Title', 'Status chip', 'Legend chip', 'Slide number', 'Credit line', 'Table')
CHIP_WORDS = {'AGREED', 'PROPOSED', 'PROVISIONAL', 'UNDECIDED', 'BACKGROUND'}


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)


def frames(sh):
    if sh.has_text_frame and sh.text_frame.text.strip():
        yield sh, sh.text_frame
    if getattr(sh, 'has_table', False) and sh.has_table:
        for r in sh.table.rows:
            for c in r.cells:
                yield sh, c.text_frame


def main(path):
    prs = Presentation(path)
    print('slides:', len(prs.slides))
    rows = []
    problems = []
    for n, slide in enumerate(prs.slides, 1):
        words = 0
        minpt = 99
        for sh in walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                d = sh._element.xpath('.//p:cNvPr')[0].get('descr')
                if not d:
                    problems.append('slide %d: no alt text on %s' % (n, sh.name))
            is_table = getattr(sh, 'has_table', False) and sh.has_table
            for owner, tf in frames(sh):
                text = tf.text
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.size:
                            pt = r.font.size.pt
                            minpt = min(minpt, pt) if sh.name not in ('Credit line',) and not is_table else minpt
                if not is_table and sh.name not in SKIP_COUNT and not re.fullmatch(r'[\d\s]+', text.strip()):
                    t = text.replace('\v', ' ')
                    toks = [w for w in re.split(r'\s+', t) if w and w not in CHIP_WORDS
                            and not re.fullmatch(r'[·•|–-]+', w)]
                    words += len(toks)
                # fit estimate
                if not is_table and sh.width and sh.height:
                    wi = (sh.width - tf.margin_left - tf.margin_right) / 914400.0
                    hi = (sh.height - tf.margin_top - tf.margin_bottom) / 914400.0
                    need = 0.0
                    for p in tf.paragraphs:
                        sz = max([r.font.size.pt for r in p.runs if r.font.size] or [14])
                        bold = any(r.font.bold for r in p.runs)
                        ptxt = ''.join('\v' if False else r.text for r in p.runs)
                        # line breaks are separate elements; approximate with element text
                        ptxt = p.text
                        if tf.word_wrap is False:
                            if text_w(ptxt.replace('\v', ' '), sz, bold) > wi + 0.02:
                                problems.append('slide %d: "%s" wider than its chip/box (%s)' % (n, ptxt[:30], sh.name))
                            need += sz * 1.2 / 72
                        else:
                            need += wrap_lines(ptxt, sz, wi, bold) * sz * 1.2 / 72
                    if need > hi + 0.05 and sh.name != 'Title':
                        problems.append('slide %d: text may overflow "%s" (need %.2f in, have %.2f): %s' %
                                        (n, sh.name, need, hi, text[:40].replace('\n', ' / ')))
                    if sh.name == 'Title' and n > 1 and wrap_lines(text, 36, wi, True) > 1:
                        problems.append('slide %d: title wraps' % n)
        notes = slide.has_notes_slide
        rows.append((n, words, minpt if minpt < 99 else None, notes))
    print('slide words minpt notes')
    for n, w, m, nt in rows:
        flag = ''
        if w > 45 and n not in (2, 25):
            flag = '  <-- OVER 45'
        elif w > 25 and n not in (2, 25):
            flag = '  (over the 25-word aim)'
        print('%5d %5d %5s %s%s' % (n, w, m, 'NOTES' if nt else '-', flag))
    print('problems:', len(problems))
    for p in problems:
        print('  ', p)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT)
