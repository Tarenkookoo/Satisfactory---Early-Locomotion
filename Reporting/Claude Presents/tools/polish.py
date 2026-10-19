"""Add smooth slide transitions and automatic fade-in builds to a deck.

usage: polish.py in.pptx out.pptx

- Every slide gets a 0.7 s Fade transition.
- On each slide, top-level shapes (diagram groups, pictures, captions, rules) fade in
  automatically, in z-order, with no clicks. The title, tag chip, slide number, credit
  line and series line stay put so the slide never looks empty on arrival.
- Viewers that ignore animations simply show the finished slide.
"""
import sys
from lxml import etree
from pptx import Presentation

P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'

STATIC_PREFIXES = ('chip', 'status chip', 'slide number', 'credit', 'series line', 'section label')
STATIC_EXACT = ('num',)
FADE_MS = 500
MAX_SPAN_MS = 1800
MAX_STEP_MS = 300

TRANSITION = (
    f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p="{P}">'
    f'<mc:Choice xmlns:p14="{P14}" Requires="p14">'
    '<p:transition spd="med" p14:dur="700"><p:fade/></p:transition></mc:Choice>'
    '<mc:Fallback><p:transition spd="med"><p:fade/></p:transition></mc:Fallback>'
    '</mc:AlternateContent>'
)


def is_static(sh):
    if sh.is_placeholder:
        return True  # title and any other placeholder
    name = sh.name.lower()
    return name.startswith(STATIC_PREFIXES) or name in STATIC_EXACT


def timing_xml(spids):
    n = len(spids)
    step = min(MAX_STEP_MS, MAX_SPAN_MS // max(n, 1))
    nid = [4]

    def next_id():
        nid[0] += 1
        return nid[0]

    effects = []
    for i, spid in enumerate(spids):
        node = 'afterEffect' if i == 0 else 'withEffect'
        a, b, c = next_id(), next_id(), next_id()
        effects.append(
            '<p:par>'
            f'<p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}">'
            f'<p:stCondLst><p:cond delay="{i * step}"/></p:stCondLst><p:childTnLst>'
            '<p:set><p:cBhvr>'
            f'<p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
            '<p:to><p:strVal val="visible"/></p:to></p:set>'
            '<p:animEffect transition="in" filter="fade"><p:cBhvr>'
            f'<p:cTn id="{c}" dur="{FADE_MS}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            '</p:childTnLst></p:cTn></p:par>'
        )
    return (
        f'<p:timing xmlns:p="{P}"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        + ''.join(effects) +
        '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
        '</p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )


def polish(src, dst):
    prs = Presentation(src)
    total = 0
    for slide in prs.slides:
        sld = slide._element
        for tag in ('transition', 'timing'):
            for el in sld.findall(f'{{{P}}}{tag}'):
                sld.remove(el)
        for el in sld.findall(f'{{{MC}}}AlternateContent'):
            sld.remove(el)
        anchor = sld.find(f'{{{P}}}clrMapOvr')
        if anchor is None:
            anchor = sld.find(f'{{{P}}}cSld')
        trans = etree.fromstring(TRANSITION)
        anchor.addnext(trans)
        spids = [sh.shape_id for sh in slide.shapes if not is_static(sh)]
        if spids:
            trans.addnext(etree.fromstring(timing_xml(spids)))
            total += len(spids)
    prs.save(dst)
    return len(prs.slides), total


if __name__ == '__main__':
    n, t = polish(sys.argv[1], sys.argv[2])
    print(f'{sys.argv[2]}: {n} slides, {t} animated shapes')
