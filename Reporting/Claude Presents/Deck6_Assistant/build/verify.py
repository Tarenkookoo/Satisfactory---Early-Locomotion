"""Checks for Deck 6. Usage: python3 verify.py [first last]

Prints, per slide: word count (caption plus labels), smallest font size, shapes outside the slide,
text-overflow estimates, missing alt text. Then deck-wide checks: tag counts, notes slides,
transitions, animations, language, document properties, contrast of text on its own fill.
"""
import os
import re
import sys
import zipfile

from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from lib import text_w, wrap_count  # noqa: E402

PPTX = os.path.abspath(os.path.join(HERE, "..", "Railway_Guide_Lessons_and_Build_Menu_Explained.pptx"))
EXEMPT_WORDS = {2, 23}
EXPECTED_TAGS = {"AGREED": [3, 4, 5, 6, 7, 8, 9, 10, 12], "PROPOSED": [11, 14, 15, 16, 17, 18, 20],
                 "UNDECIDED": [13, 19, 22], "PROVISIONAL": [21], "NONE": [1, 2, 23]}
SW, SH = 13.333, 7.5
BG = (0x1E, 0x21, 0x26)
NOT_COUNTED = ("CHIP", "CHIP LEGEND", "CHIP CARD", "NUM", "TITLE")


def lum(c):
    def f(v):
        v /= 255.0
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(c[0]) + 0.7152 * f(c[1]) + 0.0722 * f(c[2])


def ratio(a, b):
    la, lb = lum(a), lum(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


def walk(shapes):
    for sp in shapes:
        yield sp
        if sp.shape_type == 6:  # group
            yield from walk(sp.shapes)


def fill_of(sp):
    try:
        if sp.fill.type == 1:
            c = sp.fill.fore_color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None


def bbox(sp):
    return (Emu(sp.left).inches, Emu(sp.top).inches, Emu(sp.left + sp.width).inches, Emu(sp.top + sp.height).inches)


def under_fill(flat, sp):
    l, t, r, b = bbox(sp)
    cx, cy = (l + r) / 2, (t + b) / 2
    idx = [i for i, o in enumerate(flat) if o is sp or o._element is sp._element]
    for o in reversed(flat[:idx[0]] if idx else flat):
        if o.shape_type == 6:
            continue
        f = fill_of(o)
        if f is None:
            continue
        ol, ot, orr, ob = bbox(o)
        if ol <= cx <= orr and ot <= cy <= ob:
            return f
    return None


def runs_of(sp):
    if not sp.has_text_frame:
        return
    for p in sp.text_frame.paragraphs:
        for r in p.runs:
            yield r


def main():
    first = int(sys.argv[1]) if len(sys.argv) > 2 else 1
    last = int(sys.argv[2]) if len(sys.argv) > 2 else 999
    prs = Presentation(PPTX)
    problems = 0
    chips_by_slide = {}
    print(f"{'#':>2} {'words':>5} {'minpt':>5}  title")
    for i, s in enumerate(prs.slides, 1):
        if i < first or i > last:
            continue
        words = 0
        minpt = 999
        warn = []
        title = s.shapes.title.text_frame.text if s.shapes.title is not None else "(no title)"
        chip = None
        flat = list(walk(s.shapes))
        for sp in flat:
            # bounds
            if sp.width is not None and sp.left is not None:
                l, t, r, b = (Emu(sp.left).inches, Emu(sp.top).inches, Emu(sp.left + sp.width).inches,
                              Emu(sp.top + sp.height).inches)
                if sp.shape_type != 6 and (l < -0.01 or t < -0.01 or r > SW + 0.01 or b > SH + 0.01):
                    warn.append(f"outside slide: {sp.name} ({l:.2f},{t:.2f},{r:.2f},{b:.2f})")
            if sp.name == "CHIP":
                chip = sp.text_frame.text
            # alt text
            if sp.shape_type == 6:
                d = sp._element.xpath("./p:nvGrpSpPr/p:cNvPr")[0].get("descr")
                if not d:
                    warn.append(f"group without alt text: {sp.name}")
            if sp.has_text_frame and sp.text_frame.text.strip():
                txt = sp.text_frame.text
                sizes = [r.font.size.pt for r in runs_of(sp) if r.font.size]
                if sizes:
                    minpt = min(minpt, min(sizes))
                    if min(sizes) < 14 and sp.name != "CREDIT":
                        warn.append(f"font below 14 pt: '{txt[:30]}' {min(sizes)}")
                if sp.name not in NOT_COUNTED and not re.fullmatch(r"\d+", txt.strip()):
                    words += len(txt.split())
                # language
                for r in runs_of(sp):
                    lang = r._r.get_or_add_rPr().get("lang")
                    if lang != "en-GB":
                        warn.append(f"language {lang} on '{r.text[:20]}'")
                        break
                # overflow estimate
                if sizes and sp.name != "TITLE":
                    tf = sp.text_frame
                    w = Emu(sp.width).inches - Emu(tf.margin_left).inches - Emu(tf.margin_right).inches
                    h = Emu(sp.height).inches - Emu(tf.margin_top).inches - Emu(tf.margin_bottom).inches
                    sz = max(sizes)
                    bold = any(r.font.bold for r in runs_of(sp))
                    n = sum(wrap_count(p.text, sz, w, bold) for p in tf.paragraphs)
                    need = n * sz * 1.2 / 72.0
                    if need > h + 0.04:
                        warn.append(f"overflow? '{txt[:28]}' needs {need:.2f} in, has {h:.2f} (lines {n})")
                    if len(tf.paragraphs) == 1 and text_w(max(txt.split(), key=len), sz, bold) > w + 0.02:
                        warn.append(f"long word wider than box: '{txt[:28]}'")
                # contrast against own fill, else the nearest filled shape underneath, else background
                f = fill_of(sp) or under_fill(flat, sp) or BG
                for r in runs_of(sp):
                    try:
                        c = r.font.color.rgb
                    except Exception:
                        continue
                    cr = ratio((c[0], c[1], c[2]), f)
                    if cr < 4.5:
                        warn.append(f"contrast {cr:.1f}: '{r.text[:24]}' on fill {f}")
                        break
            if sp.shape_type == 13 and not sp._element.xpath("./p:nvPicPr/p:cNvPr")[0].get("descr"):
                warn.append("picture without alt")
        chips_by_slide[i] = chip
        flag = ""
        if i not in EXEMPT_WORDS and words > 45:
            flag = "  <-- OVER 45"
            problems += 1
        elif i not in EXEMPT_WORDS and words > 25:
            flag = "  (over the 25-word aim)"
        print(f"{i:>2} {words:>5} {minpt:>5}  {title}{flag}")
        for w in warn:
            print("     !", w)
    # deck-wide
    if first == 1 and last == 999:
        print("\nTag check (slide chip):")
        got = {}
        for i, c in chips_by_slide.items():
            got.setdefault(c or "NONE", []).append(i)
        for k, v in EXPECTED_TAGS.items():
            ok = sorted(got.get(k, [])) == v
            print(f"  {k:12} {'OK ' if ok else 'BAD'} expected {v} got {sorted(got.get(k, []))}")
            problems += 0 if ok else 1
    z = zipfile.ZipFile(PPTX)
    names = z.namelist()
    notes = [n for n in names if "notesSlide" in n]
    print("\nNotes slides:", len(notes), "(want 0)")
    trans = anim = 0
    for n in names:
        if re.match(r"ppt/slides/slide\d+\.xml", n):
            x = z.read(n).decode("utf8")
            trans += x.count("<p:transition")
            anim += x.count("<p:timing")
    print("Transitions:", trans, " Animations:", anim, "(want 0, 0)")
    pics = [n for n in names if n.startswith("ppt/media/")]
    print("Media files:", len(pics), "(want 0)")
    cp = prs.core_properties
    print("Title:", cp.title, "| Author:", cp.author, "| Language:", cp.language)
    print("Slides:", len(prs.slides), " File size: %.2f MB" % (os.path.getsize(PPTX) / 1e6))
    bad = [n for n in names if n.endswith(".xml")]
    from lxml import etree
    for n in bad:
        etree.fromstring(z.read(n))
    print("XML parts parsed:", len(bad))
    print("PROBLEMS:", problems)


if __name__ == "__main__":
    main()
