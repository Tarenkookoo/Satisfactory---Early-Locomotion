"""Shape library for Deck 6 ("Your Railway Guide: Lessons and the Build Menu, Explained").

All helpers take `sh`, a shapes collection (slide.shapes or group.shapes), so the
same call works on a slide or inside a group. Sizes are in inches.
This deck uses no pictures: every drawing is a native shape.
"""
import math
import re

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# --- style tokens (OUTLINE.md) -------------------------------------------
BG = "#1E2126"
PANEL = "#2B2F36"
PANEL_ALT = "#343A43"
LINE = "#4A515B"
TEXT = "#F2F2F2"
MUTED = "#A9B0B8"
ORANGE = "#F28C28"
RED = "#E5484D"
GREEN = "#3FB950"
AMBER = "#F2C14E"
WATER = "#4DA3FF"
ELEC = "#6CA0DC"
COAL = "#1B1D21"
COAL_LINE = "#6B7078"
STEEL = "#9AA0A6"
SLEEPER = "#6B5B4B"
PURPLE = "#B084F5"
DARKMETAL = "#3A3F47"
DARK = "#101418"  # text on bright fills
WAGON = "#6C7480"
COOL = "#5B6672"  # cool grey (ADA bubble)
TABLE_WHITE = "#F4F5F6"

CHIPS = {
    "Agreed": (GREEN, DARK, None),
    "Proposed": (PURPLE, DARK, None),
    "Provisional": (WATER, DARK, None),
    "Undecided": (None, TEXT, MUTED),
    "Background": (LINE, TEXT, None),
}

FONT = "Arial"
FONT_R = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
_fcache = {}


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def I(v):
    return Inches(v)


# --- text measuring (Liberation Sans has Arial metrics) --------------------
def text_w(s, size, bold=False):
    if bold not in _fcache:
        _fcache[bold] = ImageFont.truetype(FONT_B if bold else FONT_R, 200)
    return _fcache[bold].getlength(s) * size / 200.0 / 72.0


def wrap_count(s, size, width, bold=False):
    """Number of lines `s` takes at `size` pt in `width` inches."""
    lines = 0
    for para in s.split("\n"):
        cur = ""
        n = 1
        for w in para.split(" "):
            t = (cur + " " + w).strip()
            if text_w(t, size, bold) <= width or not cur:
                cur = t
            else:
                n += 1
                cur = w
        lines += n
    return lines


# --- deck ------------------------------------------------------------------
def new_deck(title, subject=""):
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    cp = prs.core_properties
    cp.title = title
    cp.author = "Claude (Early Locomotion project)"
    cp.last_modified_by = "Claude (Early Locomotion project)"
    cp.language = "en-GB"
    cp.subject = subject
    cp.keywords = "Satisfactory; Early Locomotion; railway guide; lessons; build menu"
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    xml = theme_part.blob.decode("utf8")
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r"\1Arial\2", xml)
    theme_part._blob = xml.encode("utf8")
    return prs


def new_slide(prs, title, chip_kind=None, title_size=36, title_w=10.2, title_box=None):
    """Title Only layout: real title placeholder (good for navigation), then the chip."""
    layout = prs.slide_layouts[5]
    s = prs.slides.add_slide(layout)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    t = s.shapes.title
    bx = title_box or (0.6, 0.4, title_w, 0.8)
    t.left, t.top, t.width, t.height = I(bx[0]), I(bx[1]), I(bx[2]), I(bx[3])
    t.name = "TITLE"
    tf = t.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    _font(r, title_size, TEXT, True)
    if chip_kind:
        chip(s.shapes, chip_kind, 13.333 - 0.6 - 1.7, 0.58, name="CHIP")
    return s


def finish(slide, caption, n):
    """Caption bar (20 pt) and slide number, added last (reading order)."""
    sh = slide.shapes
    bar_y, bar_h = 6.3, 0.8
    rect(sh, 0.6, bar_y, 12.13, bar_h, PANEL, None, radius=0.08, name="CAPTION BAR")
    rect(sh, 0.6, bar_y + 0.06, 0.08, bar_h - 0.12, ORANGE, None, name="CAPTION RULE")
    tb = label(sh, 0.85, bar_y, 11.75, bar_h, caption, 20, TEXT, align="l", anchor="m", name="CAPTION")
    if n is not None:
        label(sh, 11.73, 7.14, 1.0, 0.3, str(n), 14, MUTED, align="r", name="NUM")
    return tb


# --- low-level styling ------------------------------------------------------
def _font(run, size, color, bold=False, italic=False, strike=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.language_id = MSO_LANGUAGE_ID.ENGLISH_UK
    if strike:
        run._r.get_or_add_rPr().set("strike", "sngStrike")


def _strip_style(shape):
    el = shape._element
    for st in el.findall(qn("p:style")):
        el.remove(st)


def _set_alpha(fill_parent_elm, alpha):
    clr = fill_parent_elm.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))


def style(shape, fill=None, line=None, lw=1.5, dash=None, alpha=None):
    _strip_style(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _set_alpha(shape._element.spPr.find(qn("a:solidFill")), alpha)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
        if dash:
            shape.line.dash_style = dash


def set_alt(shape, descr, name=None):
    el = shape._element
    found = el.xpath(
        "./p:nvSpPr/p:cNvPr | ./p:nvGrpSpPr/p:cNvPr | ./p:nvPicPr/p:cNvPr | "
        "./p:nvCxnSpPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr"
    )
    c = found[0]
    c.set("descr", descr)
    if name:
        c.set("name", name)


def _fill_tf(tf, text, size, color, bold, align, anchor, margins, spacing=None):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [I(m) for m in margins]
    paras = text if isinstance(text, list) else [text]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        runs = para if isinstance(para, list) else [para]
        for rn in runs:
            t, o = rn if isinstance(rn, tuple) else (rn, {})
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False),
                  o.get("strike", False))


# --- primitives --------------------------------------------------------------
def rect(sh, x, y, w, h, fill=None, line=None, lw=1.5, radius=None, dash=None, alpha=None, name=None,
         text=None, size=14, color=TEXT, bold=False, align="c", anchor="m", margins=(0.06, 0.03, 0.06, 0.03)):
    shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    if text is not None:
        _fill_tf(shp.text_frame, text, size, color, bold, align, anchor, margins)
    return shp


def shape(sh, kind, x, y, w, h, fill=None, line=None, lw=1.5, dash=None, alpha=None, name=None, adj=None, rot=None):
    shp = sh.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        for i, a in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
            shp.adjustments[i] = a
    style(shp, fill, line, lw, dash, alpha)
    if rot:
        shp.rotation = rot
    if name:
        shp.name = name
    return shp


def oval(sh, x, y, w, h, fill=None, line=None, lw=1.5, alpha=None, name=None, text=None, size=14, color=TEXT,
         bold=False, dash=None):
    shp = sh.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(w), I(h))
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    if text is not None:
        _fill_tf(shp.text_frame, text, size, color, bold, "c", "m", (0, 0, 0, 0))
    return shp


def label(sh, x, y, w, h, text, size=14, color=TEXT, bold=False, align="c", anchor="m", name=None,
          margins=(0.02, 0.01, 0.02, 0.01), spacing=None):
    tb = sh.add_textbox(I(x), I(y), I(w), I(h))
    _fill_tf(tb.text_frame, text, size, color, bold, align, anchor, margins, spacing)
    if name:
        tb.name = name
    return tb


def poly(sh, pts, fill=None, line=None, lw=1.5, closed=True, alpha=None, name=None, dash=None):
    """Freeform polygon / polyline through `pts` (inches)."""
    e = [(int(Inches(px)), int(Inches(py))) for px, py in pts]
    fb = sh.build_freeform(e[0][0], e[0][1], scale=1.0)
    fb.add_line_segments(e[1:], close=closed)
    shp = fb.convert_to_shape()
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    return shp


def _ends(conn, head, tail):
    ln = conn.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "med")
            e.set("len", "med")


def line(sh, x1, y1, x2, y2, color=STEEL, w=2.25, dash=None, head=False, tail=False, name=None, round_cap=False):
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    if round_cap:
        c.line._get_or_add_ln().set("cap", "rnd")
    _ends(c, head, tail)
    if name:
        c.name = name
    return c


def arrow(sh, x1, y1, x2, y2, color=ORANGE, w=2.25, dash=None, name=None):
    return line(sh, x1, y1, x2, y2, color, w, dash, False, True, name)


def curve(sh, p0, p1, p2, color=ORANGE, w=2.25, n=28, head=True, dash=None, name=None):
    """Quadratic Bezier polyline p0 -> p2 bending towards p1, optional arrow head at the end."""
    pts = []
    for i in range(n + 1):
        t = i / n
        pts.append(((1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * p1[0] + t * t * p2[0],
                    (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * p1[1] + t * t * p2[1]))
    shp = poly(sh, pts, None, color, w, closed=False, name=name, dash=dash)
    if head:
        _ends(shp, False, True)
    return shp


def arc_pts(cx, cy, r, a0, a1, n=24):
    """Points on a circle from angle a0 to a1 (degrees, 0 = right, counter-clockwise up)."""
    return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / n)),
             cy - r * math.sin(math.radians(a0 + (a1 - a0) * i / n))) for i in range(n + 1)]


def ring_pts(cx, cy, r_out, r_in, a0, a1, n=24):
    """Closed ring segment between two radii."""
    return arc_pts(cx, cy, r_out, a0, a1, n) + arc_pts(cx, cy, r_in, a1, a0, n)


def group(sh, name, alt):
    g = sh.add_group_shape()
    g.name = name
    set_alt(g, alt, name)
    return g


# --- chips --------------------------------------------------------------------
def chip(sh, kind, x, y, w=1.7, h=0.4, size=14, text=None, name="CHIP", align="c"):
    fill, tcol, outline = CHIPS[kind]
    t = text if text is not None else kind.upper()
    return rect(sh, x, y, w, h, fill, outline, 2.0, radius=min(0.12, h / 2), name=name, text=t, size=size,
                color=tcol, bold=True, align=align, margins=(0.06, 0.0, 0.06, 0.0))


def ichip(sh, kind, x, y, text, w=None, h=0.34, align="c"):
    """Inline chip with its own words (14 pt: nothing in this deck is smaller except 10 pt credits)."""
    if w is None:
        w = text_w(text, 14, True) + 0.34
    return chip(sh, kind, x, y, w, h, 14, text, name="ICHIP", align=align)


def ichip_w(text):
    return text_w(text, 14, True) + 0.34


# --- shared icons --------------------------------------------------------------
def tick(sh, cx, cy, d=0.6, fill=GREEN, name="tick"):
    oval(sh, cx - d / 2, cy - d / 2, d, d, fill, None, name=name + " disc")
    u = d / 10
    poly(sh, [(cx - 3 * u, cy + 0.2 * u), (cx - 0.8 * u, cy + 2.6 * u), (cx + 3.4 * u, cy - 2.6 * u)], None, DARK,
         d * 7, closed=False, name=name + " mark")


def cross(sh, cx, cy, d=0.6, fill=RED, name="cross"):
    oval(sh, cx - d / 2, cy - d / 2, d, d, fill, None, name=name + " disc")
    u = d / 10
    line(sh, cx - 2.6 * u, cy - 2.6 * u, cx + 2.6 * u, cy + 2.6 * u, TEXT, d * 6, name=name + " bar", round_cap=True)
    line(sh, cx - 2.6 * u, cy + 2.6 * u, cx + 2.6 * u, cy - 2.6 * u, TEXT, d * 6, name=name + " bar", round_cap=True)


def padlock(sh, cx, cy, w=0.5, body=MUTED, name="padlock"):
    bw, bh = w, w * 0.78
    sw = w * 0.62
    # shackle: arch over the body
    poly(sh, ring_pts(cx, cy - bh * 0.1, sw / 2, sw / 2 - w * 0.12, 0, 180), body, None, name=name + " shackle")
    rect(sh, cx - sw / 2, cy - bh * 0.1, w * 0.12, bh * 0.35, body, None, name=name + " shackle leg")
    rect(sh, cx + sw / 2 - w * 0.12, cy - bh * 0.1, w * 0.12, bh * 0.35, body, None, name=name + " shackle leg")
    rect(sh, cx - bw / 2, cy + bh * 0.1, bw, bh, body, None, radius=w * 0.1, name=name + " body")
    oval(sh, cx - w * 0.07, cy + bh * 0.35, w * 0.14, w * 0.14, DARK, None, name=name + " keyhole")


def play_tri(sh, cx, cy, s=0.4, fill=DARK, name="play"):
    poly(sh, [(cx - s * 0.35, cy - s * 0.5), (cx - s * 0.35, cy + s * 0.5), (cx + s * 0.5, cy)], fill, None,
         name=name)


def loco_side(sh, x, y, w, color=ORANGE, name="engine"):
    """Small generic side-view engine from shapes, facing right; w is the width, height is 0.55 w."""
    u = w / 10.0
    X = lambda a: x + a * u
    Y = lambda b: y + b * u
    rect(sh, X(0), Y(1.4), 3.0 * u, 2.6 * u, color, None, radius=0.3 * u, name=name + " cab")
    rect(sh, X(2.8), Y(2.2), 5.6 * u, 1.8 * u, color, None, radius=0.8 * u, name=name + " boiler")
    rect(sh, X(7.0), Y(0.6), 0.9 * u, 1.7 * u, DARKMETAL, STEEL, 0.75, name=name + " chimney")
    rect(sh, X(0), Y(4.0), 8.4 * u, 0.4 * u, DARKMETAL, STEEL, 0.75, name=name + " frame")
    for cx in (1.5, 3.6, 5.7, 7.5):
        oval(sh, X(cx) - 0.7 * u, Y(4.6) - 0.7 * u, 1.4 * u, 1.4 * u, COAL, STEEL, 1.25, name=name + " wheel")


def semaphore(sh, bx, by, h, state="stop", name="semaphore", on_white=False):
    """Signal post with a pivoting arm; (bx, by) is the foot of the post, h the post height.
    stop = arm flat; clear = arm raised 45 degrees. A lamp circle shows the state too."""
    pw = h * 0.1
    post = "#E8E8E8" if not on_white else "#6B7078"
    rect(sh, bx - pw / 2, by - h, pw, h, post, "#6B7078" if not on_white else None, 0.75, name=name + " post")
    rect(sh, bx - h * 0.22, by - h * 0.06, h * 0.44, h * 0.06, DARKMETAL, None, name=name + " foot")
    aw, ah = h * 0.62, h * 0.12
    piv = (bx + pw / 2, by - h + h * 0.12)
    if state == "stop":
        rect(sh, piv[0], piv[1] - ah / 2, aw, ah, RED, None, name=name + " arm")
        rect(sh, piv[0] + aw * 0.55, piv[1] - ah / 2, aw * 0.12, ah, TEXT, None, name=name + " arm stripe")
    else:
        ang = math.radians(45)
        tx, ty = piv[0] + aw * math.cos(ang), piv[1] - aw * math.sin(ang)
        th = ah
        nx, ny = math.sin(ang) * th / 2, math.cos(ang) * th / 2
        poly(sh, [(piv[0] - nx, piv[1] - ny), (tx - nx, ty - ny), (tx + nx, ty + ny), (piv[0] + nx, piv[1] + ny)],
             RED, None, name=name + " arm")
    lamp = RED if state == "stop" else GREEN
    oval(sh, bx - h * 0.1, by - h - h * 0.02, h * 0.2, h * 0.2, lamp, "#2B2F36", 1.0, name=name + " lamp")
    oval(sh, piv[0] - h * 0.04, piv[1] - h * 0.04, h * 0.08, h * 0.08, DARKMETAL, None, name=name + " pivot")
    return piv


def signal_box_building(sh, x, y, w, h, name="signal box", on_white=False):
    """Small signal box: lower brick-coloured body, upper window floor, roof."""
    body = "#8C5A3C"
    rect(sh, x, y + h * 0.52, w, h * 0.48, body, None, name=name + " lower")
    rect(sh, x - w * 0.04, y + h * 0.22, w * 1.08, h * 0.32, "#C9CED3" if not on_white else "#AEB4BB", None,
         name=name + " upper")
    rect(sh, x + w * 0.08, y + h * 0.28, w * 0.84, h * 0.18, "#4DA3FF", None, name=name + " windows", alpha=0.75)
    poly(sh, [(x - w * 0.08, y + h * 0.22), (x + w * 0.5, y), (x + w * 1.08, y + h * 0.22)], "#4A515B", None,
         name=name + " roof")
    rect(sh, x + w * 0.4, y + h * 0.66, w * 0.2, h * 0.34, "#3A3F47", None, name=name + " door")


def station_building(sh, x, y, w, h, name="station"):
    rect(sh, x, y + h * 0.3, w, h * 0.7, "#C9CED3", "#6B7078", 0.75, name=name + " body")
    poly(sh, [(x - w * 0.06, y + h * 0.3), (x + w * 0.5, y), (x + w * 1.06, y + h * 0.3)], "#6B7078", None,
         name=name + " roof")
    for k in range(3):
        rect(sh, x + w * (0.1 + 0.3 * k), y + h * 0.5, w * 0.2, h * 0.28, "#4DA3FF", None, alpha=0.75,
             name=name + " window")


def mini_train(sh, x, y, w, name="mini train"):
    """Tiny train facing right: engine plus two wagons; w is the full width, height is 0.3 w."""
    u = w / 10.0
    X = lambda a: x + a * u
    Y = lambda b: y + b * u
    rect(sh, X(0), Y(0.6), 3.4 * u, 2.2 * u, ORANGE, None, radius=0.4 * u, name=name + " engine")
    rect(sh, X(2.5), Y(0.0), 0.7 * u, 1.0 * u, DARKMETAL, None, name=name + " chimney")
    rect(sh, X(0), Y(2.0), 3.4 * u, 0.8 * u, COAL, None, name=name + " frame")
    for k in range(2):
        rect(sh, X(3.6 + k * 3.2), Y(0.9), 3.0 * u, 1.9 * u, WAGON, "#4A515B", 0.75, radius=0.2 * u,
             name=name + " wagon")
    for cx in (0.9, 2.5, 4.4, 6.0, 7.6, 9.2):
        oval(sh, X(cx) - 0.38 * u, Y(2.8) - 0.38 * u, 0.76 * u, 0.76 * u, COAL, STEEL, 0.75, name=name + " wheel")


# --- the placeholder avatar ---------------------------------------------------------
AVATAR_ALT = ("Placeholder avatar for the railway guide: a plain round head with no face, a headset and "
              "orange shoulders. Her look is undecided (OQ15).")


def avatar(sh, cx, top, d, name="Placeholder avatar", alt=AVATAR_ALT):
    """Neutral placeholder: featureless head circle with a headset, on orange shoulders.
    (cx, top) = centre-x and top of the headset band; d = head diameter. Returns the bottom y."""
    g = group(sh, name, alt)
    gs = g.shapes
    R = d / 2
    band_t = d * 0.13  # headset band sits above the head
    hy = top + band_t + R  # head centre y
    # shoulders
    sw, shh = d * 2.0, d * 0.95
    sy = hy + R * 0.95
    shape(gs, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, cx - sw / 2, sy, sw, shh, ORANGE, None, name="shoulders", adj=[0.5, 0])
    rect(gs, cx - d * 0.2, hy + R * 0.7, d * 0.4, d * 0.4, "#B9BEC4", None, name="neck")
    oval(gs, cx - R, hy - R, d, d, "#D5D9DE", "#9AA0A6", 1.5, name="head")
    # headset: band over the top, ear cups at the sides, mic boom
    poly(gs, ring_pts(cx, hy, R + band_t, R + band_t * 0.12, 8, 172), DARKMETAL, STEEL, 1.5, name="headset band")
    for sgn in (-1, 1):
        rect(gs, cx + sgn * (R + band_t * 0.55) - d * 0.09, hy - d * 0.14, d * 0.18, d * 0.34, DARKMETAL, STEEL, 1.5,
             radius=d * 0.07, name="ear cup")
    curve(gs, (cx - R - band_t * 0.5, hy + d * 0.18), (cx - R - band_t * 0.5, hy + R * 0.95), (cx - d * 0.17, hy + R * 0.8),
          STEEL, max(2.0, d * 3.2), head=False, name="mic boom")
    oval(gs, cx - d * 0.17 - d * 0.06, hy + R * 0.8 - d * 0.06, d * 0.12, d * 0.12, STEEL, None, name="mic tip")
    return sy + shh


def speech_bubble(sh, x, y, w, h, text, fill, tcolor, size=20, bold=True, tail="bl", line_c=None, name="bubble",
                  tail_len=0.4, tail_w=0.4, align="c"):
    """Rounded bubble with a triangular tail. tail: 'bl','br','bottom-left' style corners or 'tl','tr'."""
    rect(sh, x, y, w, h, fill, line_c, 2.0, radius=min(0.3, h / 3), name=name + " body", text=text, size=size,
         color=tcolor, bold=bold, align=align, margins=(0.12, 0.05, 0.12, 0.05))
    if tail in ("bl", "br"):
        tx = x + (0.55 if tail == "bl" else w - 0.55 - tail_w)
        pts = [(tx, y + h - 0.02), (tx + tail_w, y + h - 0.02),
               (tx + (tail_w * 0.15 if tail == "bl" else tail_w * 0.85), y + h + tail_len)]
    else:  # tl / tr
        tx = x + (0.55 if tail == "tl" else w - 0.55 - tail_w)
        pts = [(tx, y + 0.02), (tx + tail_w, y + 0.02),
               (tx + (tail_w * 0.15 if tail == "tl" else tail_w * 0.85), y - tail_len)]
    poly(sh, pts, fill, line_c, 2.0 if line_c else 1.0, name=name + " tail")


# --- miniature model-railway table (slides 8, 9, 10) ---------------------------------
TABLE_W, TABLE_H = 10.1, 3.6


def stadium(sh, x, y, w, h, line_c, lw, name="loop"):
    shp = rect(sh, x, y, w, h, None, line_c, lw, radius=h / 2, name=name)
    shp.adjustments[0] = 0.5
    return shp


def miniature_table(sh, x, y, k=1.0, labels=True, sem_state="stop", label_color_on_table=PANEL):
    """White table top with a thin front edge, a track loop, a train, a semaphore, a signal box and a station.
    (x, y) top-left of the table top; k scales everything (base 10.1 x 3.6 in).
    Returns dict of key points (slide coordinates)."""
    S = lambda v: v * k
    X = lambda v: x + v * k
    Y = lambda v: y + v * k
    rect(sh, x, y + S(TABLE_H) - S(0.02), S(TABLE_W), S(0.2), "#B5BBC2", None, radius=0.04, name="table edge")
    rect(sh, x, y, S(TABLE_W), S(TABLE_H), TABLE_WHITE, "#D5D9DE", 1.0, radius=0.12 * k, name="table top")
    # track loop (two concentric stadiums = two rails) with sleepers on the straights
    lx, ly, lw_, lh_ = 0.6, 0.75, 6.0, 2.3
    off = 0.15
    for sx in [lx + 1.15 + i * 0.22 for i in range(int((lw_ - 2.3) / 0.22) + 1)]:
        for ty in (ly - 0.02, ly + lh_ - off - 0.04):
            rect(sh, X(sx), Y(ty), S(0.07), S(off + 0.06), "#B9A58F", None, name="sleeper")
    stadium(sh, X(lx), Y(ly), S(lw_), S(lh_), "#7A828C", max(1.5, 3 * k), name="outer rail")
    stadium(sh, X(lx + off), Y(ly + off), S(lw_ - 2 * off), S(lh_ - 2 * off), "#7A828C", max(1.5, 3 * k),
            name="inner rail")
    # train on the top straight
    mini_train(sh, X(2.3), Y(ly - 0.17), S(1.9), name="mini train")
    # station inside the loop
    rect(sh, X(2.0), Y(2.52), S(2.6), S(0.2), "#D5D9DE", "#8C939B", 0.75, name="platform")
    station_building(sh, X(2.7), Y(1.8), S(1.2), S(0.72), name="station")
    # semaphore beside the loop
    piv = semaphore(sh, X(7.6), Y(3.0), S(1.5), sem_state, name="semaphore", on_white=True)
    # signal box
    signal_box_building(sh, X(8.55), Y(2.05), S(1.1), S(0.95), name="signal box", on_white=True)
    pts = {"semaphore": (X(7.6), Y(3.0) - S(1.5) * 0.45), "pivot": piv, "sem_top": Y(3.0) - S(1.5),
           "train": (X(3.25), Y(ly - 0.17)), "station": (X(3.3), Y(1.7)), "box": (X(9.1), Y(2.05)),
           "foot": Y(3.0)}
    if labels:
        def pill(cx, ty, text):
            w = text_w(text, 14, True) + 0.3
            rect(sh, cx - w / 2, ty, w, 0.3, PANEL, None, radius=0.1, name="label pill", text=text, size=14,
                 color=TEXT, bold=True, margins=(0.04, 0, 0.04, 0))
        pill(X(3.25), Y(0.22), "Train")
        pill(X(3.3), Y(1.42), "Station")
        pill(X(7.6), Y(3.1), "Semaphore")
        pill(X(9.1), Y(3.1), "Signal box")
    return pts
