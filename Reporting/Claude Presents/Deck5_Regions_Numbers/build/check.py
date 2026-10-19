"""Per-slide audit: word counts, smallest font, text-overflow estimate, alt text, notes, and
text/text and line/text collisions.

Word-count convention (same as Deck 3): caption + diagram labels + inline chips that carry their own words.
Not counted: title, plain status chips (AGREED and so on), slide numbers, 10 pt credit lines, bare numerals.
Open-question codes such as "(OQ06)" ARE counted unless they sit inside a card chip (slide 25).

usage: python3 check.py [--collisions]
"""
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402
from lib import wrap_lines, text_w  # noqa: E402

PATH = os.path.abspath(os.path.join(HERE, "..", "Regions_Numbers_and_Fleet_Identity_Explained.pptx"))
SKIP_COUNT = ("TITLE", "CHIP", "CHIP LEGEND", "CHIP CARD", "NUM", "CREDIT")
EXEMPT_WORDS = (2, 26)
CREDIT_SLIDE = 26


def walk(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:  # group
            yield from walk(s.shapes)


def words(t):
    """Whitespace tokens that contain a letter (bare numerals are not words)."""
    return len([w for w in t.split() if any(c.isalpha() for c in w)])


def inches(v):
    return Emu(v).inches


def text_box(s):
    """Approximate rendered extent (x0, y0, x1, y1) in inches of the text in shape `s`."""
    tf = s.text_frame
    ml, mr, mt, mb = (inches(tf.margin_left), inches(tf.margin_right), inches(tf.margin_top),
                      inches(tf.margin_bottom))
    x, y, w, h = inches(s.left), inches(s.top), inches(s.width), inches(s.height)
    inner_w = w - ml - mr
    lines = []
    size = 14
    bold = False
    align = "l"
    for p in tf.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        if sizes:
            size = max(sizes)
        bold = any(r.font.bold for r in p.runs) or bold
        if p.alignment is not None:
            align = {1: "l", 2: "c", 3: "r"}.get(int(p.alignment), "l")
        lines += wrap_lines(p.text if p.text else " ", size, inner_w, bold)
    n = len(lines)
    th = n * size * 1.2 / 72
    tw = max((text_w(l, size, bold) for l in lines), default=0)
    anchor = tf.vertical_anchor
    if anchor is None or int(anchor) == 1:
        ty = y + mt
    elif int(anchor) == 3:
        ty = y + (h - th) / 2
    else:
        ty = y + h - mb - th
    if align == "c":
        tx = x + ml + (inner_w - tw) / 2
    elif align == "r":
        tx = x + w - mr - tw
    else:
        tx = x + ml
    return tx, ty, tx + tw, ty + th


def seg_hits_box(p, q, b, pad=0.0):
    """Does segment p-q cross box b=(x0,y0,x1,y1)? (Liang-Barsky)"""
    x0, y0, x1, y1 = b[0] - pad, b[1] - pad, b[2] + pad, b[3] + pad
    dx, dy = q[0] - p[0], q[1] - p[1]
    t0, t1 = 0.0, 1.0
    for pp, qq in ((-dx, p[0] - x0), (dx, x1 - p[0]), (-dy, p[1] - y0), (dy, y1 - p[1])):
        if pp == 0:
            if qq < 0:
                return False
        else:
            r = qq / pp
            if pp < 0:
                if r > t1:
                    return False
                t0 = max(t0, r)
            else:
                if r < t0:
                    return False
                t1 = min(t1, r)
    return True


def collisions(sl):
    """Text/text overlaps and straight lines that cut through text. Returns messages."""
    items = []
    for s in walk(sl.shapes):
        if s.shape_type == 6 or s.has_text_frame is False:
            continue
        if not s.text_frame.text.strip():
            continue
        if s.name in ("CAPTION", "NUM", "TITLE"):
            continue
        items.append((s, text_box(s)))
    msgs = []
    for i in range(len(items)):
        for j in range(i + 1, len(items)):
            a, b = items[i][1], items[j][1]
            ox = min(a[2], b[2]) - max(a[0], b[0])
            oy = min(a[3], b[3]) - max(a[1], b[1])
            if ox > 0.03 and oy > 0.03:
                msgs.append("text overlap: '%s' / '%s'" % (items[i][0].text_frame.text[:22],
                                                           items[j][0].text_frame.text[:22]))
    for s in walk(sl.shapes):
        if s.shape_type == 9 or s.name in ("lead pointer", "leader", "change arrow"):  # LINE / connector
            try:
                p = (inches(s.begin_x), inches(s.begin_y))
                q = (inches(s.end_x), inches(s.end_y))
            except Exception:
                continue
            for it, b in items:
                if seg_hits_box(p, q, b, 0.02):
                    msgs.append("line '%s' crosses text '%s'" % (s.name, it.text_frame.text[:22]))
    return msgs


def main():
    want_col = "--collisions" in sys.argv
    prs = Presentation(PATH)
    print("slides:", len(prs.slides), " size KB:", os.path.getsize(PATH) // 1024)
    print("%3s %-36s %-11s %5s %5s %6s %s" % ("#", "title", "chip", "words", "total", "minpt", "issues"))
    for i, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title.text_frame.text if sl.shapes.title is not None else "(none)"
        chip, body, total, minpt = "-", 0, 0, 99
        issues = []
        if sl.has_notes_slide:
            issues.append("HAS NOTES")
        for s in walk(sl.shapes):
            nm = s.name or ""
            if s.shape_type in (13, 6):
                c = s._element.xpath("./p:nvPicPr/p:cNvPr | ./p:nvGrpSpPr/p:cNvPr")
                if c and not c[0].get("descr"):
                    issues.append("no alt: " + nm)
            if getattr(s, "has_table", False) and s.has_table:
                for r in s.table.rows:
                    for c in r.cells:
                        if c.text_frame.text.strip():
                            total += words(c.text_frame.text)
                            body += words(c.text_frame.text)
                            for p in c.text_frame.paragraphs:
                                for rn in p.runs:
                                    if rn.font.size:
                                        minpt = min(minpt, rn.font.size.pt)
                continue
            if not (s.has_text_frame and s.text_frame.text.strip()):
                continue
            tf = s.text_frame
            t = tf.text
            if nm == "CHIP" and chip == "-":
                chip = t
            if nm not in ("NUM", "CREDIT"):
                total += words(t)
            if nm not in SKIP_COUNT:
                body += words(t)
            sizes = [rn.font.size.pt for p in tf.paragraphs for rn in p.runs if rn.font.size]
            if sizes and nm != "CREDIT":
                minpt = min(minpt, min(sizes))
            if sizes:
                size = max(sizes)
                bold = any(rn.font.bold for p in tf.paragraphs for rn in p.runs)
                inner_w = inches(s.width - tf.margin_left - tf.margin_right)
                inner_h = inches(s.height - tf.margin_top - tf.margin_bottom)
                nlines = sum(len(wrap_lines(p.text if p.text else " ", size, inner_w, bold)) for p in tf.paragraphs)
                need = nlines * size * 1.2 / 72
                longest = max((text_w(wd, size, bold) for p in tf.paragraphs for wd in p.text.split()), default=0)
                if need > inner_h + 0.02:
                    issues.append("overflow? %s '%s' need %.2f have %.2f" % (nm, t[:24], need, inner_h))
                if longest > inner_w + 0.01:
                    issues.append("word too wide: %s '%s'" % (nm, t[:24]))
                x, y, w, h = inches(s.left), inches(s.top), inches(s.width), inches(s.height)
                if x < -0.01 or y < -0.01 or x + w > 13.34 or y + h > 7.51:
                    issues.append("off slide: %s '%s'" % (nm, t[:24]))
        if want_col:
            issues += collisions(sl)
        flag = ""
        if i not in EXEMPT_WORDS:
            flag = " >45!" if body > 45 else (" >25" if body > 25 else "")
        if minpt < 14 and i != CREDIT_SLIDE and minpt != 99:
            issues.append("font below 14 pt")
        print("%3d %-36s %-11s %5d %5d %6s %s%s" % (i, title[:36], chip[:11], body, total,
                                                    ("%g" % minpt) if minpt < 99 else "-", "; ".join(issues), flag))


if __name__ == "__main__":
    main()
