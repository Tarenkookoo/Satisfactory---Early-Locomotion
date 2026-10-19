"""Shape library for the Deck 7 build (python-pptx). Native shapes only.

All coordinates are in inches on a 13.333 x 7.5 canvas. Child shapes inside a group
use slide coordinates (python-pptx keeps chOff == off), so the audit can read them directly.
"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from lxml import etree
from PIL import ImageFont, Image

FONT = 'Arial'
LANG = 'en-GB'
AUTHOR = 'Claude (Early Locomotion project)'
DECK_TITLE = 'After Launch, Part 1: Packs, Locomotives and History, Explained'

BG, PANEL, PANEL_ALT, LINE = '1E2126', '2B2F36', '343A43', '4A515B'
TEXT, MUTED, ORANGE = 'F2F2F2', 'A9B0B8', 'F28C28'
RED, GREEN, AMBER, WATER, ELEC = 'E5484D', '3FB950', 'F2C14E', '4DA3FF', '6CA0DC'
COAL, STEEL, SLEEPER, PURPLE, INK = '1B1D21', '9AA0A6', '6B5B4B', 'B084F5', '101418'
WAGON = '6C7480'
IMG_DIR = None  # set by build script

FONT_REG = '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'
FONT_BOLD = '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'
_fcache = {}


def _pil_font(bold):
    if bold not in _fcache:
        _fcache[bold] = ImageFont.truetype(FONT_BOLD if bold else FONT_REG, 200)
    return _fcache[bold]


def text_w(text, size, bold=False):
    """Width of a single line in inches."""
    return _pil_font(bold).getlength(text) * size / 200.0 / 72.0


def wrap_lines(text, size, bold, width):
    """Greedy word wrap (as PowerPoint does). Returns the list of lines."""
    out = []
    for para in text.split('\n'):
        words = para.split(' ')
        cur = ''
        for wd in words:
            trial = wd if not cur else cur + ' ' + wd
            if text_w(trial, size, bold) <= width or not cur:
                cur = trial
            else:
                out.append(cur)
                cur = wd
        out.append(cur)
    return out


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- deck + slide
def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cp = prs.core_properties
    cp.title = DECK_TITLE
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.language = LANG
    cp.subject = 'Explainer deck for the Early Locomotion project (Satisfactory mod)'
    cp.keywords = ''
    cp.comments = ''
    # theme: Arial everywhere, readable hyperlink colours on a dark background
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode('utf-8')
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r'\1Arial\2', xml)
    xml = re.sub(r'(<a:hlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}(")', r'\1F28C28\2', xml)
    xml = re.sub(r'(<a:folHlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}(")', r'\1F2B36B\2', xml)
    theme._blob = xml.encode('utf-8')
    return prs


def _strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def set_alt(shape, descr, name=None):
    el = shape._element.xpath('./*[1]/p:cNvPr')[0]
    if descr is not None:
        el.set('descr', descr)
    if name:
        el.set('name', name)


def name_it(shape, name):
    el = shape._element.xpath('./*[1]/p:cNvPr')[0]
    el.set('name', name)


def new_slide(prs, title, n, chip=None, title_w=10.4):
    layout = [l for l in prs.slide_layouts if l.name == 'Title Only'][0]
    s = prs.slides.add_slide(layout)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.45), Inches(title_w), Inches(0.8)
    put_text(t, title, size=36, bold=True, color=TEXT, align='l', anchor='m', inset=(0, 0, 0, 0), wrap=True)
    name_it(t, 'Title')
    if chip:
        slide_chip(s, chip)
    s._el_number = n
    return s


STATUS_STYLE = {
    'Agreed': (GREEN, INK, None),
    'Proposed': (PURPLE, INK, None),
    'Provisional': (WATER, INK, None),
    'Undecided': (None, TEXT, MUTED),
    'Background': (LINE, TEXT, None),
}


def slide_chip(slide, status):
    fill, tc, outline = STATUS_STYLE[status]
    sp = rect(slide.shapes, 11.2, 0.5, 1.5, 0.4, fill=fill, line=outline, lw=2, radius=0.1)
    put_text(sp, status.upper(), size=14, bold=True, color=tc, align='c', anchor='m', inset=(0, 0, 0, 0), wrap=False)
    name_it(sp, 'Status chip: ' + status)
    return sp


def inline_chip(S, x, y, status, label, size=14, h=0.34, name=None):
    """Chip beside a single statement. Adds the status word when the label lacks it."""
    fill, tc, outline = STATUS_STYLE[status]
    word = status.upper()
    has_word = status.lower() in label.lower()
    runs = [(label, {'bold': True})] if has_word else [(word + ' · ', {'bold': True}), (label, {})]
    total = sum(text_w(t, size, True) for t, _ in runs)
    w = total + 0.34
    sp = rect(S, x, y, w, h, fill=fill, line=outline, lw=2 if outline else 1.5, radius=0.1)
    put_text(sp, [runs], size=size, color=tc, align='c', anchor='m', inset=(0.05, 0, 0.05, 0), wrap=False)
    name_it(sp, name or ('Inline chip: ' + status))
    return sp, w


def caption(slide, text):
    lines = wrap_lines(text, 20, False, 11.45)
    two = len(lines) > 1
    y, h = (6.2, 0.8) if two else (6.4, 0.6)
    S = slide.shapes
    bar = rect(S, 0.6, y, 12.13, h, fill=PANEL, radius=0.08)
    put_text(bar, text, size=20, color=TEXT, align='l', anchor='m', inset=(0.3, 0.04, 0.2, 0.04))
    name_it(bar, 'Caption')
    rule = rect(S, 0.6, y, 0.08, h, fill=ORANGE)
    name_it(rule, 'Caption rule')
    return two


def slide_number(slide):
    n = slide._el_number
    tb = slide.shapes.add_textbox(Inches(12.0), Inches(7.08), Inches(0.73), Inches(0.3))
    put_text(tb, str(n), size=14, color=MUTED, align='r', anchor='m', inset=(0, 0, 0, 0))
    name_it(tb, 'Slide number')


def credit(S, x, y, w, text, h=0.25, align='l'):
    tb = S.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(tb, text, size=10, color=MUTED, align=align, anchor='t', inset=(0, 0, 0, 0))
    name_it(tb, 'Credit line')
    return tb


# ---------------------------------------------------------------- primitives
def rect(S, x, y, w, h, fill=None, line=None, lw=1.5, radius=0.0, dash=None, shape=None, alpha=None, rot=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sp = S.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(sp)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE and radius:
        sp.adjustments[0] = min(0.5, radius / min(w, h))
    _paint(sp, fill, line, lw, dash, alpha)
    if rot is not None:
        sp.rotation = rot
    return sp


def oval(S, x, y, w, h, fill=None, line=None, lw=1.5, dash=None, alpha=None):
    return rect(S, x, y, w, h, fill, line, lw, dash=dash, shape=MSO_SHAPE.OVAL, alpha=alpha)


def _paint(sp, fill, line, lw, dash, alpha):
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = sp._element.spPr.find(qn('a:solidFill'))[0]
            a = etree.SubElement(clr, qn('a:alpha'))
            a.set('val', str(int(alpha * 1000)))
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = dash
    else:
        sp.line.fill.background()


DASH = MSO_LINE_DASH_STYLE.DASH
DOT = MSO_LINE_DASH_STYLE.ROUND_DOT


def poly(S, pts, fill=None, line=None, lw=1.5, closed=True, dash=None):
    emu = [(Inches(px), Inches(py)) for px, py in pts]
    fb = S.build_freeform(emu[0][0], emu[0][1], scale=1.0)
    fb.add_line_segments(emu[1:], close=closed)
    sp = fb.convert_to_shape()
    _strip_style(sp)
    _paint(sp, fill, line, lw, dash, None)
    return sp


def line(S, x1, y1, x2, y2, color=ORANGE, lw=2.25, dash=None, head=False, tail=False):
    c = S.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    ln = c._element.spPr.find(qn('a:ln'))
    if head:
        e = etree.SubElement(ln, qn('a:headEnd'))
        e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    if tail:
        e = etree.SubElement(ln, qn('a:tailEnd'))
        e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')
    return c


def arrow(S, x1, y1, x2, y2, color=ORANGE, lw=2.25, dash=None, both=False):
    return line(S, x1, y1, x2, y2, color, lw, dash, head=both, tail=True)


def put_text(sp, paras, size=14, color=TEXT, bold=False, align='l', anchor='t',
             inset=(0.08, 0.04, 0.08, 0.04), wrap=True, italic=False, spacing=None, space_after=None):
    """paras: str | list of paragraphs; a paragraph is str | list of runs; a run is str | (str, opts)."""
    tf = sp.text_frame
    body = tf._txBody
    for extra in body.findall(qn('a:p'))[1:]:
        body.remove(extra)
    p0 = body.find(qn('a:p'))
    for child in list(p0):
        if child.tag != qn('a:pPr'):
            p0.remove(child)
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(v) for v in inset]
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = paras.split('\n')
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        if space_after is not None:
            p.space_after = Pt(space_after)
        runs = [para] if isinstance(para, (str, tuple)) else para
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = p.add_run()
            run.text = t
            f = run.font
            f.name = FONT
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', italic)
            f.color.rgb = rgb(o.get('color', color))
            run._r.get_or_add_rPr().set('lang', LANG)
        end = p._p.get_or_add_endParaRPr()
        end.set('lang', LANG)
        end.set('sz', str(int(o.get('size', size) * 100)) if runs else str(int(size * 100)))


def label(S, x, y, w, h, text, size=14, color=TEXT, bold=False, align='l', anchor='m', italic=False,
          inset=(0, 0, 0, 0), wrap=True, name=None):
    tb = S.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(tb, text, size=size, color=color, bold=bold, align=align, anchor=anchor, italic=italic,
             inset=inset, wrap=wrap)
    if name:
        name_it(tb, name)
    return tb


def boxed(S, x, y, w, h, text, size=14, fill=PANEL, line=None, lw=1.5, color=TEXT, bold=False,
          radius=0.08, dash=None, align='c', anchor='m', inset=(0.1, 0.05, 0.1, 0.05), alpha=None):
    sp = rect(S, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius, dash=dash, alpha=alpha)
    put_text(sp, text, size=size, color=color, bold=bold, align=align, anchor=anchor, inset=inset)
    return sp


def group(slide, name, descr):
    g = slide.shapes.add_group_shape()
    set_alt(g, descr, name)
    return g


# ---------------------------------------------------------------- pictures
def picture(S, path, x, y, w, h, descr, focus=0.5, radius=0.08, name=None):
    """Cover-fit crop (never stretch). focus: 0..1 horizontal position of the crop window."""
    with Image.open(path) as im:
        iw, ih = im.size
    box_ar, img_ar = w / h, iw / ih
    pic = S.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if img_ar > box_ar:  # too wide: crop sides
        keep = box_ar / img_ar
        extra = 1 - keep
        pic.crop_left = extra * focus
        pic.crop_right = extra * (1 - focus)
    else:
        keep = img_ar / box_ar
        extra = 1 - keep
        pic.crop_top = extra * 0.5
        pic.crop_bottom = extra * 0.5
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    geom = pic._element.spPr.find(qn('a:prstGeom'))
    av = geom.find(qn('a:avLst'))
    if av is None:
        av = etree.SubElement(geom, qn('a:avLst'))
    for g in list(av):
        av.remove(g)
    gd = etree.SubElement(av, qn('a:gd'))
    gd.set('name', 'adj')
    gd.set('fmla', 'val %d' % int(radius / min(w, h) * 100000))
    set_alt(pic, descr, name or 'Picture')
    return pic


# ---------------------------------------------------------------- reusable objects
def engine(S, x, y, h, n_wheels=3, boiler=1.15, tender=0.0, tank=False, body=ORANGE, outline=None,
           dashed=False, simple=False, alpha=None):
    """Side view, facing right. (x, y) top-left; h is the full height; returns the width used.
    boiler: boiler length in units of h; tender: tender length in units of h (0 = none)."""
    u = h
    ln = outline
    lw = 1.5
    d = DASH if dashed else None
    fill_body = body
    if dashed and alpha is None:
        alpha = 35
    ox = x + (tender + 0.08) * u if tender else x
    cab_w = 0.60
    boil_x0 = ox + cab_w * u
    boil_x1 = boil_x0 + boiler * u
    total = (boil_x1 - x) / u + 0.05
    base = y + u
    r = 0.15 * u
    # tender
    if tender:
        tb = rect(S, x, y + 0.34 * u, tender * u, 0.34 * u, fill=PANEL, line=fill_body if not dashed else (ln or fill_body),
                  lw=lw, dash=d, radius=0.03)
        poly(S, [(x + 0.08 * u, y + 0.34 * u), (x + 0.2 * u, y + 0.22 * u), (x + (tender - 0.2) * u, y + 0.20 * u),
                 (x + (tender - 0.08) * u, y + 0.34 * u)], fill=COAL, line=STEEL, lw=1)
        for cx in (0.28, tender - 0.28):
            oval(S, x + (cx - 0.12) * u, base - 0.24 * u, 0.24 * u, 0.24 * u, fill=PANEL_ALT, line=STEEL, lw=1)
        line(S, x + tender * u, y + 0.62 * u, ox, y + 0.62 * u, color=STEEL, lw=2)
    # wheels
    span = (boil_x1 - 0.12 * u) - (ox + 0.28 * u)
    wx = [ox + 0.28 * u + span * i / max(1, n_wheels - 1) for i in range(n_wheels)] if n_wheels > 1 else [ox + 0.5 * u]
    # chassis plate
    rect(S, ox, y + 0.62 * u, boil_x1 - ox, 0.10 * u, fill=COAL, line=STEEL, lw=1)
    for cx in wx:
        oval(S, cx - r, base - 2 * r, 2 * r, 2 * r, fill=PANEL_ALT, line=STEEL, lw=1.25)
        if not simple:
            oval(S, cx - 0.03 * u, base - r - 0.03 * u, 0.06 * u, 0.06 * u, fill=STEEL)
    if not simple and n_wheels > 1:
        line(S, wx[0], base - r, wx[-1], base - r, color=STEEL, lw=1.5)
    # boiler
    rect(S, boil_x0, y + 0.24 * u, boiler * u, 0.38 * u, fill=fill_body, line=ln, lw=lw, dash=d, radius=0.06, alpha=alpha)
    if tank:
        rect(S, boil_x0 + 0.04 * u, y + 0.38 * u, (boiler - 0.28) * u, 0.26 * u, fill=fill_body, line=COAL, lw=1.25, radius=0.04)
    # smokebox, chimney
    rect(S, boil_x1 - 0.18 * u, y + 0.24 * u, 0.18 * u, 0.38 * u, fill=COAL, line=STEEL, lw=1)
    rect(S, boil_x1 - 0.14 * u, y + 0.06 * u, 0.10 * u, 0.20 * u, fill=COAL, line=STEEL, lw=1)
    rect(S, boil_x1 - 0.18 * u, y + 0.02 * u, 0.18 * u, 0.06 * u, fill=COAL, line=STEEL, lw=1)
    if not simple:
        rect(S, boil_x0 + 0.30 * u, y + 0.14 * u, 0.20 * u, 0.12 * u, fill=fill_body, line=ln, lw=1, dash=d, radius=0.04, alpha=alpha)
    # cab
    rect(S, ox, y + 0.10 * u, cab_w * u, 0.54 * u, fill=fill_body, line=ln, lw=lw, dash=d, alpha=alpha)
    rect(S, ox - 0.03 * u, y + 0.05 * u, (cab_w + 0.06) * u, 0.07 * u, fill=COAL, line=STEEL, lw=1)
    rect(S, ox + 0.14 * u, y + 0.22 * u, 0.26 * u, 0.20 * u, fill=COAL)
    return total * u


def person(S, x, y, h, body=TEXT, head=None):
    """Simple figure: head circle + rounded body. (x, y) top-left; width about 0.7h."""
    hd = 0.34 * h
    w = 0.74 * h
    oval(S, x + (w - hd) / 2, y, hd, hd, fill=head or body)
    rect(S, x, y + 0.40 * h, w, 0.60 * h, fill=body, radius=0.25 * w, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    return w


def rails(S, x, y, length, sleepers=True):
    if sleepers:
        n = int(length / 0.35)
        for i in range(n + 1):
            rect(S, x + i * 0.35, y - 0.06, 0.06, 0.40, fill=SLEEPER)
    rect(S, x, y, length, 0.06, fill=STEEL)
    rect(S, x, y + 0.28, length, 0.06, fill=STEEL)


def wagon(S, x, y, w=1.1, h=0.5, fill=WAGON, line_c=None, dashed=False):
    body = rect(S, x, y, w, h, fill=fill, line=line_c, lw=1.25, radius=0.04, dash=DASH if dashed else None)
    for cx in (x + 0.22, x + w - 0.22):
        oval(S, cx - 0.09, y + h - 0.09, 0.18, 0.18, fill=PANEL_ALT, line=STEEL, lw=1)
    return body


def pack_box_icon(S, x, y, w=0.55, h=0.45, color=ORANGE):
    rect(S, x, y, w, h, fill=None, line=color, lw=1.75, dash=DASH, radius=0.05)
    rect(S, x + w * 0.42, y, w * 0.16, h, fill=color)
    rect(S, x, y + h * 0.4, w, h * 0.16, fill=color)


def question_mark(S, cx, cy, d, color=MUTED, size=28, fill=None, outline=True):
    o = oval(S, cx - d / 2, cy - d / 2, d, d, fill=fill, line=color if outline else None, lw=2)
    put_text(o, '?', size=size, bold=True, color=color, align='c', anchor='m', inset=(0, 0, 0, 0))
    return o


def engine_geom(x, h, boiler=1.15, tender=0.0):
    """Key x positions of an engine() drawn at x with height h."""
    u = h
    ox = x + ((tender + 0.08) * u if tender else 0)
    bx0 = ox + 0.60 * u
    bx1 = bx0 + boiler * u
    return dict(ox=ox, boil_x0=bx0, boil_x1=bx1, width=(bx1 - x) + 0.05 * u, dome_x=bx0 + 0.4 * u)
