"""Audit a built deck: words, font sizes, alt text, notes, transitions, language, text fit and collisions.
usage: python3 audit.py deck.pptx [first_slide last_slide]
"""
import sys, re
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
sys.path.insert(0, __file__.rsplit('/', 1)[0])
from el_lib import text_w, wrap_lines

EXCLUDE_WORDS = ('Title', 'Status chip', 'Slide number', 'Credit line')
INCH = 914400.0


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)


def runs_of(sh):
    out = []
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            out.append([r for r in p.runs])
    return out


def inches(v):
    return (v or 0) / INCH


def text_rect(sh):
    """Estimated rectangle (x0, y0, x1, y1) in inches occupied by the shape's text, plus required height."""
    tf = sh.text_frame
    ml, mt, mr, mb = [inches(v) for v in (tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom)]
    x, y, w, h = inches(sh.left), inches(sh.top), inches(sh.width), inches(sh.height)
    iw = w - ml - mr
    total_h = 0.0
    max_w = 0.0
    wrap = tf.word_wrap is not False
    for p in tf.paragraphs:
        rs = [r for r in p.runs if r.text]
        if not rs:
            continue
        txt = ''.join(r.text for r in rs)
        size = max((r.font.size.pt if r.font.size else 18) for r in rs)
        bold = any(bool(r.font.bold) for r in rs)
        if wrap:
            lines = wrap_lines(txt, size, bold, iw)
        else:
            lines = [txt]
        for ln in lines:
            max_w = max(max_w, text_w(ln, size, bold))
        total_h += len(lines) * size * 1.2 / 72
    anchor = tf.vertical_anchor
    avail = h - mt - mb
    if anchor is None or 'TOP' in str(anchor):
        ty = y + mt
    elif 'MIDDLE' in str(anchor):
        ty = y + mt + (avail - total_h) / 2
    else:
        ty = y + h - mb - total_h
    algn = tf.paragraphs[0].alignment
    s = str(algn)
    tw = min(max_w, iw) if wrap else max_w
    if 'CENTER' in s:
        tx = x + ml + (iw - tw) / 2
    elif 'RIGHT' in s:
        tx = x + ml + iw - tw
    else:
        tx = x + ml
    return (tx, ty, tx + tw, ty + total_h), total_h, avail, (max_w > iw + 0.01 if not wrap else False)


def seg_hits_rect(x1, y1, x2, y2, r, pad=0.03):
    rx0, ry0, rx1, ry1 = r[0] - pad, r[1] - pad, r[2] + pad, r[3] + pad
    n = max(2, int(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5 / 0.02))
    for i in range(n + 1):
        t = i / n
        px, py = x1 + (x2 - x1) * t, y1 + (y2 - y1) * t
        if rx0 < px < rx1 and ry0 < py < ry1:
            return True
    return False


def rects_overlap(a, b, pad=0.0):
    return not (a[2] <= b[0] + pad or b[2] <= a[0] + pad or a[3] <= b[1] + pad or b[3] <= a[1] + pad)


def lum(h):
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4 for v in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(a, b):
    la, lb = lum(a), lum(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


def fill_hex(sh):
    try:
        sp = sh._element.spPr
    except AttributeError:
        return None
    sf = sp.find(qn('a:solidFill'))
    if sf is None:
        return None
    clr = sf[0]
    alpha = clr.find(qn('a:alpha'))
    hx = clr.get('val')
    if alpha is not None:
        a = int(alpha.get('val')) / 100000.0
        bg = (0x1E, 0x21, 0x26)
        c = [int(hx[i:i + 2], 16) for i in (0, 2, 4)]
        hx = ''.join('%02X' % round(c[i] * a + bg[i] * (1 - a)) for i in range(3))
    return hx


def audit(path, lo=None, hi=None):
    prs = Presentation(path)
    problems = []
    rows = []
    for idx, slide in enumerate(prs.slides, start=1):
        if lo and idx < lo or hi and idx > hi:
            continue
        P = lambda msg: problems.append('slide %d: %s' % (idx, msg))
        sx = slide._element
        if sx.find('.//' + qn('p:transition')) is not None or sx.find('.//' + qn('p:timing')) is not None:
            P('has transition or timing')
        if slide.has_notes_slide:
            P('has notes slide')
        words_all = 0
        words_label = 0
        min_size, min_where = 999, ''
        text_rects = []
        connectors = []
        order = [sh.name for sh in slide.shapes]
        for sh in walk(slide.shapes):
            nm = sh.name
            if sh.shape_type in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE) or sh.has_chart if hasattr(sh, 'has_chart') else False:
                pass
            # alt text
            cnv = sh._element.xpath('./*[1]/p:cNvPr')
            descr = cnv[0].get('descr') if cnv else None
            if sh.shape_type in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE) and not descr:
                P('missing alt text on %s (%s)' % (nm, sh.shape_type))
            if getattr(sh, 'has_table', False) and sh.has_table and not descr:
                P('missing alt text on table %s' % nm)
            # bounds
            if sh.shape_type != MSO_SHAPE_TYPE.GROUP:
                x0, y0, x1, y1 = inches(sh.left), inches(sh.top), inches(sh.left + sh.width), inches(sh.top + sh.height)
                if x0 < -0.01 or y0 < -0.01 or x1 > 13.34 or y1 > 7.51:
                    P('off-slide: %s (%.2f,%.2f)-(%.2f,%.2f)' % (nm, x0, y0, x1, y1))
            if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh._element.tag == qn('p:cxnSp'):
                connectors.append((nm, inches(sh.begin_x), inches(sh.begin_y), inches(sh.end_x), inches(sh.end_y)))
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size:
                                    if r.font.size.pt < min_size:
                                        min_size, min_where = r.font.size.pt, nm + ' (table)'
                                if r._r.find(qn('a:rPr')) is not None and r._r.find(qn('a:rPr')).get('lang') != 'en-GB':
                                    P('table run without en-GB: %s' % r.text[:20])
                continue
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text.strip()
            if not txt:
                continue
            toks = [t for t in re.split(r'\s+', txt) if re.search(r'[A-Za-z0-9]', t)]
            words_all += len(toks)
            is_credit = nm.startswith('Credit line')
            if not any(nm.startswith(e) for e in EXCLUDE_WORDS):
                words_label += len(toks)
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    sz = r.font.size.pt if r.font.size else None
                    if sz is None:
                        P('run without explicit size in %s' % nm)
                    elif sz < min_size and not is_credit:
                        min_size, min_where = sz, nm
                    elif is_credit and sz < 10:
                        P('credit line below 10 pt')
                    rpr = r._r.find(qn('a:rPr'))
                    if rpr is None or rpr.get('lang') != 'en-GB':
                        P('run without en-GB lang in %s' % nm)
                    if r.font.name != 'Arial':
                        P('run font not Arial in %s' % nm)
            rect, need, avail, overflow_w = text_rect(sh)
            if need > avail + 0.03:
                P('text overflows its box: %s needs %.2f in, has %.2f (%r)' % (nm, need, avail, txt[:40]))
            if overflow_w:
                P('unwrapped text wider than its box: %s (%r)' % (nm, txt[:40]))
            text_rects.append((nm, rect, txt[:30]))
            # contrast
            fh = fill_hex(sh) if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else None
            bg = fh or '1E2126'
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    try:
                        col = str(r.font.color.rgb)
                    except Exception:
                        continue
                    if fh is None and sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                        bgc = '1E2126'
                    else:
                        bgc = bg
                    cr = contrast(col, bgc)
                    if cr < 4.5 and not (fh is None and sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and False):
                        P('low contrast %.1f: %s (%s on %s)' % (cr, nm, col, bgc))
                    break
                break
        # text vs text overlaps
        for i in range(len(text_rects)):
            for j in range(i + 1, len(text_rects)):
                a, b = text_rects[i], text_rects[j]
                if rects_overlap(a[1], b[1], pad=0.01):
                    P('text overlap: "%s" (%s) and "%s" (%s)' % (a[2], a[0], b[2], b[0]))
        for (cn, x1, y1, x2, y2) in connectors:
            for (tn, r, t) in text_rects:
                if seg_hits_rect(x1, y1, x2, y2, r):
                    P('line "%s" crosses text "%s" (%s)' % (cn, t, tn))
        # reading order
        order_chk = [n for n in order]
        if order_chk and order_chk[0] != 'Title':
            P('first shape is not the title: %s' % order_chk[0])
        if 'Caption' in order_chk and 'Slide number' in order_chk and order_chk.index('Caption') > order_chk.index('Slide number'):
            P('caption after slide number')
        chip = [n for n in order if n.startswith('Status chip')]
        rows.append((idx, ','.join(c.split(': ')[1] for c in chip) or '-', words_label, words_all, min_size, min_where))
    return rows, problems


if __name__ == '__main__':
    path = sys.argv[1]
    lo = int(sys.argv[2]) if len(sys.argv) > 2 else None
    hi = int(sys.argv[3]) if len(sys.argv) > 3 else None
    rows, problems = audit(path, lo, hi)
    print('slide  chip         words(caption+labels)  words(all)  min pt  (where)')
    for r in rows:
        print('%5d  %-11s  %21d  %10d  %6.1f  %s' % r)
    print()
    if problems:
        print('PROBLEMS (%d):' % len(problems))
        for p in problems:
            print(' -', p)
    else:
        print('No problems found.')
