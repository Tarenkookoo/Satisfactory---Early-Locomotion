"""Checks for the built deck: word counts, smallest font, alt text, text fit,
status chips, language tags, notes/transitions and picture aspect ratios.

Usage: python3 check_deck.py DECK.pptx [first_slide last_slide]
"""
import collections
import os
import re
import sys
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import wrap_lines  # noqa: E402

EMU = 914400.0
SKIP_WORDS = ('Status chip', 'Credit', 'Section label', 'Slide number', 'Badge')
EXEMPT = {2, 13, 32, 41, 42}
SMALL_OK = ('Credit', 'Credit table')
STATUS = {'AGREED', 'PROPOSED', 'PROVISIONAL', 'UNDECIDED', 'BACKGROUND'}


def iter_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:      # group
            yield from iter_shapes(s.shapes)


def text_of(s):
    if not s.has_text_frame:
        return ''
    return '\n'.join(p.text for p in s.text_frame.paragraphs)


def fit_problem(s):
    """Rough overflow estimate with Liberation Sans metrics (same as Arial)."""
    if not s.has_text_frame or not text_of(s).strip() or s.rotation:
        return None
    tf = s.text_frame
    bp = tf._txBody.find(qn('a:bodyPr'))
    ins = lambda k, d: int(bp.get(k, d)) / EMU
    w = s.width / EMU - ins('lIns', 91440) - ins('rIns', 91440)
    h = s.height / EMU - ins('tIns', 45720) - ins('bIns', 45720)
    geom = s._element.find('.//' + qn('a:prstGeom'))
    if geom is not None and geom.get('prst') == 'roundRect':
        gd = geom.find('.//' + qn('a:gd'))
        adj = int(gd.get('fmla').split()[1]) if gd is not None else 16667
        inset = 0.29289 * min(s.width, s.height) / EMU * adj / 100000
        w -= 2 * inset
        h -= 2 * inset
    elif geom is not None and geom.get('prst') == 'ellipse':
        w *= 0.7071
        h *= 0.7071
    total = 0.0
    widest = 0.0
    for p in tf.paragraphs:
        size = max([r.font.size.pt for r in p.runs if r.font.size] or [18])
        bold = any(r.font.bold for r in p.runs)
        txt = ''.join(r.text for r in p.runs)
        if s.text_frame.word_wrap is False:
            lines = [txt]
        else:
            lines = wrap_lines(txt, size, max(w, 0.05), bold)
        from lib import text_width_in
        for ln in lines:
            widest = max(widest, text_width_in(ln, size, bold) - w)
            for word in ln.split():
                if text_width_in(word, size, bold) > w + 0.01:
                    return 'word "%s" wider than box' % word
        sp = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        total += len(lines) * size * 1.17 * sp / 72
        if p.space_after is not None:
            total += p.space_after.pt / 72
    if total > h + 0.04:
        return 'height %.2f needed, %.2f available' % (total, h)
    return None


def main(path, first=1, last=999):
    prs = Presentation(path)
    z = zipfile.ZipFile(path)
    names = z.namelist()
    notes = [n for n in names if 'notesSlide' in n]
    print('slides:', len(prs.slides), '| notes parts:', len(notes),
          '| file size: %.2f MB' % (os.path.getsize(path) / 1e6))
    trans = [n for n in names if n.startswith('ppt/slides/slide') and n.endswith('.xml')
             and (b'<p:transition' in z.read(n) or b'<p:timing' in z.read(n))]
    print('slides with transitions/animations:', trans or 'none')
    tags = collections.Counter()
    problems = 0
    for idx, slide in enumerate(prs.slides, 1):
        if not first <= idx <= last:
            continue
        words_a = words_b = 0
        min_size = 99
        min_where = ''
        issues = []
        chip = None
        for s in iter_shapes(slide.shapes):
            name = s.name
            if name == 'Status chip':
                chip = text_of(s)
            txt = text_of(s)
            if s.has_text_frame:
                is_title = s.is_placeholder and s.placeholder_format.type == 1
                # reference codes such as "(OQ03)" are not counted as words
                n = len([w for w in txt.split() if not re.fullmatch(r'\(?OQ\d+\)?[.,:]?', w)])
                if is_title:
                    words_b += n
                elif name.startswith('Inline chip') and txt.strip().upper() in STATUS:
                    pass        # a bare status tag, like the corner chip
                elif not name.startswith(SKIP_WORDS):
                    words_a += n
                    words_b += n
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.text.strip():
                            if r.font.size.pt < min_size and not name.startswith(SMALL_OK):
                                min_size, min_where = r.font.size.pt, '%s "%s"' % (name, r.text[:20])
                        if r._r.find(qn('a:rPr')) is None or r._r.find(qn('a:rPr')).get('lang') != 'en-GB':
                            issues.append('run without en-GB: %r' % r.text[:20])
                prob = fit_problem(s)
                if prob:
                    issues.append('FIT %s "%s": %s' % (name, txt[:30].replace('\n', ' '), prob))
                if s.is_placeholder and not txt.strip():
                    issues.append('empty placeholder')
            if s.shape_type == 13:     # picture
                if not s._element[0][0].get('descr'):
                    issues.append('picture without alt: ' + name)
                img = Image.open(__import__('io').BytesIO(s.image.blob))
                iw, ih = img.size
                cw = iw * (1 - s.crop_left - s.crop_right)
                ch = ih * (1 - s.crop_top - s.crop_bottom)
                if abs((cw / ch) / (s.width / s.height) - 1) > 0.01:
                    issues.append('stretched picture ' + name)
            if s.shape_type == 6 and not s._element[0][0].get('descr'):
                issues.append('group without alt: ' + name)
            if s.has_text_frame is False and s.shape_type == 19 and not s._element[0][0].get('descr'):
                issues.append('table without alt: ' + name)
        # top-level shapes with no text must be grouped (alt) or decorative
        for s in slide.shapes:
            if s.shape_type in (1, 9, 5, None) and not text_of(s).strip():
                el = s._element[0][0]
                if not el.get('descr') and b'decorative' not in __import__('lxml').etree.tostring(el):
                    issues.append('ungrouped shape without alt/decorative: ' + s.name)
        tags[chip or '(none)'] += 1
        cap = 45 if idx not in EXEMPT else 999
        flag = ''
        if words_a > cap:
            flag = '  <-- OVER 45'
        elif words_a > 25 and idx not in EXEMPT:
            flag = '  (over 25 aim)'
        print('slide %2d | chip %-12s | words(caption+labels) %3d | incl. title %3d | min pt %4.1f %s%s'
              % (idx, chip or '-', words_a, words_b, min_size if min_size < 99 else 0, min_where if min_size < 14 else '', flag))
        if min_size < 14:
            issues.append('text below 14 pt: ' + min_where)
        for i in issues:
            print('        ! ' + i)
            problems += 1
    print('chip tally:', dict(tags))
    print('problems:', problems)


if __name__ == '__main__':
    a = sys.argv[1:]
    main(a[0], *(int(v) for v in a[1:3]))
