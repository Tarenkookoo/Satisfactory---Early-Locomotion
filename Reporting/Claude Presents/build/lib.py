"""Shape library for the "Early Locomotion, Explained" deck.

All coordinates are in inches on a 13.333 x 7.5 in slide. Every helper takes a
`shapes` collection (slide.shapes or a group's .shapes) so the same objects can
be drawn on a slide or inside a group. Style tokens come from OUTLINE.md.
"""
import math
import os
import re

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE, MSO_PATTERN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                      # Reporting/Claude Presents
IMG_DIR = os.path.join(ROOT, 'images')
MANIFEST = os.path.join(ROOT, 'ASSET_MANIFEST.md')
FONT_DIR = '/usr/share/fonts/truetype/liberation'

# ---------------------------------------------------------------- tokens
C = dict(
    bg='1E2126', panel='2B2F36', panel_alt='343A43', line='4A515B',
    text='F2F2F2', muted='A9B0B8', orange='F28C28', red='E5484D',
    green='3FB950', amber='F2C14E', water='4DA3FF', electric='6CA0DC',
    coal='1B1D21', steel='9AA0A6', sleeper='6B5B4B',
    # derived helpers (not new style tokens): dark ink for text on bright
    # fills, a darker red for white-text STOP pills (5.8:1), post white,
    # wagon grey from RESEARCH_NOTES, coal outline from SPEC common colours.
    ink='101418', red_dark='B3262B', post='E8E8E8', wagon='6C7480',
    coal_line='6B7078', purple='B084F5', ballast='7A7F86',
)

CHIPS = {
    'Agreed':      dict(fill=C['green'], text=C['ink'], outline=None),
    'Proposed':    dict(fill=C['purple'], text=C['ink'], outline=None),
    'Provisional': dict(fill=C['water'], text=C['ink'], outline=None),
    'Undecided':   dict(fill=None, text=C['text'], outline=C['muted']),
    'Background':  dict(fill=C['line'], text=C['text'], outline=None),
}

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
RIGHT = SLIDE_W - MARGIN          # 12.733
DECOR_URI = '{C183D7F6-B498-43B3-948B-1728B52AA6E4}'
DECOR_NS = 'http://schemas.microsoft.com/office/drawing/2017/decorative'


def I(v):
    return Emu(int(round(v * 914400)))


def rgb(h):
    return RGBColor.from_string(h)


def blend(fg, bg=C['bg'], a=0.25):
    """Solid colour equal to `fg` at opacity `a` over `bg`."""
    f = [int(fg[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(bg[i:i + 2], 16) for i in (0, 2, 4)]
    return ''.join('%02X' % round(a * f[k] + (1 - a) * b[k]) for k in range(3))


# ---------------------------------------------------------------- text measure
_fonts = {}


def font(size, bold=False):
    key = (round(size * 4), bold)
    if key not in _fonts:
        name = 'LiberationSans-Bold.ttf' if bold else 'LiberationSans-Regular.ttf'
        _fonts[key] = ImageFont.truetype(os.path.join(FONT_DIR, name), max(1, round(size * 4)))
    return _fonts[key]


def text_width_in(s, size, bold=False):
    return font(size, bold).getlength(s) / 4 / 72


def wrap_lines(s, size, width_in, bold=False):
    lines = []
    for para in s.split('\n'):
        words = para.split(' ')
        cur = ''
        for w in words:
            t = (cur + ' ' + w).strip()
            if text_width_in(t, size, bold) <= width_in or not cur:
                cur = t
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


# ---------------------------------------------------------------- xml helpers
def _strip_style(shape):
    st = shape._element.find(qn('p:style'))
    if st is not None:
        shape._element.remove(st)


def cnvpr(shape):
    return shape._element[0][0]


def set_alt(shape, text):
    cnvpr(shape).set('descr', text)
    return shape


def mark_decorative(shape):
    c = cnvpr(shape)
    ext_lst = c.find(qn('a:extLst'))
    if ext_lst is None:
        ext_lst = etree.SubElement(c, qn('a:extLst'))
    ext = etree.SubElement(ext_lst, qn('a:ext'), uri=DECOR_URI)
    etree.SubElement(ext, '{%s}decorative' % DECOR_NS, nsmap={'adec': DECOR_NS}, val='1')
    return shape


def _fill(shape, fill):
    if fill is None:
        shape.fill.background()
    elif isinstance(fill, tuple) and fill[0] == 'hatch':
        _, fg, bgc = fill
        shape.fill.patterned()
        shape.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
        shape.fill.fore_color.rgb = rgb(fg)
        shape.fill.back_color.rgb = rgb(bgc)
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)


def _line(shape, line, lw=1.5, dash=None):
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
        if dash:
            shape.line.dash_style = {'dash': MSO_LINE.DASH, 'dot': MSO_LINE.ROUND_DOT,
                                     'sqdot': MSO_LINE.SQUARE_DOT,
                                     'longdash': MSO_LINE.LONG_DASH}[dash]


def _arrowheads(shape, head, size='med'):
    """head: None, 'end', 'start' or 'both'."""
    if not head:
        return
    ln = shape.line._get_or_add_ln()
    if head in ('start', 'both'):
        etree.SubElement(ln, qn('a:headEnd'), type='triangle', w=size, len=size)
    if head in ('end', 'both'):
        etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w=size, len=size)


# ---------------------------------------------------------------- text
def _apply_text(tf, text, size=18, bold=False, color=C['text'], align='l',
                anchor='t', margins=(0.08, 0.04, 0.08, 0.04), spacing=None,
                space_after=0, italic=False):
    """text: str (\\n = new paragraph) or list of paragraphs; a paragraph is a
    str or a list of runs; a run is a str or (str, dict(bold,color,size,italic))."""
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    l, t, r, b = margins
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = I(l), I(t), I(r), I(b)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    paras = text.split('\n') if isinstance(text, str) else text
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        if space_after and i < len(paras) - 1:
            p.space_after = Pt(space_after)
        runs = [para] if isinstance(para, (str, tuple)) else para
        for run in runs:
            s, opt = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = s
            f = r.font
            f.name = 'Arial'
            f.size = Pt(opt.get('size', size))
            f.bold = opt.get('bold', bold)
            f.italic = opt.get('italic', italic)
            f.color.rgb = rgb(opt.get('color', color))


def box(shapes, x, y, w, h, fill=None, line=None, lw=1.5, radius=None,
        shape=None, text=None, size=18, bold=False, color=C['text'], align='c',
        anchor='m', margins=(0.08, 0.04, 0.08, 0.04), dash=None, rot=0,
        spacing=None, name=None, deco=False, adj=None, space_after=0):
    """Autoshape (rectangle by default) with optional text."""
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = shapes.add_shape(st, I(x), I(y), I(w), I(h))
    _strip_style(s)
    _fill(s, fill)
    _line(s, line, lw, dash)
    if radius:
        s.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    if adj:
        for k, v in enumerate(adj):
            s.adjustments[k] = v
    if rot:
        s.rotation = rot
    if name:
        s.name = name
    if text is not None:
        _apply_text(s.text_frame, text, size, bold, color, align, anchor, margins, spacing, space_after)
    if deco:
        mark_decorative(s)
    return s


def label(shapes, x, y, w, h, text, size=14, bold=False, color=C['text'],
          align='l', anchor='t', margins=(0, 0, 0, 0), spacing=None, name=None,
          italic=False, space_after=0):
    s = shapes.add_textbox(I(x), I(y), I(w), I(h))
    _apply_text(s.text_frame, text, size, bold, color, align, anchor, margins, spacing,
                space_after, italic)
    if name:
        s.name = name
    return s


def oval(shapes, cx, cy, d, fill=None, line=None, lw=1.5, dh=None, **kw):
    return box(shapes, cx - d / 2, cy - (dh or d) / 2, d, dh or d, fill=fill, line=line,
               lw=lw, shape=MSO_SHAPE.OVAL, **kw)


def line(shapes, x1, y1, x2, y2, color=C['orange'], lw=2.25, head=None, dash=None, name=None,
         hs='med'):
    s = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _strip_style(s)
    _line(s, color, lw, dash)
    _arrowheads(s, head, hs)
    if name:
        s.name = name
    return s


def arrow(shapes, x1, y1, x2, y2, color=C['orange'], lw=2.25, head='end', dash=None, name=None,
          hs='med'):
    return line(shapes, x1, y1, x2, y2, color, lw, head, dash, name, hs)


def poly(shapes, pts, fill=None, line=None, lw=1.5, closed=True, head=None, dash=None, name=None,
         hs='med'):
    e = [(int(round(px * 914400)), int(round(py * 914400))) for px, py in pts]
    fb = shapes.build_freeform(e[0][0], e[0][1], scale=1.0)
    fb.add_line_segments(e[1:], close=closed)
    s = fb.convert_to_shape()
    _strip_style(s)
    _fill(s, fill if closed else None)
    _line(s, line, lw, dash)
    _arrowheads(s, head, hs)
    if name:
        s.name = name
    return s


def polyline(shapes, pts, color=C['orange'], lw=2.25, head=None, dash=None, name=None, hs='med'):
    return poly(shapes, pts, None, color, lw, closed=False, head=head, dash=dash, name=name, hs=hs)


def arc_pts(cx, cy, rx, ry, a0, a1, n=24):
    """Points on an ellipse; angles in degrees, 0 = right, 90 = up (screen)."""
    out = []
    for k in range(n + 1):
        a = math.radians(a0 + (a1 - a0) * k / n)
        out.append((cx + rx * math.cos(a), cy - ry * math.sin(a)))
    return out


def curve_arrow(shapes, cx, cy, rx, ry, a0, a1, color=C['orange'], lw=2.25, head='end', dash=None,
                hs='lg'):
    return polyline(shapes, arc_pts(cx, cy, rx, ry, a0, a1), color, lw, head, dash, hs=hs)


def group(shapes, name, alt=None):
    g = shapes.add_group_shape()
    g.name = name
    if alt:
        set_alt(g, alt)
    return g


def rot_pt(p, c, ang_ccw):
    a = math.radians(ang_ccw)
    dx, dy = p[0] - c[0], p[1] - c[1]
    return (c[0] + dx * math.cos(a) + dy * math.sin(a), c[1] - dx * math.sin(a) + dy * math.cos(a))


# ---------------------------------------------------------------- manifest
def _md_rows(md, header_start):
    i = md.index(header_start)
    rows = []
    for ln in md[i:].splitlines()[2:]:
        if not ln.startswith('|'):
            break
        rows.append([c.strip() for c in ln.strip().strip('|').split('|')])
    return rows


def load_manifest():
    md = open(MANIFEST, encoding='utf-8').read()
    imgs = {}
    for r in _md_rows(md, '| Slot | Bundle path | Original path (repo-relative) | What it shows (also the alt text) | Credit'):
        imgs[r[0]] = dict(file=r[1].strip('`').replace('images/', ''), alt=r[3], credit_full=r[4], url=r[5], kind='photo')
    for r in _md_rows(md, '| Slot | Slide credit text | Link |'):
        imgs[r[0]]['credit'] = r[1]
        assert imgs[r[0]]['url'] == r[2], r[0]
    for r in _md_rows(md, '| Slot | Bundle path | Original path (repo-relative) | What it shows (also the alt text) | Slides |'):
        imgs[r[0]] = dict(file=r[1].strip('`').replace('images/', ''), alt=r[3],
                          credit='Early Locomotion project render. Early test model, not final art.',
                          url='', kind='render', slides=r[4])
    m = re.search(r'Credit line to print: "([^"]+)"', md)
    for k, v in imgs.items():
        if v['kind'] == 'render':
            v['credit'] = m.group(1)
    return imgs


IMAGES = load_manifest()


# ---------------------------------------------------------------- pictures
def picture(shapes, slot, x, y, w, h, focus=(0.5, 0.5), radius=0.08, alt=None, name=None):
    """Place a bundle image to cover the box (crop, never stretch)."""
    info = IMAGES[slot]
    path = os.path.join(IMG_DIR, info['file'])
    iw, ih = Image.open(path).size
    pic = shapes.add_picture(path, I(x), I(y), I(w), I(h))
    box_ar, img_ar = w / h, iw / ih
    if img_ar > box_ar:            # too wide: crop left/right
        keep = box_ar / img_ar
        extra = 1 - keep
        left = min(max(focus[0] - keep / 2, 0), extra)
        pic.crop_left, pic.crop_right = left, extra - left
    elif img_ar < box_ar:          # too tall: crop top/bottom
        keep = img_ar / box_ar
        extra = 1 - keep
        top = min(max(focus[1] - keep / 2, 0), extra)
        pic.crop_top, pic.crop_bottom = top, extra - top
    if radius:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = pic._element.spPr.find(qn('a:prstGeom'))
        av = geom.find(qn('a:avLst'))
        if av is None:
            av = etree.SubElement(geom, qn('a:avLst'))
        for gd in list(av):
            av.remove(gd)
        etree.SubElement(av, qn('a:gd'), name='adj', fmla='val %d' % round(radius / min(w, h) * 100000))
    pic.name = name or ('Picture ' + slot)
    set_alt(pic, alt or info['alt'])
    return pic


def credit_text(slot, prefix=None):
    info = IMAGES[slot]
    if info['kind'] == 'photo':
        base = 'Photo: %s, Wikimedia Commons.' % info['credit']
    else:
        base = info['credit']
    return (prefix + ' ' + base) if prefix else base


def credit(shapes, x, y, w, text, h=None, align='l'):
    lines = len(wrap_lines(text, 10, w))
    s = label(shapes, x, y, w, h or (0.03 + 0.17 * lines), text, size=10, color=C['muted'],
              align=align, name='Credit')
    return s


def photo(shapes, slot, x, y, w, h, prefix=None, focus=(0.5, 0.5), credit_w=None,
          credit_align='l', alt=None, extra=None):
    """Picture plus its credit line directly underneath."""
    p = picture(shapes, slot, x, y, w, h, focus=focus, alt=alt)
    txt = credit_text(slot, prefix)
    if extra:
        txt = txt + '\n' + extra
    credit(shapes, x, y + h + 0.03, credit_w or w, txt, align=credit_align)
    return p


# ---------------------------------------------------------------- chips etc.
def chip(shapes, kind, x, y, w=1.7, h=0.4, size=14, text=None, name=None):
    st = CHIPS[kind]
    s = box(shapes, x, y, w, h, fill=st['fill'], line=st['outline'], lw=2, radius=h / 2,
            text=(text or kind).upper(), size=size, bold=True, color=st['text'],
            margins=(0.06, 0.02, 0.06, 0.02), name=name or ('Chip ' + kind))
    return s


def inline_chip(shapes, kind, x, y, text=None, h=0.36, size=14, w=None):
    t = (text or kind).upper()
    w = w or chip_w(t, size)
    return chip(shapes, kind, x, y, w, h, size, t, name='Inline chip ' + kind)


def chip_w(kind_or_text, size=14):
    # text + margins + the rounded-rectangle text inset at both ends
    return text_width_in(kind_or_text.upper(), size, True) + 0.4


def pill(shapes, x, y, w, h, text, fill, color=C['ink'], size=14, bold=True, line=None, lw=1.5, dash=None):
    return box(shapes, x, y, w, h, fill=fill, line=line, lw=lw, radius=h / 2, text=text,
               size=size, bold=bold, color=color, margins=(0.06, 0.02, 0.06, 0.02), dash=dash)


def badge(shapes, cx, cy, n, d=0.42, size=16):
    return oval(shapes, cx, cy, d, fill=C['orange'], text=str(n), size=size, bold=True,
                color=C['ink'], margins=(0, 0, 0, 0))


def state_pill(shapes, x, y, state, w=0.95, h=0.32, size=14):
    if state == 'stop':
        return pill(shapes, x, y, w, h, 'STOP', C['red_dark'], C['text'], size)
    return pill(shapes, x, y, w, h, 'CLEAR', C['green'], C['ink'], size)


# ---------------------------------------------------------------- railway objects
def rails(shapes, x, y, length, sleepers=True, gap=0.28, thick=0.06, pitch=0.35, name='Track'):
    """Top-down track per RESEARCH_NOTES: two steel rails 0.28 in apart, sleepers every 0.35 in."""
    if sleepers:
        n = int(length / pitch)
        x0 = x + (length - n * pitch) / 2 + pitch / 2 - 0.03
        for k in range(n):
            box(shapes, x0 + k * pitch, y - 0.06, 0.06, gap + 0.06 + 0.06 + thick - 0.06,
                fill=C['sleeper'])
    box(shapes, x, y - thick / 2, length, thick, fill=C['steel'])
    box(shapes, x, y + gap - thick / 2, length, thick, fill=C['steel'])


def track_group(shapes, x, y, length, alt='Railway track: two rails on sleepers.', **kw):
    g = group(shapes, 'Track', alt)
    rails(g.shapes, x, y, length, **kw)
    return g


def _arm_pts(hx, hy, L, t, ang, side, fishtail=False):
    if fishtail:
        n = 0.16
        loc = [(0, -t / 2), (L, -t / 2), (L - n, 0), (L, t / 2), (0, t / 2)]
    else:
        loc = [(0, -t / 2), (L, -t / 2), (L, t / 2), (0, t / 2)]
    out = []
    for (u, v) in loc:
        p = rot_pt((hx + u, hy + v), (hx, hy), ang)
        out.append((hx - (p[0] - hx), p[1]) if side < 0 else p)
    return out


def _stripe_pts(hx, hy, L, t, ang, side, u0, u1):
    loc = [(u0, -t / 2), (u1, -t / 2), (u1, t / 2), (u0, t / 2)]
    out = []
    for (u, v) in loc:
        p = rot_pt((hx + u, hy + v), (hx, hy), ang)
        out.append((hx - (p[0] - hx), p[1]) if side < 0 else p)
    return out


def signal(shapes, x, y_base, state='stop', kind='stop', h=1.1, side=1, arm_len=0.7,
           arm_t=0.14, label_pos='left', show_label=True, alt=None, name=None, scale=1.0):
    """Semaphore on a post. x = post centre, y_base = foot of the post.
    The arm stays red (stop) or amber (distant); state is shown by angle, lamp and a label."""
    h, arm_len, arm_t = h * scale, arm_len * scale, arm_t * scale
    kind_word = 'distant' if kind == 'distant' else 'stop'
    if alt is None:
        if kind == 'distant':
            alt = 'Distant signal, arm %s: %s.' % ('flat' if state == 'stop' else 'raised',
                                                   'be ready to stop' if state == 'stop' else 'next signal clear')
        else:
            alt = 'Semaphore %s signal, arm %s: %s.' % (kind_word, 'flat' if state == 'stop' else 'raised',
                                                       'STOP' if state == 'stop' else 'CLEAR')
    g = group(shapes, name or ('Signal ' + state), alt)
    s = g.shapes
    pw = 0.12 * scale
    box(s, x - 0.17 * scale, y_base - 0.07 * scale, 0.34 * scale, 0.07 * scale, fill=C['steel'])
    box(s, x - pw / 2, y_base - h, pw, h, fill=C['post'])
    oval(s, x, y_base - h - 0.04 * scale, 0.14 * scale, fill=C['post'])
    hx, hy = x + side * pw / 2, y_base - h + 0.1 * scale + arm_t / 2 + 0.06 * scale
    ang = 0 if state == 'stop' else 45
    if kind == 'distant':
        poly(s, _arm_pts(hx, hy, arm_len, arm_t + 0.02, ang, side, fishtail=True),
             fill=C['amber'], line=C['ink'], lw=0.75)
        poly(s, _stripe_pts(hx, hy, arm_len, arm_t + 0.02, ang, side, 0.45 * arm_len, 0.55 * arm_len),
             fill=C['ink'])
    else:
        poly(s, _arm_pts(hx, hy, arm_len, arm_t, ang, side), fill=C['red'], line=C['ink'], lw=0.75)
        poly(s, _stripe_pts(hx, hy, arm_len, arm_t, ang, side, 0.62 * arm_len, 0.74 * arm_len),
             fill='FFFFFF')
    lamp = C['red'] if state == 'stop' else C['green']
    if kind == 'distant' and state == 'stop':
        lamp = C['amber']
    lx = x - side * (pw / 2 + 0.11 * scale)
    oval(s, lx, hy, 0.17 * scale, fill=lamp, line=C['ink'], lw=0.75)
    if show_label:
        pw_ = 0.95
        if label_pos == 'left':
            lx0 = x - pw_ - 0.25 * scale
            ly0 = y_base - h * 0.55
        elif label_pos == 'right':
            lx0 = x + 0.2 * scale
            ly0 = y_base - h * 0.55
        else:   # below
            lx0 = x - pw_ / 2
            ly0 = y_base + 0.08
        state_pill(s, lx0, ly0, state)
    return g


def train(shapes, x, y, n_wagons=2, scale=1.0, alt='A steam train: an engine pulling wagons.',
          name='Train', loco_color=C['orange'], tender=False):
    """Steam train facing right. x = left end, y = rail level (wheel bottoms)."""
    g = group(shapes, name, alt)
    s = g.shapes
    k = scale
    cx = x
    for _ in range(n_wagons):
        box(s, cx, y - 0.57 * k, 0.8 * k, 0.45 * k, fill=C['wagon'], line=C['ink'], lw=0.75)
        for wx in (0.2, 0.6):
            oval(s, cx + wx * k, y - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
        box(s, cx + 0.8 * k, y - 0.25 * k, 0.08 * k, 0.05 * k, fill=C['steel'])
        cx += 0.88 * k
    if tender:
        box(s, cx, y - 0.6 * k, 0.6 * k, 0.48 * k, fill=loco_color, line=C['ink'], lw=0.75, radius=0.04 * k)
        box(s, cx + 0.05 * k, y - 0.68 * k, 0.5 * k, 0.12 * k, fill=C['coal'])
        for wx in (0.15, 0.45):
            oval(s, cx + wx * k, y - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
        box(s, cx + 0.6 * k, y - 0.25 * k, 0.08 * k, 0.05 * k, fill=C['steel'])
        cx += 0.68 * k
    lx = cx
    box(s, lx, y - 0.65 * k, 1.0 * k, 0.55 * k, fill=loco_color, line=C['ink'], lw=0.75, radius=0.06 * k)
    box(s, lx + 0.03 * k, y - 0.62 * k, 0.26 * k, 0.2 * k, fill=C['coal'], name='Bunker')
    box(s, lx + 0.34 * k, y - 0.58 * k, 0.14 * k, 0.14 * k, fill=C['panel'])
    box(s, lx + 0.78 * k, y - 0.93 * k, 0.15 * k, 0.3 * k, fill=C['coal'], name='Chimney')
    for wx in (0.2, 0.5, 0.8):
        oval(s, lx + wx * k, y - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
    return g


def train_width(n_wagons=2, scale=1.0, tender=False):
    return (n_wagons * 0.88 + (0.68 if tender else 0) + 1.0) * scale


def block_band(shapes, x, y, w, state, text, h=0.38, size=16):
    col = C['red'] if state == 'occupied' else C['green']
    return box(shapes, x, y, w, h, fill=blend(col, C['bg'], 0.25), line=col, lw=1.5,
               text=text, size=size, bold=True, color=C['text'], margins=(0.04, 0, 0.04, 0))


def person(shapes, cx, top, hgt=0.7, color=C['post'], name='Person'):
    g = group(shapes, name, 'A signaller.')
    oval(g.shapes, cx, top + hgt * 0.16, hgt * 0.3, fill=color)
    box(g.shapes, cx - hgt * 0.24, top + hgt * 0.36, hgt * 0.48, hgt * 0.64, fill=color, radius=hgt * 0.2)
    return g


def lever_row(shapes, x, y, colours, w=0.12, h=0.9, gap=0.22):
    for k, col in enumerate(colours):
        lx = x + k * (w + gap)
        if col == 'points':
            box(shapes, lx, y, w, h, fill='5A6068', line=C['ink'], lw=1, radius=0.03)
        else:
            box(shapes, lx, y, w, h, fill=col, line=C['ink'], lw=0.75, radius=0.03)
        oval(shapes, lx + w / 2, y + h, w * 1.8, fill=C['steel'])
        box(shapes, lx - 0.02, y - 0.06, w + 0.04, 0.12, fill=C['steel'], radius=0.03)


def tank(shapes, x, y, w, h, pct, fill, outline=C['text'], lw=2, name='Tank'):
    box(shapes, x, y, w, h, fill=C['panel'], line=outline, lw=lw, radius=0.08)
    fh = (h - 0.08) * pct
    if fh > 0.02:
        box(shapes, x + 0.04, y + h - 0.04 - fh, w - 0.08, fh, fill=fill, radius=0.05)


def bar_row(shapes, x, y, label_text, track_w, pct, color, label_w=2.2, h=0.32, size=16,
            hatch=False, track_col=C['panel_alt']):
    label(shapes, x, y - 0.04, label_w - 0.1, h + 0.08, label_text, size=size, anchor='m')
    box(shapes, x + label_w, y, track_w, h, fill=track_col, radius=0.06)
    f = ('hatch', color, C['panel_alt']) if hatch else color
    box(shapes, x + label_w, y, max(0.05, track_w * pct), h, fill=f, radius=0.06,
        line=color if hatch else None, lw=1.5)


def box_building(shapes, x, y, w, h, outline=C['orange'], with_person=True, roof=True, windows=True):
    """Simple signal box: body, roof, window strip, signaller."""
    if roof:
        poly(shapes, [(x - 0.1, y), (x + w / 2, y - h * 0.3), (x + w + 0.1, y)], fill=C['line'], line=outline, lw=1.5)
    box(shapes, x, y, w, h, fill=C['panel'], line=outline, lw=1.5)
    if windows:
        box(shapes, x + 0.1, y + 0.12, w - 0.2, h * 0.28, fill=blend(C['water'], C['panel'], 0.35),
            line=C['steel'], lw=0.75)
    if with_person:
        person(shapes, x + w * 0.25, y + h * 0.3, h * 0.55)


# ---------------------------------------------------------------- slide base
class Base:
    pass


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    master = prs.slide_master
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = rgb(C['bg'])
    # theme: Arial fonts, readable hyperlink colours on the dark background
    theme = master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode('utf-8')
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*"', r'\1Arial"', xml)
    xml = re.sub(r'<a:hlink>.*?</a:hlink>', '<a:hlink><a:srgbClr val="7DB8FF"/></a:hlink>', xml, flags=re.S)
    xml = re.sub(r'<a:folHlink>.*?</a:folHlink>', '<a:folHlink><a:srgbClr val="C9A7FF"/></a:folHlink>', xml, flags=re.S)
    theme._blob = xml.encode('utf-8')
    # keep only the "Title Only" layout
    keep = [l for l in prs.slide_layouts if l.name == 'Title Only'][0]
    for l in list(prs.slide_layouts):
        if l is not keep:
            prs.slide_layouts.remove(l)
    cp = prs.core_properties
    cp.title = 'Early Locomotion, Explained'
    cp.author = 'Early Locomotion project'
    cp.subject = 'A read-alone explainer of the Early Locomotion steam railway mod for Satisfactory'
    cp.keywords = 'Satisfactory; Early Locomotion; steam railway; explainer'
    cp.language = 'en-GB'
    cp.last_modified_by = 'Early Locomotion project'
    cp.comments = 'Plans, not a finished mod. Status tags follow Design/GAME_DESIGN.md.'
    return prs


def caption_lines(caption):
    return len(wrap_lines(caption, 20, 12.13 - 0.3 - 0.2))


def add_base(prs, n, section, title, chip_kind, caption, title_size=36):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    b = Base()
    b.slide, b.shapes, b.n = slide, slide.shapes, n
    # title placeholder (first in reading order)
    t = slide.shapes.title
    title_w = (RIGHT - 1.7 - 0.25 - MARGIN) if chip_kind else (RIGHT - MARGIN)
    t.left, t.top, t.width, t.height = I(MARGIN), I(0.45), I(title_w), I(0.8)
    _apply_text(t.text_frame, title, title_size, True, C['text'], 'l', 'm', (0, 0, 0, 0))
    b.title = t
    if chip_kind:
        b.chip = chip(slide.shapes, chip_kind, RIGHT - 1.7, 0.5, name='Status chip')
    # caption bar (moved to the end of the reading order in finish())
    nl = caption_lines(caption)
    bh = 0.6 if nl == 1 else 0.88
    by = 7.0 - bh
    b.bar_y = by
    b.bottom = by - 0.12
    b._tail = []
    panel = box(slide.shapes, MARGIN, by, 12.13, bh, fill=C['panel'], radius=0.08, text=caption,
                size=20, color=C['text'], align='l', anchor='m', margins=(0.3, 0.03, 0.2, 0.03),
                name='Caption')
    rule = box(slide.shapes, MARGIN, by, 0.08, bh, fill=C['orange'], name='Caption rule', deco=True)
    sec = label(slide.shapes, MARGIN, 0.13, 8, 0.3, section.upper(), size=14, color=C['muted'],
                name='Section label')
    num = label(slide.shapes, RIGHT - 1.0, 7.06, 1.0, 0.3, str(n), size=14, color=C['muted'],
                align='r', name='Slide number')
    _slide_number_field(num, n)
    b._tail = [panel, rule, sec, num]
    return b


def _slide_number_field(shape, n):
    p = shape.text_frame.paragraphs[0]._p
    r = p.find(qn('a:r'))
    rpr = r.find(qn('a:rPr'))
    fld = etree.SubElement(p, qn('a:fld'), id='{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}', type='slidenum')
    fld.append(rpr)
    t = etree.SubElement(fld, qn('a:t'))
    t.text = str(n)
    p.remove(r)


def finish(b):
    """Put caption, section label and slide number last in the reading order."""
    tree = b.slide.shapes._spTree
    for s in b._tail:
        el = s._element
        tree.remove(el)
        tree.append(el)


def finalize(prs):
    """Language tags on every run; no notes, no transitions are ever created."""
    for slide in prs.slides:
        root = slide._element
        for r in root.iter(qn('a:r'), qn('a:fld')):
            rpr = r.find(qn('a:rPr'))
            if rpr is None:
                rpr = etree.Element(qn('a:rPr'))
                r.insert(0, rpr)
            rpr.set('lang', 'en-GB')
        for e in root.iter(qn('a:endParaRPr')):
            e.set('lang', 'en-GB')
