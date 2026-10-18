"""Chunk 3: slides 21-31 (Large, wagons, range, tenders, stations, servicing, signalling)."""
from pptx.enum.shapes import MSO_SHAPE

from lib import (C, MARGIN, RIGHT, add_base, arc_pts, arrow, badge, blend, block_band, box,
                 box_building, chip, chip_w, coal_heap, credit, drop, finish, group, icon,
                 inline_chip, label, lever_row, line, oval, person, photo, picture, pill, poly,
                 polyline, rails, set_alt, signal, state_pill, table, tank, text_width_in, train,
                 train_width)
from slides_b import role_slide


# ------------------------------------------------------------------ 21
def s21(prs, n):
    role_slide(prs, n, 'Large: the heavy hauler',
               'Large pulls the heaviest trains. It is thirsty and slow to speed up. It suits heavy freight.',
               'P12', 'Reference: BR Standard 9F',
               [('weights', 'Most powerful', ()), ('drop_coal', 'High coal and water demand', ()),
                ('aux_tender', 'Auxiliary tender for very long routes', ()),
                ('slow', 'Slower to speed up', ('Proposed',))],
               photo_h=3.6)


# ------------------------------------------------------------------ 22
def _wagon(s, x, y, w, h, col, rows, cols):
    """Side-view wagon with a grid of item stacks. y = top of body."""
    box(s, x, y, w, h, fill=blend(col, C['bg'], 0.22), line=col, lw=2.5, radius=0.06)
    sq, gap = 0.26, 0.07
    gw = cols * sq + (cols - 1) * gap
    gh = rows * sq + (rows - 1) * gap
    x0, y0 = x + (w - gw) / 2, y + (h - gh) / 2
    for r in range(rows):
        for c in range(cols):
            box(s, x0 + c * (sq + gap), y0 + r * (sq + gap), sq, sq, fill=col, radius=0.03)
    box(s, x - 0.1, y + h, w + 0.2, 0.08, fill=C['steel'])
    for f in (0.18, 0.82):
        oval(s, x + w * f, y + h + 0.22, 0.32, fill=C['coal'], line=C['steel'], lw=1.5)
    box(s, x - 0.3, y + h + 0.38, w + 0.6, 0.06, fill=C['steel'])


def s22(prs, n):
    b = add_base(prs, n, 'Locomotives', 'Wagons', 'Provisional',
                 'Steam wagons are cheaper and hold less. About 20 item stacks, against 32 in a vanilla wagon.')
    sh = b.shapes
    g = group(sh, 'Wagon comparison', 'Two wagons side by side. The steam wagon holds a grid of 20 '
              'squares (about 20 item stacks); the vanilla wagon holds 32.')
    _wagon(g.shapes, 1.0, 1.9, 4.6, 1.6, C['orange'], 4, 5)
    _wagon(g.shapes, 7.13, 1.9, 4.6, 1.6, C['electric'], 4, 8)
    t1 = 'Steam wagon: about 20 stacks'
    label(sh, 1.0, 4.15, 4.6, 0.4, t1, size=20, bold=True)
    inline_chip(sh, 'Provisional', 1.0, 4.62)
    label(sh, 7.13, 4.15, 4.6, 0.4, 'Vanilla wagon: 32 stacks', size=20, bold=True)
    lg = group(sh, 'Legend', 'Legend: one square is one stack, a pile of one item.')
    box(lg.shapes, 1.0, 5.33, 0.26, 0.26, fill=C['orange'], radius=0.03)
    label(lg.shapes, 1.4, 5.25, 4.5, 0.42, 'Stack = a pile of one item.', size=18, anchor='m')
    label(sh, 7.13, 5.25, 2.6, 0.42, 'Wagon catalogue:', size=18, anchor='m')
    inline_chip(sh, 'Undecided', 7.13 + text_width_in('Wagon catalogue:', 18) + 0.15, 5.28)
    finish(b)


# ------------------------------------------------------------------ 23
def s23(prs, n):
    b = add_base(prs, n, 'Locomotives', 'How far can an engine go?', 'Agreed',
                 'Range comes from real tank sizes and real use. Hills and heavy loads use more. Coasting uses less.')
    sh = b.shapes
    g = group(sh, 'Engine gauges', 'Gauges for one engine: coal about two-thirds full, water about '
              'half full, and a readout card showing range left with a dial and no number.')
    s = g.shapes
    box(s, MARGIN, 1.5, 5.6, 3.3, fill=C['panel'], radius=0.1)
    tank(s, 1.0, 2.0, 1.0, 2.3, 0.65, C['coal'], outline=C['coal_line'])
    coal_heap(s, 1.5, 1.85, 0.6, 0.22)
    tank(s, 2.3, 2.0, 1.0, 2.3, 0.5, C['water'])
    drop(s, 2.8, 1.72, 0.3)
    # readout card
    box(s, 3.65, 2.0, 2.3, 2.3, fill=C['bg'], line=C['orange'], lw=2, radius=0.12)
    label(s, 3.65, 2.12, 2.3, 0.4, 'Range left', size=18, bold=True, align='c')
    from lib import curve_arrow
    polyline(s, arc_pts(4.8, 3.75, 0.75, 0.75, 180, 0, 24), color=C['line'], lw=8)
    polyline(s, arc_pts(4.8, 3.75, 0.75, 0.75, 180, 70, 16), color=C['orange'], lw=8)
    line(s, 4.8, 3.75, 4.8 + 0.55 * 0.342, 3.75 - 0.55 * 0.94, C['text'], 3)
    oval(s, 4.8, 3.75, 0.16, fill=C['text'])
    inline_chip(sh, 'Agreed', 3.65, 4.95, text='Agreed direction')
    # table of route-dependent targets
    label(sh, 6.75, 1.45, 3.6, 0.42, 'Route-dependent targets', size=20, bold=True, anchor='m')
    inline_chip(sh, 'Provisional', 6.75 + text_width_in('Route-dependent targets', 20, True) + 0.2, 1.48)
    table(sh, 6.75, 2.05, [2.0, 2.0, 2.0], 0.62,
          [['', 'Water (km)', 'Coal (km)'], ['Small', '2–3', '4–5'], ['Medium', '4–5', 'about 7'],
           ['Large', '6–7', 'about 10']], size=20, aligns=['l', 'c', 'c'],
          alt='Table of provisional, route-dependent range targets. Small: water 2 to 3 km, coal 4 to 5 km. '
              'Medium: water 4 to 5 km, coal about 7 km. Large: water 6 to 7 km, coal about 10 km.')
    note = 'Compacted Coal burns more efficiently.'
    label(sh, 6.75, 4.85, 5.2, 0.42, note, size=18, anchor='m')
    inline_chip(sh, 'Agreed', 6.75 + text_width_in(note, 18) + 0.2, 4.88)
    finish(b)


# ------------------------------------------------------------------ 24
def s24(prs, n):
    b = add_base(prs, n, 'Locomotives', 'Tenders', 'Agreed',
                 'A normal tender belongs to the engine. An auxiliary tender is an extra wagon: more range, less pulling power.')
    sh = b.shapes
    k = 1.5
    ry = 2.95
    g = group(sh, 'Train with tenders', 'A train facing left: the engine and its normal tender form one '
              'vehicle; behind them a hatched auxiliary tender, a separate wagon; then three cargo wagons.')
    s = g.shapes
    x = 0.9
    # engine (facing left): chimney at the left end
    box(s, x, ry - 0.65 * k, 1.0 * k, 0.55 * k, fill=C['orange'], line=C['ink'], lw=0.75, radius=0.06 * k)
    box(s, x + 0.07 * k, ry - 0.93 * k, 0.15 * k, 0.3 * k, fill=C['coal'])
    box(s, x + 0.52 * k, ry - 0.58 * k, 0.14 * k, 0.14 * k, fill=C['panel'])
    for wx in (0.2, 0.5, 0.8):
        oval(s, x + wx * k, ry - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
    x += 1.0 * k
    box(s, x, ry - 0.25 * k, 0.08 * k, 0.05 * k, fill=C['steel'])
    x += 0.08 * k
    # normal tender
    box(s, x, ry - 0.6 * k, 0.62 * k, 0.48 * k, fill=C['orange'], line=C['ink'], lw=0.75)
    box(s, x + 0.05 * k, ry - 0.68 * k, 0.52 * k, 0.12 * k, fill=C['coal'])
    for wx in (0.15, 0.47):
        oval(s, x + wx * k, ry - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
    tender_end = x + 0.62 * k
    x = tender_end
    box(s, x, ry - 0.25 * k, 0.3 * k, 0.05 * k, fill=C['steel'])
    x += 0.3 * k
    aux_x = x
    box(s, x, ry - 0.57 * k, 0.8 * k, 0.45 * k, fill=('hatch', C['orange'], C['panel']), line=C['orange'], lw=2)
    for wx in (0.2, 0.6):
        oval(s, x + wx * k, ry - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
    x += 0.8 * k
    for _ in range(3):
        box(s, x, ry - 0.25 * k, 0.08 * k, 0.05 * k, fill=C['steel'])
        x += 0.08 * k
        box(s, x, ry - 0.57 * k, 0.8 * k, 0.45 * k, fill=C['wagon'], line=C['ink'], lw=0.75)
        for wx in (0.2, 0.6):
            oval(s, x + wx * k, ry - 0.1 * k, 0.2 * k, fill=C['coal'], line=C['steel'], lw=1)
        x += 0.8 * k
    box(s, 0.7, ry, x - 0.5, 0.06, fill=C['steel'])
    # bracket under engine + tender
    bx0, bx1 = 0.9, tender_end
    bk = group(sh, 'Brackets', 'Brackets: one under the engine and normal tender, one under the auxiliary tender.')
    polyline(bk.shapes, [(bx0, ry + 0.15), (bx0, ry + 0.27), (bx1, ry + 0.27), (bx1, ry + 0.15)],
             color=C['orange'], lw=2)
    label(sh, bx0, ry + 0.35, 2.4, 0.4, 'One vehicle', size=18, bold=True, anchor='m')
    inline_chip(sh, 'Proposed', bx0, ry + 0.85)
    # auxiliary tender labels
    ax = aux_x - 0.1
    polyline(bk.shapes, [(aux_x, ry + 0.15), (aux_x, ry + 0.27), (aux_x + 0.8 * k, ry + 0.27),
                         (aux_x + 0.8 * k, ry + 0.15)], color=C['orange'], lw=2)
    label(sh, ax, ry + 0.35, 5.0, 0.4, 'Auxiliary tender: a separate wagon', size=18, bold=True, anchor='m')
    inline_chip(sh, 'Proposed', ax, ry + 0.85, text='Water only')
    inline_chip(sh, 'Proposed', ax + chip_w('Water only') + 0.12, ry + 0.85, text='Coal and water')
    from lib import list_row
    dg = group(sh, 'Downsides', 'Downsides of an auxiliary tender: it carries no freight and adds weight.')
    list_row(dg.shapes, ax, ry + 1.4, 4.2, 0.66, 'no_tender', 'Carries no freight', 18)
    list_row(dg.shapes, ax, ry + 2.16, 4.2, 0.66, 'weight', 'Adds weight', 18)
    label(sh, 9.33, 3.0, 3.4, 0.34, 'A real tender, for illustration', size=14, color=C['text'])
    photo(sh, 'P18', 9.33, 3.36, 3.4, 2.27)
    finish(b)


# ------------------------------------------------------------------ 25
def s25(prs, n):
    b = add_base(prs, n, 'Stations', 'A steam station', 'Agreed',
                 'A steam station is big. It has a water tower, coal loading, a signal box and a platform.')
    sh = b.shapes
    ty = 4.0                                    # main line (top-down)
    g = group(sh, 'Station plan', 'Top-down plan of a steam station: main line with a docked engine, '
              'a coaling stage fed by a conveyor, a water tower fed by a pipe, a signal box and a '
              'platform for cargo. A dashed outline shows a large footprint; a small vanilla station '
              'is shown for comparison.')
    s = g.shapes
    box(s, 0.7, 1.95, 9.45, 3.8, fill=None, line=C['muted'], lw=1.5, radius=0.15, dash='dash')
    rails(s, 0.5, ty, 12.25, gap=0.24)
    # coaling stage + conveyor
    box(s, 4.6, 2.4, 1.8, 0.85, fill=C['panel_alt'], line=C['steel'], lw=1.5)
    coal_heap(s, 5.5, 3.08, 1.2, 0.45)
    line(s, 5.75, 3.25, 5.75, ty - 0.02, C['steel'], 4)
    box(s, 0.85, 2.73, 3.75, 0.2, fill=C['coal'], line=C['steel'], lw=1, dash='dash')
    kx = 1.05
    while kx < 4.4:
        box(s, kx, 2.76, 0.12, 0.14, fill=C['orange'], shape=MSO_SHAPE.CHEVRON)
        kx += 0.38
    # water tower + pipe
    oval(s, 8.1, 2.8, 1.0, fill=C['water'], line='FFFFFF', lw=2)
    drop(s, 8.1, 2.83, 0.4, col='FFFFFF')
    line(s, 7.75, 3.15, 7.05, ty - 0.02, C['steel'], 4)
    polyline(s, [(8.1, 1.55), (8.1, 2.3)], color=C['water'], lw=9)
    # docked engine with bunker and hatch
    box(s, 5.2, ty - 0.12, 2.2, 0.48, fill=C['orange'], line=C['ink'], lw=1, radius=0.1)
    oval(s, 5.75, ty + 0.12, 0.2, fill=C['coal'], line=C['text'], lw=1.5)
    oval(s, 7.05, ty + 0.12, 0.2, fill=C['water'], line=C['text'], lw=1.5)
    # platform and signal box
    box(s, 4.5, 4.55, 3.6, 0.55, fill=C['panel_alt'], line=C['steel'], lw=1.5,
        text='Platform and cargo loading', size=14)
    box(s, 2.2, 4.55, 1.5, 1.0, fill=C['panel_alt'], line=C['orange'], lw=1.5)
    line(s, 2.2, 4.55, 3.7, 5.55, C['orange'], 1.5)
    label(s, 0.85, 3.32, 2.3, 0.3, 'Conveyor: coal in', size=14)
    label(sh, 0.8, 4.9, 1.3, 0.3, 'Signal box', size=14, align='r')
    label(sh, 8.25, 1.45, 1.9, 0.32, 'Pipe: water in', size=14, anchor='m')
    label(sh, 4.0, 3.4, 1.65, 0.3, 'Coal bunker', size=14, align='r')
    label(sh, 7.72, 3.42, 1.4, 0.3, 'Water hatch', size=14)
    label(sh, 4.5, 5.3, 5.5, 0.36, 'Large footprint. Exact size undecided.', size=16, align='l')
    box(sh, 11.1, 4.75, 1.6, 0.75, fill=C['panel'], line=C['electric'], lw=1.5,
        text='Vanilla station (for comparison)', size=14, name='Vanilla station')
    photo(sh, 'P17', 10.33, 1.5, 2.4, 1.4)
    finish(b)


# ------------------------------------------------------------------ 26
def _step_icon(s, k, cx, cy):
    o = C['orange']
    if k == 1:
        arrow(s, cx - 0.35, cy, cx + 0.35, cy, o, 3, hs='lg')
    elif k == 2:
        polyline(s, [(cx - 0.35, cy - 0.2), (cx - 0.35, cy + 0.2), (cx + 0.35, cy + 0.2), (cx + 0.35, cy - 0.2)],
                 color=o, lw=3)
        box(s, cx - 0.22, cy - 0.05, 0.44, 0.2, fill=o, radius=0.05)
    elif k == 3:
        for dx, dy, d in ((-0.22, 0.12, 0.26), (0.0, -0.02, 0.32), (0.24, -0.16, 0.38)):
            oval(s, cx + dx, cy + dy, d, fill=blend('F2F2F2', C['panel'], 0.85), dh=d * 0.65)
    elif k == 4:
        poly(s, [(cx - 0.35, cy - 0.25), (cx - 0.15, cy - 0.25), (cx + 0.3, cy + 0.15), (cx + 0.1, cy + 0.15)],
             fill=C['steel'])
        coal_heap(s, cx - 0.2, cy - 0.25, 0.36, 0.14)
    elif k == 5:
        polyline(s, arc_pts(cx - 0.05, cy + 0.2, 0.4, 0.4, 170, 40, 14), color=o, lw=3, head='end', hs='lg')
        drop(s, cx + 0.3, cy + 0.12, 0.22)
    elif k == 6:
        for dx in (-0.2, 0.2):
            oval(s, cx + dx, cy, 0.3, fill=None, line=o, lw=2)
            line(s, cx + dx - 0.2, cy, cx + dx + 0.2, cy, o, 1.5)
            line(s, cx + dx, cy - 0.2, cx + dx, cy + 0.2, o, 1.5)
    elif k == 7:
        oval(s, cx, cy, 0.5, fill=None, line=o, lw=2.5)
        line(s, cx, cy, cx, cy - 0.17, C['text'], 2.5)
        line(s, cx, cy, cx + 0.13, cy + 0.06, C['text'], 2.5)
    elif k == 8:
        arrow(s, cx + 0.3, cy - 0.15, cx - 0.3, cy - 0.15, o, 3, hs='lg')
        box(s, cx - 0.35, cy + 0.02, 0.7, 0.14, fill=C['steel'])
    elif k == 9:
        box(s, cx - 0.35, cy - 0.15, 0.55, 0.3, fill=o, radius=0.05)
        box(s, cx + 0.05, cy - 0.3, 0.1, 0.16, fill=C['coal'])
        arrow(s, cx + 0.22, cy, cx + 0.42, cy, C['text'], 2.5, hs='lg')


def s26(prs, n):
    b = add_base(prs, n, 'Stations', 'Servicing takes time', 'Agreed',
                 'The machinery finds each engine\'s hatches. Servicing takes visible time. Then it folds away.')
    sh = b.shapes
    steps = ['Approach slowly', 'Settle into dock', 'Long steam hiss', 'Coal chutes in', 'Water arm swings',
             'Finds the hatches', 'Wait while it fills', 'Folds away', 'Depart']
    g = group(sh, 'Servicing steps', 'Nine servicing steps: approach slowly, settle into the dock, a '
              'long hiss of steam, coal chutes in, the water arm swings round, the machinery finds '
              'the hatches, wait while it fills, the machinery folds away (proposed), depart.')
    s = g.shapes
    W, H, gap = 2.2, 1.45, 0.28
    cols = [MARGIN + i * (W + gap) for i in range(5)]
    pos = [(cols[i], 1.45) for i in range(5)] + [(cols[4 - i], 3.25) for i in range(4)]
    for k, ((x, y), t) in enumerate(zip(pos, steps), 1):
        box(s, x, y, W, H, fill=C['panel'], radius=0.08)
        badge(s, x + 0.3, y + 0.3, k, d=0.4, size=16)
        _step_icon(s, k, x + W / 2 + 0.25, y + 0.42)
        label(s, x + 0.12, y + 0.78, W - 0.24, 0.6, t, size=16, align='c', anchor='m')
        if k == 8:
            inline_chip(s, 'Proposed', x + W - chip_w('Proposed') - 0.06, y - 0.18)
    for k in range(4):
        arrow(s, cols[k] + W + 0.03, 2.17, cols[k + 1] - 0.03, 2.17, hs='lg')
    arrow(s, cols[4] + W / 2, 1.45 + H + 0.03, cols[4] + W / 2, 3.25 - 0.03, hs='lg')
    for k in range(3):
        arrow(s, cols[4 - k] - 0.03, 3.97, cols[3 - k] + W + 0.03, 3.97, hs='lg')
    # time bars (qualitative)
    tg = group(sh, 'Stop time', 'Qualitative stop-time bars: a long stop for empty tanks, a short stop '
               'for a small top-up. Rates and times are proposed or provisional.')
    t = tg.shapes
    from lib import icon as _icon
    oval(t, 0.95, 5.2, 0.42, fill=None, line=C['orange'], lw=2.5)
    line(t, 0.95, 5.2, 0.95, 5.06, C['text'], 2)
    line(t, 0.95, 5.2, 1.06, 5.25, C['text'], 2)
    label(t, 1.35, 4.9, 1.8, 0.3, 'Empty tanks', size=14, anchor='m')
    box(t, 3.1, 4.95, 6.0, 0.2, fill=C['panel_alt'], radius=0.04)
    box(t, 3.1, 4.95, 5.4, 0.2, fill=C['orange'], radius=0.04)
    label(t, 1.35, 5.25, 1.8, 0.3, 'Small top-up', size=14, anchor='m')
    box(t, 3.1, 5.3, 6.0, 0.2, fill=C['panel_alt'], radius=0.04)
    box(t, 3.1, 5.3, 1.3, 0.2, fill=C['orange'], radius=0.04)
    lbl = 'Rates and times:'
    label(sh, 9.35, 4.88, 2.2, 0.32, lbl, size=14, anchor='m')
    inline_chip(sh, 'Proposed', 9.35, 5.22)
    inline_chip(sh, 'Provisional', 9.35 + chip_w('Proposed') + 0.1, 5.22)
    finish(b)


# ------------------------------------------------------------------ 27
def _tile_icon(s, kind, cx, cy):
    if kind in ('water', 'both', 'full'):
        drop(s, cx - (0.35 if kind != 'water' else 0), cy, 0.55)
    if kind in ('coal', 'both', 'full'):
        coal_heap(s, cx + (0.35 if kind != 'coal' else 0), cy + 0.25, 0.75, 0.42)
    if kind == 'full':
        box(s, cx - 0.9, cy + 0.42, 1.8, 0.14, fill=C['steel'])
        box_building(s, cx + 0.62, cy - 0.15, 0.42, 0.42, with_person=False, windows=False)


def s27(prs, n):
    b = add_base(prs, n, 'Stations', 'Service points', 'Proposed',
                 'Separate coal and water stops make long routes possible. Not every stop needs a full station.',
                 caption_chip='Agreed')
    sh = b.shapes
    tiles = [('Water Tower', 'water'), ('Coaling Stage', 'coal'), ('Steam Service Point', 'both'),
             ('Full Station', 'full')]
    g = group(sh, 'Service point types', 'Four proposed service point types: water tower, coaling stage, '
              'steam service point with water and coal, and full station with platform, water, coal '
              'and signal box.')
    for k, (name, kind) in enumerate(tiles):
        x = MARGIN + k * (2.8 + 0.31)
        box(g.shapes, x, 1.5, 2.8, 2.4, fill=C['panel'], radius=0.1)
        _tile_icon(g.shapes, kind, x + 1.4, 2.55)
        label(g.shapes, x + 0.1, 3.2, 2.6, 0.55, name, size=18, bold=True, align='c', anchor='m')
        inline_chip(g.shapes, 'Proposed', x + 2.8 - chip_w('Proposed') - 0.1, 1.6)
    rg = group(sh, 'Route', 'A route from a mine to a factory with a water stop and a water and coal '
               'stop on the way; a train is at the water stop.')
    s = rg.shapes
    y = 4.75
    line(s, 1.0, y, 12.3, y, C['steel'], 5)
    stops = [(1.3, 'Mine'), (4.75, 'Water stop'), (8.2, 'Water and coal stop'), (11.95, 'Factory')]
    for x, t in stops:
        oval(s, x, y, 0.36, fill=C['bg'], line=C['orange'], lw=3)
        label(s, x - 1.3, y + 0.3, 2.6, 0.36, t, size=16, align='c')
    train(s, 3.75, y - 0.05, 1, 0.7, name='Train at water stop')
    finish(b)


# ------------------------------------------------------------------ 28
def s28(prs, n):
    b = add_base(prs, n, 'Signals', 'Blocks: one train at a time', 'Agreed',
                 'Signals split the track into blocks. Only one train may be in a block. The next train waits.')
    sh = b.shapes
    ty = 3.65
    posts = (3.7, 6.7, 9.7)
    g = group(sh, 'Block diagram', 'Track split into four blocks, A to D, by three signals. A train '
              'occupies block B, so the signal at its entrance shows STOP and a second train waits in '
              'block A. Blocks C and D are free and their signals show CLEAR.')
    s = g.shapes
    rails(s, 0.8, ty, 11.7)
    edges = [0.8, *posts, 12.5]
    states = ['occupied', 'occupied', 'free', 'free']
    for i, (st, letter) in enumerate(zip(states, 'ABCD')):
        block_band(s, edges[i] + 0.04, 4.35, edges[i + 1] - edges[i] - 0.08, st,
                   'Block %s: %s' % (letter, st), h=0.38, size=16)
    train(s, 4.0, ty, 2, 2.0 / train_width(2), name='Train in block B')
    train(s, 0.9, ty, 2, 0.62, name='Waiting train in block A')
    for x, st in zip(posts, ('stop', 'clear', 'clear')):
        signal(s, x, 4.2, st, label_pos='topright')
    # "One block" dimension over block C
    line(s, 6.7, 2.25, 9.7, 2.25, C['text'], 1.5, head='both')
    line(s, 6.7, 2.1, 6.7, 2.4, C['text'], 1.5)
    line(s, 9.7, 2.1, 9.7, 2.4, C['text'], 1.5)
    label(s, 6.7, 1.83, 3.0, 0.38, 'One block', size=18, bold=True, align='c', anchor='b')
    # "Waits" callout
    label(s, 0.9, 2.05, 1.4, 0.4, 'Waits', size=18, bold=True, anchor='m')
    arrow(s, 1.7, 2.45, 2.05, 2.98, C['orange'], 2.25, hs='lg')
    # legend
    lg = group(sh, 'Legend', 'Legend: arm flat means stop; arm raised means go.')
    box(lg.shapes, MARGIN, 4.95, 5.0, 0.95, fill=C['panel'], radius=0.08)
    signal(lg.shapes, 1.05, 5.82, 'stop', scale=0.6, show_label=False)
    label(lg.shapes, 1.65, 5.05, 1.6, 0.75, 'Arm flat: stop', size=16, anchor='m')
    signal(lg.shapes, 3.45, 5.82, 'clear', scale=0.6, show_label=False)
    label(lg.shapes, 4.0, 5.05, 1.6, 0.75, 'Arm raised: go', size=16, anchor='m')
    finish(b)


# ------------------------------------------------------------------ 29
def s29(prs, n):
    b = add_base(prs, n, 'Signals', 'Semaphore signals', 'Agreed',
                 'Arm flat means stop. Arm raised means go. On real railways, a broken wire drops the arm to stop.')
    sh = b.shapes
    inline_chip(sh, 'Undecided', RIGHT - chip_w('Arm style: undecided'), 0.98, text='Arm style: undecided')
    extra = 'Blender model, not yet in the game.'
    for slot, x, txt, tgt in (('R03', MARGIN, 'STOP: arm flat', (0.475, 0.215)),
                              ('R04', 4.95, 'CLEAR: arm raised', (0.53, 0.16))):
        photo(sh, slot, x, 1.45, 4.1, 3.075, extra=extra)
        cg = group(sh, 'Label ' + txt, 'Label: ' + txt)
        box(cg.shapes, x + 0.15, 1.6, 2.75, 0.45, fill=C['bg'], line=C['orange'], lw=1.5, radius=0.08,
            text=txt, size=18, bold=True)
        tx, ty_ = x + tgt[0] * 4.1, 1.45 + tgt[1] * 3.075
        arrow(cg.shapes, x + 2.9, 1.82, tx + 0.15, ty_ - 0.02, C['orange'], 2, hs='lg')
    # distant signal mini-diagram
    label(sh, 9.35, 1.42, 3.4, 0.42, 'Distant signal', size=20, bold=True, color=C['orange'])
    g = group(sh, 'Distant signal', 'Distant signal with a yellow fishtail arm. Arm flat: be ready to '
              'stop at the next signal. Arm raised: next signal is clear.')
    s = g.shapes
    signal(s, 10.0, 3.0, 'stop', kind='distant', side=-1, show_label=False, scale=0.85)
    label(s, 10.35, 2.05, 2.4, 0.9, 'Be ready to stop at the next signal', size=16, anchor='m')
    signal(s, 10.0, 4.3, 'clear', kind='distant', side=-1, show_label=False, scale=0.85)
    label(s, 10.35, 3.45, 2.4, 0.8, 'Next signal is clear', size=16, anchor='m')
    photo(sh, 'P13', 10.33, 4.45, 2.0, 1.2)
    finish(b)


# ------------------------------------------------------------------ 30
def s30(prs, n):
    b = add_base(prs, n, 'Signals', 'The signal box', 'Agreed',
                 'A signaller pulls levers. Rods and wires carry the movement out to the signals.')
    sh = b.shapes
    photo(sh, 'P14', MARGIN, 1.45, 3.3, 4.4, focus=(0.5, 0.45))
    g = group(sh, 'Signal box cutaway', 'Cutaway of a two-floor signal box. Upstairs, on the operating '
              'floor, a signaller stands at a row of eight coloured levers under a track diagram. '
              'Downstairs, in the locking room, rods leave the wall and run out to the signals and points.')
    s = g.shapes
    X, Y, W, Hu, Hl = 5.25, 1.5, 5.2, 2.15, 1.3
    box(s, X, Y, W, Hu + Hl, fill=C['panel'], line=C['orange'], lw=1.5)
    line(s, X, Y + Hu, X + W, Y + Hu, C['orange'], 1.5)
    label(s, X + 0.12, Y + 0.08, 2.0, 0.32, 'Operating floor', size=14, bold=True, color=C['orange'])
    box(s, X + 1.75, Y + 0.12, W - 1.9, 0.3, fill=blend(C['water'], C['panel'], 0.35), line=C['steel'], lw=0.75)
    box(s, X + 1.0, Y + 0.55, 2.6, 0.24, fill='F2F2F2')
    polyline(s, [(X + 1.1, Y + 0.67), (X + 2.2, Y + 0.67), (X + 2.5, Y + 0.6), (X + 3.5, Y + 0.6)],
             color=C['ink'], lw=1.5)
    person(s, X + 0.42, Y + 0.95, 1.1)
    cols = [C['red']] * 3 + [C['amber']] * 2 + ['points'] * 2 + [C['electric']]
    lx0, lw_, lg_ = X + 1.0, 0.12, 0.32
    lever_row(s, lx0, Y + 0.95, cols, w=lw_, h=0.85, gap=lg_)
    groups = [('Stop', 0, 3), ('Distant', 3, 5), ('Points', 5, 7), ('Lock', 7, 8)]
    for t, a, z in groups:
        x0 = lx0 + a * (lw_ + lg_) - 0.1
        x1 = lx0 + (z - 1) * (lw_ + lg_) + lw_ + 0.1
        label(s, x0 - 0.15, Y + 1.88, x1 - x0 + 0.3, 0.27, t, size=14, align='c')
    label(s, X + 3.75, Y + 0.38, 1.4, 0.48, 'Example colours only', size=14, color=C['muted'])
    label(s, X + 0.12, Y + Hu + 0.06, 2.0, 0.32, 'Locking room', size=14, bold=True, color=C['orange'])
    for k in range(4):
        box(s, X + 1.6, Y + Hu + 0.45 + k * 0.2, 3.2, 0.06, fill=C['steel'])
    for k in range(8):
        box(s, X + 1.75 + k * 0.38, Y + Hu + 0.38, 0.05, 0.8, fill=C['line'])
    # rods out of the left wall, down to the track
    ry = 5.55
    for k in range(8):
        yy = Y + Hu + 0.42 + k * 0.1
        xx = X - 0.25 - k * 0.07
        polyline(s, [(X + 1.6, yy), (xx, yy), (xx, ry)], color=C['steel'], lw=1.25)
    box(s, 4.2, ry, 6.25, 0.06, fill=C['steel'])
    label(sh, X, Y + Hu + Hl + 0.08, W, 0.36, 'Rods and wires run out to signals and points', size=14,
          align='c')
    box(sh, 10.7, 1.5, 2.03, 1.5, fill=C['bg'], line=C['orange'], lw=1.5, radius=0.08,
        text='Levers are locked so two clashing signals cannot both be clear', size=16, align='l',
        margins=(0.12, 0.06, 0.08, 0.06), name='Interlocking note')
    photo(sh, 'P15', 10.53, 4.05, 2.2, 1.5)
    finish(b)


# ------------------------------------------------------------------ 31
def s31(prs, n):
    b = add_base(prs, n, 'Signals', 'How it works together', 'Agreed',
                 'The signaller watches the blocks, sets the signals and keeps trains apart.')
    sh = b.shapes
    ty = 4.25
    sig_x = (4.6, 7.6, 10.4)
    g = group(sh, 'Whole signalling system', 'A signal box controls three signals through rods that run '
              'beside the track, each ending at a counterweight that returns the arm to STOP. A train is in block 2, so signal 1 is held at STOP behind it while '
              'signals 2 and 3 show CLEAR. A dotted line marks the box territory. A second box further '
              'along is linked by bell codes.')
    s = g.shapes
    # territory
    box(s, 0.62, 1.72, 11.0, 3.7, fill=None, line=C['orange'], lw=1.5, radius=0.15, dash='sqdot')
    label(s, 8.3, 1.8, 3.2, 0.36, 'Box territory: limited', size=16, bold=True, color=C['orange'],
          align='r')
    # bell-code link between boxes
    polyline(s, [(2.1, 1.82), (2.1, 1.58), (12.22, 1.58), (12.22, 2.08)], color=C['steel'], lw=2,
             head='both', dash='dash', hs='lg')
    label(s, 6.0, 1.25, 3.6, 0.3, 'Bell codes', size=16, anchor='m', align='c')
    # box 1
    box_building(s, 0.85, 2.25, 2.5, 1.45, with_person=False)
    person(s, 1.25, 2.75, 0.85)
    lever_row(s, 2.0, 2.78, [C['red'], C['red'], C['amber']], w=0.12, h=0.7, gap=0.22)
    # box 2 (smaller, steel)
    box_building(s, 11.82, 2.3, 0.8, 0.7, outline=C['steel'], with_person=False)
    # track, blocks, train
    rails(s, 0.8, ty, 11.7)
    for x0, x1, t in ((0.8, 4.6, 'Block 1'), (4.6, 7.6, 'Block 2'), (7.6, 10.4, 'Block 3')):
        box(s, x0 + 0.04, 4.98, x1 - x0 - 0.08, 0.34, fill=C['panel_alt'], line=C['steel'], lw=1,
            text=t, size=14, bold=True)
    train(s, 4.8, ty, 2, 2.15 / train_width(2), name='Train in block 2')
    # rods along the ground to each post, with counterweights
    for i, (x, st, lp) in enumerate(zip(sig_x, ('stop', 'clear', 'clear'), ('topright', 'topright', 'topright'))):
        ry = 4.66 + i * 0.085
        polyline(s, [(3.35, 2.95 + i * 0.25), (3.6 + i * 0.12, 2.95 + i * 0.25), (3.6 + i * 0.12, ry),
                     (x - 0.12, ry)], color=C['orange'], lw=2.25, dash='dash')
        box(s, x - 0.26, ry - 0.07, 0.14, 0.14, fill=C['steel'], line=C['ink'], lw=0.75)
        signal(s, x, 4.72, st, label_pos=lp)
    # steps strip
    steps = ['Train in block', 'Signal at STOP', 'Train leaves', 'Lever pulled: arm lifts']
    sg = group(sh, 'Steps', 'Four steps: 1, the train is in the block. 2, the signaller keeps the signal '
               'behind it at stop. 3, the train leaves and the block is free. 4, the signaller pulls the '
               'lever and the arm lifts.')
    for k, t in enumerate(steps):
        x = MARGIN + k * 3.05
        badge(sg.shapes, x + 0.22, 5.65, k + 1, d=0.4, size=16)
        label(sg.shapes, x + 0.5, 5.45, 2.5, 0.4, t, size=18, anchor='m')
    box(sh, MARGIN, 5.92, 12.13, 0.34, fill=C['amber'], radius=0.06,
        text='In the mod: no player wiring.', size=16, bold=True, color=C['ink'], align='l',
        margins=(0.2, 0, 0.1, 0), name='In the mod')
    inline_chip(sh, 'Agreed', MARGIN + 0.2 + text_width_in('In the mod: no player wiring.', 16, True) + 0.3,
                5.92, h=0.34)
    finish(b)


SLIDES = [s21, s22, s23, s24, s25, s26, s27, s28, s29, s30, s31]
