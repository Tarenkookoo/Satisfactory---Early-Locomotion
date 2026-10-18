"""Chunk 1: slides 1-10 (welcome, primers, water)."""
from lib import (C, I, MARGIN, RIGHT, add_base, arc_pts, arrow, badge, blend, box,
                 chip, credit, curve_arrow, finish, group, inline_chip, label, line,
                 mark_decorative, oval, photo, picture, poly, polyline, tank, train,
                 credit_text, _apply_text, chip_w)
from pptx.enum.shapes import MSO_SHAPE


# ------------------------------------------------------------------ 1
def s01(prs, n):
    """Title slide: own layout (no caption bar, no chip)."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    sh = slide.shapes
    t = sh.title
    t.left, t.top, t.width, t.height = I(MARGIN), I(1.35), I(6.3), I(2.1)
    _apply_text(t.text_frame, 'Early Locomotion', 54, True, C['text'], 'l', 'b', (0, 0, 0, 0),
                spacing=0.9)
    box(sh, MARGIN, 3.65, 2.6, 0.1, fill=C['orange'], name='Title rule', deco=True)
    label(sh, MARGIN, 3.95, 6.2, 1.8,
          'A steam railway mod for Satisfactory, explained for newcomers. '
          'Everything here is a plan, not a finished mod.', size=24, color=C['text'],
          spacing=1.05, name='Caption')
    picture(sh, 'R01', 7.13, 0.6, 5.6, 5.85, focus=(0.47, 0.5))
    credit(sh, 7.13, 6.48, 5.6, 'Early Locomotion project render. Early test model. Not final art.')
    sec = label(sh, MARGIN, 0.13, 8, 0.3, 'WELCOME', size=14, color=C['muted'], name='Section label')
    num = label(sh, RIGHT - 1.0, 7.06, 1.0, 0.3, str(n), size=14, color=C['muted'], align='r',
                name='Slide number')


# ------------------------------------------------------------------ 2
def s02(prs, n):
    b = add_base(prs, n, 'Welcome', 'How to read this guide', None,
                 'Every idea carries a tag in the top-right corner.')
    sh = b.shapes
    rows = [('Agreed', 'Decided.'), ('Proposed', 'Suggested, not confirmed.'),
            ('Provisional', 'A test number that may change.'),
            ('Undecided', 'Still an open question.'), ('Background', 'A real-world fact.')]
    g = group(sh, 'Tag legend', 'Legend of the five tags used in this guide and what each means.')
    y = 1.6
    for kind, meaning in rows:
        box(g.shapes, 1.4, y, 10.53, 0.72, fill=C['panel'], radius=0.08)
        chip(g.shapes, kind, 1.65, y + 0.12, w=2.3, h=0.48, size=16)
        label(g.shapes, 4.35, y, 7.4, 0.72, meaning, size=22, anchor='m')
        y += 0.88
    finish(b)


# ------------------------------------------------------------------ 3
def s03(prs, n):
    b = add_base(prs, n, 'Welcome', 'The idea in one picture', 'Agreed',
                 'Cheap track and trains. Real limits. You solve them with coal, water, signals and planning.')
    sh = b.shapes
    steps = [
        [[('Cheap to build', {'bold': True})]],
        [[('Real limits:', {'bold': True})], 'speed, range, signals, stations'],
        [[('Your answers:', {'bold': True})], 'tenders, service stops, signal boxes'],
        [[('Later:', {'bold': True})], 'electric monorail'],
    ]
    g = group(sh, 'Idea flow', 'Four steps: cheap to build, real limits, your answers, and later the electric monorail.')
    xs = [0.6, 3.777, 6.953, 10.13]
    for k, (x, txt) in enumerate(zip(xs, steps)):
        col = C['electric'] if k == 3 else C['orange']
        box(g.shapes, x, 1.95, 2.6, 1.6, fill=C['panel'], line=col, lw=2, radius=0.08,
            text=txt, size=18, align='c', anchor='m', margins=(0.15, 0.08, 0.15, 0.08))
        if k < 3:
            arrow(g.shapes, x + 2.66, 2.75, xs[k + 1] - 0.08, 2.75, lw=3)
    photo(sh, 'R01', 10.13, 4.35, 2.6, 1.5, focus=(0.47, 0.5))
    finish(b)


# ------------------------------------------------------------------ 4
def _machine(s, x, y, w, h, kind):
    box(s, x, y, w, h, fill=C['panel_alt'], line=C['steel'], lw=1.5, radius=0.06)
    if kind == 'miner':
        poly(s, [(x + w * 0.3, y + h), (x + w * 0.7, y + h), (x + w * 0.5, y + h + 0.45)],
             fill=C['steel'], line=C['ink'], lw=0.75)
        box(s, x + w * 0.2, y + 0.18, w * 0.6, 0.18, fill=C['orange'])
    elif kind == 'smelter':
        box(s, x + w * 0.62, y - 0.55, 0.28, 0.55, fill=C['steel'])
        box(s, x + w * 0.25, y + h * 0.45, w * 0.5, h * 0.3, fill=C['orange'], radius=0.05)
    elif kind == 'constructor':
        oval(s, x + w / 2, y + h * 0.5, h * 0.5, fill=None, line=C['orange'], lw=3)
        oval(s, x + w / 2, y + h * 0.5, h * 0.16, fill=C['orange'])
    elif kind == 'storage':
        line(s, x + 0.1, y + 0.1, x + w - 0.1, y + h - 0.1, C['steel'], 2)
        line(s, x + w - 0.1, y + 0.1, x + 0.1, y + h - 0.1, C['steel'], 2)


def _belt(s, x1, x2, y):
    box(s, x1, y, x2 - x1, 0.24, fill=C['coal'], line=C['steel'], lw=1)
    k = x1 + 0.2
    while k < x2 - 0.25:
        box(s, k, y + 0.04, 0.14, 0.16, fill=C['orange'], shape=MSO_SHAPE.CHEVRON)
        k += 0.36


def s04(prs, n):
    b = add_base(prs, n, 'The game', 'What is Satisfactory?', 'Background',
                 'A first-person game about building factories. You mine, make and move materials until it all runs itself.')
    sh = b.shapes
    box(sh, MARGIN, 1.55, 5.3, 0.62, fill=C['panel'], line=C['orange'], lw=1.5, radius=0.08,
        text=[[('Vanilla', {'bold': True, 'color': C['orange']}), ' = the game without mods.']],
        size=18, align='l', margins=(0.2, 0.03, 0.1, 0.03), name='Definition')
    g = group(sh, 'Factory strip', 'A simple factory line from shapes: a miner, conveyor belts, '
              'a smelter, a constructor and a storage crate, with the player standing beside it.')
    s = g.shapes
    gy = 4.55
    poly(s, [(0.9, gy), (2.3, gy), (2.1, gy + 0.55), (1.1, gy + 0.55)], fill='5C4A3A')
    box(s, 0.6, gy, 12.13, 0.05, fill=C['steel'])
    _machine(s, 1.0, gy - 1.2, 1.2, 1.2, 'miner')
    _belt(s, 2.2, 3.55, gy - 0.65)
    _machine(s, 3.55, gy - 1.2, 1.45, 1.2, 'smelter')
    _belt(s, 5.0, 6.35, gy - 0.65)
    _machine(s, 6.35, gy - 1.2, 1.45, 1.2, 'constructor')
    _belt(s, 7.8, 9.15, gy - 0.65)
    _machine(s, 9.15, gy - 1.1, 1.1, 1.1, 'storage')
    # player on a platform
    box(s, 10.85, gy - 0.35, 1.6, 0.35, fill=C['panel_alt'], line=C['steel'], lw=1)
    oval(s, 11.65, gy - 1.55, 0.32, fill=C['orange'])
    box(s, 11.45, gy - 1.37, 0.4, 0.75, fill=C['orange'], radius=0.12)
    names = [('Miner', 1.0, 1.2), ('Smelter', 3.55, 1.45), ('Constructor', 6.35, 1.45),
             ('Storage', 9.15, 1.1), ('You', 10.85, 1.6)]
    for txt, x, w in names:
        label(sh, x - 0.3, gy + 0.62, w + 0.6, 0.35, txt, size=16, align='c')
    label(sh, 2.2, gy - 1.05, 1.35, 0.32, 'Belt', size=14, align='c', color=C['muted'])
    label(sh, MARGIN, 5.75, 8, 0.32, 'By Coffee Stain Studios. Full release 2024.', size=14,
          color=C['muted'])
    finish(b)


# ------------------------------------------------------------------ 5
def _electric_train(s, x, y, scale=1.0):
    k = scale
    for j in range(2):
        cx = x + j * 1.1 * k
        box(s, cx, y - 0.6 * k, 1.02 * k, 0.48 * k, fill=C['electric'], line=C['ink'], lw=0.75,
            radius=0.12 * k)
        box(s, cx + 0.1 * k, y - 0.5 * k, 0.82 * k, 0.14 * k, fill=C['panel'])
    box(s, x - 0.05 * k, y - 0.1 * k, 2.25 * k, 0.1 * k, fill=C['steel'])


def s05(prs, n):
    b = add_base(prs, n, 'The game', 'Where the mod fits', 'Provisional',
                 "Technology unlocks in tiers. The mod adds steam railways early, before the game's electric trains.")
    sh = b.shapes
    steps = [('Tier 3', 'Early Steam Railways'), ('Tier 3 to 4', 'Railway Expansion'),
             ('Tier 4 to 5', 'Heavy Steam'), ('Tier 6', "Monorail: the game's electric trains")]
    g = group(sh, 'Tier staircase', 'A four-step staircase: Tier 3 early steam railways with a small engine, '
              'Tier 3 to 4 railway expansion with a medium engine, Tier 4 to 5 heavy steam with a large engine, '
              'and Tier 6 the electric monorail.')
    s = g.shapes
    base_y = 5.2
    w = 2.95
    for k, (tier, name) in enumerate(steps):
        x = 0.6 + k * (w + 0.11)
        top = base_y - (0.95 + 0.72 * k)
        col = C['electric'] if k == 3 else C['orange']
        box(s, x, top, w, base_y - top, fill=blend(col, C['bg'], 0.18), line=col, lw=2, radius=0.06)
        label(s, x + 0.15, top + 0.08, w - 0.3, 0.85,
              [[(tier, {'bold': True, 'size': 20, 'color': col})], name], size=16)
        if k == 0:
            train(s, x + 0.55, top - 0.04, 1, 0.85, name='Small engine')
        elif k == 1:
            train(s, x + 0.2, top - 0.04, 1, 0.8, tender=True, name='Medium engine')
        elif k == 2:
            train(s, x + 0.1, top - 0.04, 1, 0.95, tender=True, name='Large engine')
        else:
            _electric_train(s, x + 0.35, top - 0.04, 1.0)
    label(sh, MARGIN, 5.38, 3.2, 0.36, 'Tiers and prerequisites:', size=16, anchor='m')
    inline_chip(sh, 'Provisional', 3.2, 5.4)
    inline_chip(sh, 'Undecided', 3.2 + chip_w('Provisional') + 0.15, 5.4, text='Undecided (OQ01)')
    finish(b)


# ------------------------------------------------------------------ 6
def _rail_profile(cx, top):
    hw, hh = 0.2, 0.16      # head
    ww, wh = 0.07, 0.26     # web
    fw, fh = 0.27, 0.08     # foot
    y1, y2, y3 = top + hh, top + hh + wh, top + hh + wh + fh
    return [(cx - hw, top), (cx + hw, top), (cx + hw, y1), (cx + ww, y1), (cx + ww, y2),
            (cx + fw, y2), (cx + fw, y3), (cx - fw, y3), (cx - fw, y2), (cx - ww, y2),
            (cx - ww, y1), (cx - hw, y1)]


def s06(prs, n):
    b = add_base(prs, n, 'Steam railways', 'Track', 'Background',
                 'Two steel rails sit on sleepers over stone ballast. The sleepers keep the rails the right distance apart.')
    sh = b.shapes
    photo(sh, 'P01', MARGIN, 1.6, 4.2, 3.1)
    g = group(sh, 'Track cross-section', 'End-on cross-section of track: two rails on a sleeper, '
              'which sits on a bed of stone ballast. The gauge is the distance between the rails.')
    s = g.shapes
    cx = 8.75
    poly(s, [(cx - 2.75, 5.45), (cx + 2.75, 5.45), (cx + 2.15, 4.35), (cx - 2.15, 4.35)],
         fill=C['ballast'], line=C['steel'], lw=1)
    import random
    rnd = random.Random(4)
    for _ in range(26):
        yy = rnd.uniform(4.5, 5.3)
        half = 2.15 + (yy - 4.35) / 1.1 * 0.6 - 0.2
        oval(s, cx + rnd.uniform(-half, half), yy, rnd.uniform(0.1, 0.18), fill='5E636A',
             dh=rnd.uniform(0.07, 0.11))
    box(s, cx - 2.0, 4.02, 4.0, 0.36, fill=C['sleeper'], line='4E4236', lw=1)
    for d in (-1.15, 1.15):
        poly(s, _rail_profile(cx + d, 3.52), fill=C['steel'], line=C['ink'], lw=0.75)
    arrow(s, cx - 1.15 + 0.2, 3.3, cx + 1.15 - 0.2, 3.3, color=C['orange'], lw=2.25, head='both')
    label(sh, cx - 2.6, 2.55, 5.2, 0.6, [[('Gauge:', {'bold': True}), ' 1,435 mm on most railways']],
          size=18, align='c', anchor='b')
    # labels with leaders on the right
    for txt, ty, tx in (('Rail', 3.6, cx + 1.4), ('Sleeper', 4.2, cx + 2.0), ('Ballast', 4.95, cx + 2.45)):
        lg = group(sh, 'Label ' + txt, 'Label: ' + txt)
        line(lg.shapes, tx + 0.05, ty, 11.75, ty, C['orange'], 1.5)
        label(lg.shapes, 11.82, ty - 0.2, 0.95, 0.4, txt, size=18, anchor='m')
    finish(b)


# ------------------------------------------------------------------ 7
def s07(prs, n):
    b = add_base(prs, n, 'Steam railways', 'How a steam engine works', 'Background',
                 'Coal heats water. Steam pushes pistons. Rods turn the wheels.')
    sh = b.shapes
    g = group(sh, 'Steam engine cutaway', 'Side cutaway of a steam engine: a fire in the firebox heats water '
              'in the long boiler; steam drives a piston in the cylinder; rods turn the wheels; '
              'used steam leaves through the chimney.')
    s = g.shapes
    ry = 5.45                       # rail level
    box(s, 0.6, ry, 8.1, 0.08, fill=C['steel'])
    # cab and firebox (rear)
    box(s, 0.75, 2.2, 1.55, 2.3, fill=C['panel_alt'], line=C['steel'], lw=1.5)
    box(s, 0.65, 2.1, 1.75, 0.14, fill=C['steel'])
    box(s, 0.95, 2.45, 0.6, 0.5, fill=C['panel'], line=C['steel'], lw=1)
    box(s, 1.25, 3.3, 1.05, 1.15, fill='5A2E12', line=C['orange'], lw=2, radius=0.08)
    for fx, fh in ((1.45, 0.55), (1.72, 0.75), (2.0, 0.5)):
        poly(s, [(fx - 0.14, 4.4), (fx + 0.14, 4.4), (fx, 4.4 - fh)], fill=C['orange'])
        poly(s, [(fx - 0.07, 4.4), (fx + 0.07, 4.4), (fx, 4.4 - fh * 0.55)], fill=C['amber'])
    # boiler with water and steam space
    box(s, 2.3, 2.75, 5.3, 1.6, fill=C['panel'], line=C['steel'], lw=2, radius=0.3)
    box(s, 2.42, 3.55, 5.06, 0.68, fill=C['water'], radius=0.12)
    for ex, ey in ((3.1, 3.2), (4.0, 3.12), (4.9, 3.22), (5.8, 3.12), (6.7, 3.2)):
        oval(s, ex, ey, 0.5, fill=blend('F2F2F2', C['panel'], 0.85), dh=0.26)
    # smokebox and chimney
    box(s, 7.6, 2.75, 0.55, 1.6, fill=C['coal'], line=C['steel'], lw=1.5)
    box(s, 7.15, 2.05, 0.38, 0.72, fill=C['coal'], line=C['steel'], lw=1.5)
    for k, (ex, ey, d) in enumerate(((7.42, 1.82, 0.36), (7.72, 1.62, 0.46), (8.12, 1.5, 0.56))):
        oval(s, ex, ey, d, fill=blend('F2F2F2', C['bg'], 0.85), dh=d * 0.6)
    # frame, cylinder, piston
    box(s, 2.3, 4.4, 5.85, 0.12, fill=C['coal'])
    box(s, 6.95, 4.6, 1.15, 0.55, fill=C['steel'], line=C['ink'], lw=1, radius=0.05)
    box(s, 7.25, 4.65, 0.16, 0.45, fill=C['orange'])
    line(s, 6.4, 4.875, 7.25, 4.875, C['post'], 4)
    # wheels and rods
    wheels = (3.55, 4.75, 5.95)
    for wx in wheels:
        oval(s, wx, ry - 0.52, 1.04, fill=C['coal'], line=C['steel'], lw=3)
        oval(s, wx, ry - 0.52, 0.16, fill=C['steel'])
    pins = [(wx + 0.27, ry - 0.62) for wx in wheels]
    line(s, pins[0][0], pins[0][1], pins[2][0], pins[2][1], C['post'], 6)
    line(s, 6.4, 4.875, pins[1][0], pins[1][1], C['amber'], 6)
    for px, py in pins:
        oval(s, px, py, 0.12, fill=C['ink'])
    # exhaust arrow: cylinder round the front to the chimney
    curve_arrow(s, 7.75, 3.55, 0.75, 1.25, -55, 75, color=C['orange'], lw=2.5)
    # numbered badges
    badges = [(1, 0.95, 4.85, 'Fire'), (2, 2.55, 2.25, 'Steam'), (3, 7.15, 5.9, 'Piston'),
              (4, 4.45, 5.9, 'Wheels')]
    for num, bx, by_, txt in badges:
        badge(sh, bx, by_, num)
        label(sh, bx + 0.28, by_ - 0.2, 1.3, 0.4, txt, size=18, bold=True, anchor='m')
    label(sh, 8.6, 2.75, 1.2, 0.4, 'Exhaust', size=16, anchor='m', color=C['text'])
    photo(sh, 'P03', 9.13, 3.35, 3.6, 2.4)
    finish(b)


# ------------------------------------------------------------------ 8
def s08(prs, n):
    b = add_base(prs, n, 'Steam railways', 'Coal, water and steam', 'Agreed',
                 'Steam is used once and lost. So water runs out first. Coal lasts longer.')
    sh = b.shapes
    from lib import pill, text_width_in
    label(sh, 0.85, 1.42, 3.0, 0.42, 'Before', size=20, bold=True, color=C['orange'])
    label(sh, 7.85, 1.42, 3.0, 0.42, 'After', size=20, bold=True, color=C['orange'])
    g = group(sh, 'Coal and water run down', 'Before a long trip the coal and water tanks are both nearly full. '
              'After it, coal is just under half full but water is almost empty, '
              'because steam leaves the chimney and is not used again.')
    s = g.shapes
    ty, th = 2.75, 2.3
    box(s, 0.6, 1.9, 4.6, 3.75, fill=C['panel'], radius=0.1)
    box(s, 7.6, 1.9, 5.13, 3.75, fill=C['panel'], radius=0.1)
    tank(s, 1.2, ty, 1.6, th, 0.8, C['coal'], outline=C['coal_line'])
    tank(s, 3.1, ty, 1.6, th, 0.8, C['water'])
    tank(s, 8.05, ty, 1.6, th, 0.45, C['coal'], outline=C['coal_line'])
    tank(s, 9.95, ty, 1.6, th, 0.15, C['water'])
    for x in (1.2, 8.05):
        label(s, x, ty + th + 0.04, 1.6, 0.35, 'Coal', size=16, align='c')
    for x in (3.1, 9.95):
        label(s, x, ty + th + 0.04, 1.6, 0.35, 'Water', size=16, align='c')
    pill(s, 9.95 + 0.8 - 0.8, ty + th - 0.95, 1.6, 0.36, 'EMPTY FIRST', C['red_dark'], C['text'], 14)
    # chimney and steam along the top of the "after" panel
    box(s, 7.95, 2.22, 0.3, 0.42, fill=C['coal'], line=C['steel'], lw=1.5)
    for ex, ey, d in ((8.42, 2.12, 0.3), (8.8, 2.02, 0.38), (9.25, 1.98, 0.44)):
        oval(s, ex, ey, d, fill=blend('F2F2F2', C['panel'], 0.85), dh=d * 0.6)
    label(s, 9.6, 2.0, 3.0, 0.62, 'Steam leaves the chimney. It is not used again.', size=16,
          anchor='m')
    # arrow
    box(s, 5.45, 3.55, 1.95, 0.65, fill=C['orange'], shape=MSO_SHAPE.RIGHT_ARROW)
    label(s, 5.2, 2.95, 2.45, 0.5, 'One long trip', size=18, bold=True, align='c', anchor='b')
    note = 'In the mod, water is usually the tighter limit.'
    label(sh, MARGIN, 5.82, 7.5, 0.38, note, size=18, anchor='m')
    inline_chip(sh, 'Agreed', MARGIN + text_width_in(note, 18) + 0.2, 5.85)
    finish(b)


# ------------------------------------------------------------------ 9
def s09(prs, n):
    b = add_base(prs, n, 'Steam railways', 'Tender or tank?', 'Background',
                 'A tender engine pulls a coal and water wagon. A tank engine carries both itself.')
    sh = b.shapes
    from lib import set_alt
    label(sh, MARGIN, 1.42, 5.8, 0.45, 'Tender engine', size=20, bold=True, color=C['orange'])
    label(sh, 6.93, 1.42, 5.8, 0.45, 'Tank engine', size=20, bold=True, color=C['orange'])
    photo(sh, 'P04', MARGIN, 1.9, 5.8, 3.65)
    photo(sh, 'P05', 6.93, 1.9, 5.8, 3.65)
    # callouts with leader lines to the parts of each engine (positions read off the photos)
    callouts = [
        ('Tender: coal and water behind the cab', 2.25, 2.05, 3.2, 0.7, (1.59, 3.8), (2.75, 2.75)),
        ('Side tanks: water.', 7.45, 2.05, 2.15, 0.4, (9.25, 3.76), (8.9, 2.45)),
        ('Bunker: coal.', 10.55, 2.05, 1.65, 0.4, (10.9, 3.55), (11.0, 2.45)),
    ]
    for txt, x, y, w, h, tgt, src in callouts:
        g = group(sh, 'Callout', 'Label: ' + txt)
        line(g.shapes, src[0], src[1], tgt[0], tgt[1], C['orange'], 2.25)
        oval(g.shapes, tgt[0], tgt[1], 0.16, fill=C['orange'], line=C['ink'], lw=1)
        box(g.shapes, x, y, w, h, fill=C['bg'], line=C['orange'], lw=1.5, radius=0.06, text=txt,
            size=16, align='l', margins=(0.1, 0.03, 0.08, 0.03))
    finish(b)


# ------------------------------------------------------------------ 10
def s10(prs, n):
    b = add_base(prs, n, 'Steam railways', 'Taking on water', 'Background',
                 'A water crane swings over the engine. Water flows down from a tank above.')
    sh = b.shapes
    photo(sh, 'P06', MARGIN, 1.5, 3.2, 4.25, focus=(0.5, 0.38))
    g = group(sh, 'Water crane diagram', 'A water tower on legs feeds a pipe down to a water crane; '
              'the crane arm swings over the tender hatch and gravity sends the water down.')
    s = g.shapes
    ry = 5.55
    for lx in (4.75, 5.2, 5.95, 6.4):
        box(s, lx, 3.35, 0.1, 2.2, fill=C['steel'])
    box(s, 4.6, 2.05, 1.95, 1.3, fill=C['water'], line='FFFFFF', lw=2, radius=0.12)
    polyline(s, [(6.1, 3.35), (6.1, 5.25), (7.85, 5.25), (7.85, 3.75), (9.25, 3.75)],
             color=C['water'], lw=9)
    box(s, 7.75, 3.6, 0.2, 1.75, fill=C['steel'], line=C['ink'], lw=0.75)
    box(s, 7.75, 3.65, 1.55, 0.18, fill=C['steel'], line=C['ink'], lw=0.75)
    box(s, 9.17, 3.75, 0.16, 0.45, fill=C['steel'], line=C['ink'], lw=0.75)
    for dy in (4.33, 4.5):
        oval(s, 9.25, dy, 0.09, fill=C['water'], dh=0.12)
    box(s, 7.95, ry, 2.45, 0.07, fill=C['steel'])
    box(s, 8.0, 4.6, 2.4, 0.75, fill=C['wagon'], line=C['ink'], lw=1, radius=0.05)
    oval(s, 9.25, 4.6, 0.42, fill=C['coal'], line=C['steel'], lw=1.5, dh=0.14)
    for wx in (8.5, 10.1):
        oval(s, wx, ry - 0.11, 0.22, fill=C['coal'], line=C['steel'], lw=1)
    # the arm swings round the column top: arc over it, ending above the hatch
    curve_arrow(s, 8.4, 3.62, 0.95, 0.42, 165, 15, color=C['orange'], lw=2.5)
    arrow(s, 6.35, 3.6, 6.35, 4.2, C['orange'], 2.25, hs='lg')
    arrow(s, 6.35, 4.35, 6.35, 4.95, C['orange'], 2.25, hs='lg')
    label(s, 7.1, 2.62, 3.4, 0.5, 'Arm swings over the hatch', size=16, align='c', anchor='b')
    label(s, 6.5, 3.95, 1.1, 0.4, 'Gravity', size=16, anchor='m')
    photo(sh, 'P07', 10.53, 3.95, 2.2, 1.6, focus=(0.5, 0.33))
    finish(b)


SLIDES = [s01, s02, s03, s04, s05, s06, s07, s08, s09, s10]
