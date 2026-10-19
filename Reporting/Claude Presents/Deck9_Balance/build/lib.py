"""Shape library for the Deck 9 build (native shapes only).

All coordinates are inches on a 13.333 x 7.5 in canvas. Containers (`c`) are either
`slide.shapes` or `group.shapes`; child coordinates in a group are absolute because
python-pptx keeps an identity group transform.
"""
import math
import re
from lxml import etree
from PIL import ImageFont, Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

# ---------------------------------------------------------------- style tokens
BG = "1E2126"; PANEL = "2B2F36"; PANEL_ALT = "343A43"; LINE = "4A515B"
TEXT = "F2F2F2"; MUTED = "A9B0B8"; ORANGE = "F28C28"; RED = "E5484D"
GREEN = "3FB950"; AMBER = "F2C14E"; WATER = "4DA3FF"; ELEC = "6CA0DC"
COAL = "1B1D21"; STEEL = "9AA0A6"; SLEEPER = "6B5B4B"; PURPLE = "B084F5"
DARK = "101418"; WAGON = "6C7480"; ORANGE_DK = "C46E14"

FONT = "Arial"
LANG = "en-GB"
SLIDE_W, SLIDE_H = 13.333, 7.5

CHIPS = {  # status -> (label, fill, text colour, outline)
    "A": ("AGREED", GREEN, DARK, None),
    "P": ("PROPOSED", PURPLE, DARK, None),
    "V": ("PROVISIONAL", WATER, DARK, None),
    "U": ("UNDECIDED", None, TEXT, MUTED),
    "B": ("BACKGROUND", LINE, TEXT, None),
}

_REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
_BLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
_fcache = {}


def _font(size, bold):
    k = (round(size * 10), bold)
    if k not in _fcache:
        _fcache[k] = ImageFont.truetype(_BLD if bold else _REG, int(round(size * 10)))
    return _fcache[k]


def tw(text, size, bold=False):
    """Width of `text` in inches at `size` pt (Liberation Sans = Arial metrics)."""
    return _font(size, bold).getlength(text) / 10.0 / 72.0


def wrap(text, size, bold, width):
    lines, cur = [], ""
    for word in text.split(" "):
        trial = (cur + " " + word).strip()
        if tw(trial, size, bold) <= width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    return lines


def rgb(h):
    return RGBColor.from_string(h)


def E(v):
    return Emu(int(round(v * 914400)))


# ---------------------------------------------------------------- primitives
def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _fill_line(shape, fill, line, lw):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)


def shape(c, kind, x, y, w, h, fill=None, line=None, lw=1.5, name=None, radius=None,
          rot=None, dash=None):
    s = c.add_shape(kind, E(x), E(y), E(w), E(h))
    _strip_style(s)
    _fill_line(s, fill, line, lw)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if name:
        s.name = name
    if rot:
        s.rotation = rot
    if dash:
        s.line.dash_style = dash
    return s


def rect(c, x, y, w, h, **kw):
    return shape(c, MSO_SHAPE.RECTANGLE, x, y, w, h, **kw)


def rrect(c, x, y, w, h, radius=0.08, **kw):
    return shape(c, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, radius=radius, **kw)


def oval(c, x, y, w, h, **kw):
    return shape(c, MSO_SHAPE.OVAL, x, y, w, h, **kw)


def circle(c, cx, cy, r, **kw):
    return shape(c, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, **kw)


def poly(c, pts, fill=None, line=None, lw=1.5, closed=True, name=None):
    """Freeform polygon / polyline through `pts` (inches)."""
    fb = c.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=closed)
    s = fb.convert_to_shape()
    _strip_style(s)
    _fill_line(s, fill if closed else None, line, lw)
    if name:
        s.name = name
    return s


def line(c, x1, y1, x2, y2, color=LINE, lw=2.25, dash=None, head=None, tail=None, name=None):
    """Straight connector. `tail` = arrowhead at (x2, y2), `head` = at (x1, y1)."""
    cn = c.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _strip_style(cn)
    cn.line.color.rgb = rgb(color)
    cn.line.width = Pt(lw)
    if dash:
        cn.line.dash_style = dash
    ln = cn.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", on)
            e.set("w", "med")
            e.set("len", "med")
    if name:
        cn.name = name
    return cn


def _lang(el):
    for r in el.iter(qn("a:rPr"), qn("a:endParaRPr")):
        r.set("lang", LANG)


def T(text, size=18, bold=False, color=TEXT, italic=False):
    return (text, dict(size=size, bold=bold, color=color, italic=italic))


def text(c, x, y, w, h, paras, align="l", anchor="t", wrap_=True, inset=(0.05, 0.03, 0.05, 0.03),
         name=None, spacing=None):
    """Text box. `paras` is a str, a list of runs (one paragraph) or a list of lists of runs."""
    tb = c.add_textbox(E(x), E(y), E(w), E(h))
    if name:
        tb.name = name
    fill_tf(tb.text_frame, paras, align, anchor, wrap_, inset, spacing)
    return tb


def fill_tf(tf, paras, align="l", anchor="t", wrap_=True, inset=(0.05, 0.03, 0.05, 0.03), spacing=None):
    tf.word_wrap = wrap_
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(v) for v in inset]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [[T(paras)]]
    elif paras and isinstance(paras[0], tuple):
        paras = [paras]
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        if spacing:
            p.line_spacing = spacing
        for txt, f in runs:
            r = p.add_run()
            r.text = txt
            r.font.name = FONT
            r.font.size = Pt(f["size"])
            r.font.bold = f["bold"]
            r.font.italic = f.get("italic", False)
            r.font.color.rgb = rgb(f["color"])
    _lang(tf._txBody)


def label(c, x, y, w, h, s, size=14, bold=False, color=TEXT, align="c", anchor="m", **kw):
    return text(c, x, y, w, h, [[T(s, size, bold, color)]], align=align, anchor=anchor, **kw)


def box_text(shp, paras, align="c", anchor="m", inset=(0.08, 0.04, 0.08, 0.04), wrap_=True):
    fill_tf(shp.text_frame, paras, align, anchor, wrap_, inset)


# ---------------------------------------------------------------- groups & alt text
def set_alt(shp, descr, name=None):
    el = shp._element
    cnv = None
    for tag in ("p:nvSpPr", "p:nvGrpSpPr", "p:nvPicPr", "p:nvCxnSpPr", "p:nvGraphicFramePr"):
        nv = el.find(qn(tag))
        if nv is not None:
            cnv = nv.find(qn("p:cNvPr"))
            break
    if cnv is not None:
        cnv.set("descr", descr)
        if name:
            cnv.set("name", name)


def group(slide_or_group, name, descr):
    g = slide_or_group.shapes.add_group_shape() if hasattr(slide_or_group, "shapes") else slide_or_group.add_group_shape()
    set_alt(g, descr, name)
    return g


# ---------------------------------------------------------------- chips
def chip_at(c, x, y, status, w=1.6, h=0.4, name="Status chip"):
    lab, fill, tc, ol = CHIPS[status]
    s = shape(c, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=ol, lw=2, radius=0.12,
              name=name)
    box_text(s, [[T(lab, 14, True, tc)]], inset=(0, 0, 0, 0), wrap_=False)
    return s


def status_chip(slide, status):
    s = chip_at(slide.shapes, SLIDE_W - 0.6 - 1.6, 0.5, status)
    set_alt(s, "Status tag: " + CHIPS[status][0].title())
    return s


def inline_chip(c, x, y, status, txt=None, size=14, h=0.36, name=None, right=False):
    """Pill chip. With `txt` it carries a full label (e.g. 'Exact recipes: undecided (OQ02)');
    without it the status word is shown. Returns the chip's width in inches."""
    lab, fill, tc, ol = CHIPS[status]
    shown = txt if txt else lab
    bold = True
    w = tw(shown, size, bold) + 0.34
    if right:
        x = x - w
    s = shape(c, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=ol, lw=2, radius=h / 2,
              name=name or "Inline chip: " + shown)
    box_text(s, [[T(shown, size, bold, tc)]], inset=(0, 0, 0, 0), wrap_=False)
    return w


def inline_chip_width(txt, size=14):
    return tw(txt, size, True) + 0.34


# ---------------------------------------------------------------- slide scaffolding
def setup_presentation():
    prs = Presentation()
    prs.slide_width = E(SLIDE_W)
    prs.slide_height = E(SLIDE_H)
    prs.core_properties.title = "How the Mod Is Balanced, Explained"
    prs.core_properties.author = "Claude (Early Locomotion project)"
    prs.core_properties.language = LANG
    prs.core_properties.subject = "Explainer deck 9 for the Satisfactory mod Early Locomotion"
    _theme(prs)
    return prs


def _theme(prs):
    master = prs.slide_master
    # dark background and light text on the master
    cSld = master._element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg"))
    if bg is not None:
        cSld.remove(bg)
    bg = etree.fromstring(
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:bgPr><a:solidFill>'
        '<a:srgbClr val="%s"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>' % BG)
    cSld.insert(0, bg)
    cm = master._element.find(qn("p:clrMap"))
    cm.set("bg1", "dk1"); cm.set("tx1", "lt1"); cm.set("bg2", "dk2"); cm.set("tx2", "lt2")
    theme_part = master.part.part_related_by(RT.THEME)
    xml = theme_part.blob.decode("utf8")
    cols = {"dk1": BG, "lt1": TEXT, "dk2": PANEL, "lt2": MUTED, "accent1": ORANGE, "accent2": WATER,
            "accent3": GREEN, "accent4": AMBER, "accent5": PURPLE, "accent6": RED,
            "hlink": WATER, "folHlink": PURPLE}
    for k, v in cols.items():
        xml = re.sub(r"<a:%s>.*?</a:%s>" % (k, k), '<a:%s><a:srgbClr val="%s"/></a:%s>' % (k, v, k), xml,
                     flags=re.S)
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r"\1Arial\2", xml)
    theme_part._blob = xml.encode("utf8")


def new_slide(prs, title, status=None, title_size=36):
    layout = prs.slide_layouts[5]  # Title Only
    s = prs.slides.add_slide(layout)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = E(0.6), E(0.45), E(10.3), E(0.8)
    fill_tf(t.text_frame, [[T(title, title_size, True, TEXT)]], "l", "m", True, (0, 0, 0, 0))
    if status:
        status_chip(s, status)
    return s


def add_caption(slide, cap):
    g = group(slide, "Caption", "Caption: " + cap)
    y, h = 6.2, 0.9
    rrect(g.shapes, 0.6, y, 12.13, h, radius=0.08, fill=PANEL, name="Caption panel")
    rect(g.shapes, 0.6, y, 0.08, h, fill=ORANGE, name="Caption rule")
    text(g.shapes, 0.93, y, 11.6, h, [[T(cap, 20, False, TEXT)]], anchor="m", inset=(0, 0, 0, 0),
         name="Caption text")
    return g


def add_number(slide, n):
    text(slide.shapes, SLIDE_W - 0.6 - 0.8, 7.12, 0.8, 0.3, [[T(str(n), 14, False, MUTED)]], align="r",
         anchor="m", inset=(0, 0, 0, 0), name="Slide number")


def add_picture_cover(slide, path, x, y, w, h, cx=0.5, cy=0.5, descr="", radius=0.08, name="Picture"):
    """Crop (never stretch) so the picture covers the frame; (cx, cy) = crop centre 0..1."""
    iw, ih = Image.open(path).size
    fr = w / h
    if iw / ih > fr:  # too wide: crop sides
        vis_w = ih * fr
        left = min(max(cx * iw - vis_w / 2, 0), iw - vis_w) / iw
        crop = (left, 0, 1 - left - vis_w / iw, 0)
    else:
        vis_h = iw / fr
        top = min(max(cy * ih - vis_h / 2, 0), ih - vis_h) / ih
        crop = (0, top, 0, 1 - top - vis_h / ih)
    pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    prst = pic._element.spPr.find(qn("a:prstGeom"))
    av = prst.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(prst, qn("a:avLst"))
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", "val %d" % int(50000 * radius / min(w, h)))
    set_alt(pic, descr, name)
    return pic


# ---------------------------------------------------------------- reusable drawings
def rails(c, x, y, length, gauge=0.28, color=STEEL, sleepers=True, sl_color=SLEEPER, pitch=0.35):
    """Straight two-rail track with sleepers; (x, y) = left end, upper rail."""
    if sleepers:
        n = int(length // pitch)
        for i in range(n + 1):
            rect(c, x + 0.05 + i * pitch, y - 0.06, 0.07, gauge + 0.18, fill=sl_color, name="Sleeper")
    rect(c, x, y, length, 0.06, fill=color, name="Rail")
    rect(c, x, y + gauge, length, 0.06, fill=color, name="Rail")


def arc_pts(cx, cy, r, a0, a1, n=24):
    """Points on a circular arc (angles in degrees, screen coordinates)."""
    return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / n)),
             cy + r * math.sin(math.radians(a0 + (a1 - a0) * i / n))) for i in range(n + 1)]


def curved_track(c, pts, gauge=0.3, sleeper_len=0.58, n_sleepers=14, rail_w=2.5):
    """Two rails offset either side of the centreline `pts`, plus rotated sleepers."""
    def normals():
        out = []
        for i, (px, py) in enumerate(pts):
            a = pts[max(i - 1, 0)]; b = pts[min(i + 1, len(pts) - 1)]
            tx, ty = b[0] - a[0], b[1] - a[1]
            d = math.hypot(tx, ty) or 1
            out.append((-ty / d, tx / d, math.degrees(math.atan2(ty, tx))))
        return out
    nm = normals()
    # sleepers first (under the rails), evenly spaced along the polyline
    seg = [math.hypot(pts[i + 1][0] - pts[i][0], pts[i + 1][1] - pts[i][1]) for i in range(len(pts) - 1)]
    total = sum(seg)
    for k in range(n_sleepers):
        target = (k + 0.5) * total / n_sleepers
        acc = 0
        for i, sgl in enumerate(seg):
            if acc + sgl >= target:
                f = (target - acc) / sgl
                px = pts[i][0] + f * (pts[i + 1][0] - pts[i][0])
                py = pts[i][1] + f * (pts[i + 1][1] - pts[i][1])
                ang = nm[i][2]
                rect(c, px - 0.045, py - sleeper_len / 2, 0.09, sleeper_len, fill=SLEEPER, rot=ang,
                     name="Sleeper")
                break
            acc += sgl
    for side in (-1, 1):
        rp = [(p[0] + side * gauge / 2 * n[0], p[1] + side * gauge / 2 * n[1]) for p, n in zip(pts, nm)]
        poly(c, rp, line=STEEL, lw=rail_w, closed=False, name="Rail")


def wheel(c, cx, cy, r, color=COAL, rim=STEEL):
    circle(c, cx, cy, r, fill=color, line=rim, lw=1.5, name="Wheel")
    circle(c, cx, cy, r * 0.28, fill=rim, name="Hub")


def loco(c, x, y, s=1.0, tender=False, body=ORANGE, name="Locomotive"):
    """Side-view steam locomotive facing right. (x, y) = top-left of its bounding box.
    Width 2.6*s (+1.4*s with tender), height 1.4*s."""
    ox = x + (1.4 * s if tender else 0)
    if tender:
        rect(c, x, y + 0.42 * s, 1.3 * s, 0.6 * s, fill=PANEL_ALT, line=STEEL, lw=1, name="Tender")
        poly(c, [(x + 0.1 * s, y + 0.42 * s), (x + 0.3 * s, y + 0.26 * s), (x + 0.6 * s, y + 0.2 * s),
                 (x + 0.95 * s, y + 0.27 * s), (x + 1.2 * s, y + 0.42 * s)], fill=COAL, line=None,
             name="Coal")
        rect(c, x + 1.3 * s, y + 0.88 * s, 0.1 * s, 0.05 * s, fill=STEEL, name="Coupling")
        for wx in (0.3, 1.0):
            wheel(c, x + wx * s, y + 1.17 * s, 0.2 * s)
    rect(c, ox, y + 0.93 * s, 2.6 * s, 0.1 * s, fill=COAL, name="Footplate")
    rrect(c, ox + 0.9 * s, y + 0.38 * s, 1.6 * s, 0.58 * s, radius=0.15 * s, fill=body, name="Boiler")
    rect(c, ox + 2.38 * s, y + 0.38 * s, 0.14 * s, 0.58 * s, fill=COAL, name="Smokebox")
    rect(c, ox + 2.2 * s, y + 0.12 * s, 0.13 * s, 0.28 * s, fill=COAL, name="Chimney")
    rect(c, ox + 2.15 * s, y + 0.08 * s, 0.23 * s, 0.06 * s, fill=COAL, name="Chimney top")
    shape(c, MSO_SHAPE.CHORD, ox + 1.5 * s, y + 0.24 * s, 0.34 * s, 0.3 * s, fill=ORANGE_DK,
          name="Dome", rot=0)
    rect(c, ox + 0.08 * s, y + 0.18 * s, 0.9 * s, 0.78 * s, fill=ORANGE_DK, name="Cab")
    rect(c, ox + 0.02 * s, y + 0.1 * s, 1.02 * s, 0.09 * s, fill=COAL, name="Cab roof")
    rect(c, ox + 0.3 * s, y + 0.32 * s, 0.4 * s, 0.28 * s, fill=COAL, name="Cab window")
    for wx in (0.55, 1.35, 2.15):
        wheel(c, ox + wx * s, y + 1.17 * s, 0.22 * s)
    line(c, ox + 0.55 * s, y + 1.17 * s, ox + 2.15 * s, y + 1.17 * s, color=STEEL, lw=2, name="Coupling rod")
    return (2.6 + (1.4 if tender else 0)) * s


def wagon(c, x, y, s=1.0, fill=WAGON, name="Wagon"):
    """Side-view wagon, 1.5*s x 0.85*s (incl. wheels)."""
    rect(c, x, y, 1.5 * s, 0.62 * s, fill=fill, line=STEEL, lw=1, name=name)
    for i in (1, 2, 3):
        line(c, x + i * 0.375 * s, y + 0.04 * s, x + i * 0.375 * s, y + 0.58 * s, color=PANEL_ALT, lw=1)
    for wx in (0.3, 1.2):
        wheel(c, x + wx * s, y + 0.7 * s, 0.14 * s)
    return 1.5 * s


def electric_train(c, x, y, s=1.0, n_cars=1, beam=True):
    """Vanilla electric (monorail) train: blue cars on a beam. (x, y) = top-left. Width 2.4*s per car."""
    w = 2.4 * s
    for i in range(n_cars):
        cx = x + i * (w + 0.08 * s)
        rrect(c, cx, y, w, 0.82 * s, radius=0.25 * s, fill=ELEC, name="Electric car")
        rect(c, cx + 0.2 * s, y + 0.18 * s, w - 0.5 * s, 0.26 * s, fill=COAL, name="Windows")
    total = n_cars * w + (n_cars - 1) * 0.08 * s
    if beam:
        rect(c, x - 0.1 * s, y + 0.9 * s, total + 0.2 * s, 0.14 * s, fill=STEEL, name="Monorail beam")
        for px in (0.4, total - 0.5):
            rect(c, x + px * 1.0, y + 1.04 * s, 0.12 * s, 0.22 * s, fill=LINE, name="Pylon")
    return total


def coin_stack(c, x, y, n, w=0.7, h=0.14, fill=AMBER):
    """n flat coins stacked; (x, y) = top-left of the whole stack."""
    for i in range(n):
        yy = y + (n - 1 - i) * h
        shape(c, MSO_SHAPE.ROUNDED_RECTANGLE, x, yy, w, h, fill=fill, line=DARK, lw=1, radius=h / 2,
              name="Coin")
    return n * h


def tank_gauge(c, x, y, w, h, level, color, name="Gauge"):
    rrect(c, x, y, w, h, radius=0.08, fill=COAL, line=STEEL, lw=1.5, name=name + " tank")
    fh = (h - 0.1) * level
    rect(c, x + 0.05, y + h - 0.05 - fh, w - 0.1, fh, fill=color, name=name + " fill")


def numbered_badge(c, cx, cy, n, r=0.2, fill=ORANGE, size=14):
    s = circle(c, cx, cy, r, fill=fill, name="Badge " + str(n))
    box_text(s, [[T(str(n), size, True, DARK)]], inset=(0, 0, 0, 0), wrap_=False)
    return s


def x_mark(c, cx, cy, r, color=RED, lw=6):
    line(c, cx - r, cy - r, cx + r, cy + r, color=color, lw=lw, name="Cross")
    line(c, cx - r, cy + r, cx + r, cy - r, color=color, lw=lw, name="Cross")


def finish(prs, path):
    for sl in prs.slides:
        for el in sl._element.iter(qn("a:r")):
            rpr = el.find(qn("a:rPr"))
            if rpr is None:
                rpr = etree.Element(qn("a:rPr"))
                el.insert(0, rpr)
            rpr.set("lang", LANG)
        for r in sl._element.iter(qn("a:rPr"), qn("a:endParaRPr")):
            r.set("lang", LANG)
    prs.save(path)
