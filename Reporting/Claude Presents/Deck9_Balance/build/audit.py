"""Audit a built deck: word counts, smallest font, alt text, notes, chips, contrast, text fit, overlaps.

usage: audit.py deck.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from lib import tw, wrap, BG

STATUS_WORDS = {"AGREED", "PROPOSED", "PROVISIONAL", "UNDECIDED", "BACKGROUND"}
CREDIT_FREE = {1, 2, 18}          # slides exempt from the status-chip rule
EXPECT_TAG = {3: "A", 4: "A", 5: "A", 6: "P", 7: "A", 8: "A", 9: "P", 10: "U", 11: "V", 12: "A", 13: "A",
              14: "V", 15: "P", 16: "V", 17: "U"}


def lum(h):
    def f(v):
        v /= 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def contrast(a, b):
    la, lb = lum(a), lum(b)
    if la < lb:
        la, lb = lb, la
    return (la + 0.05) / (lb + 0.05)


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield sh
            yield from walk(sh.shapes)
        else:
            yield sh


def fill_of(sh):
    try:
        if sh.fill.type == 1:
            return str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return None


def runs(sh):
    if not sh.has_text_frame:
        return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            yield p, r


def seg_hits_rect(x1, y1, x2, y2, rx1, ry1, rx2, ry2):
    t0, t1 = 0.0, 1.0
    dx, dy = x2 - x1, y2 - y1
    for p, q in ((-dx, x1 - rx1), (dx, rx2 - x1), (-dy, y1 - ry1), (dy, ry2 - y1)):
        if p == 0:
            if q < 0:
                return False
        else:
            r = q / p
            if p < 0:
                if r > t1:
                    return False
                t0 = max(t0, r)
            else:
                if r < t0:
                    return False
                t1 = min(t1, r)
    return t0 <= t1


def audit(path):
    prs = Presentation(path)
    problems = []
    print("slide | words(counted) words(all-text) | min pt | chip | notes")
    for n, sl in enumerate(prs.slides, 1):
        leaves = [s for s in walk(sl.shapes)]
        counted = allw = 0
        minpt = 99
        chip = None
        texts = []   # (x1,y1,x2,y2,desc) estimated text rectangles
        for sh in leaves:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            is_title = sh.is_placeholder
            for p, r in runs(sh):
                pt = r.font.size.pt if r.font.size else None
                if pt is None:
                    problems.append((n, "run without explicit size: " + r.text[:30]))
                else:
                    if pt < minpt and not sh.name.startswith("Credit"):
                        minpt = pt
                    if pt < 14 and not sh.name.startswith("Credit"):
                        problems.append((n, "text below 14 pt: %r (%s pt)" % (r.text[:30], pt)))
                    if sh.name.startswith("Credit") and pt < 10:
                        problems.append((n, "credit below 10 pt"))
                if r.font.name != "Arial":
                    problems.append((n, "font not Arial: " + str(r.font.name)))
            txt = sh.text_frame.text if sh.has_text_frame else ""
            words = len(txt.split())
            allw += words
            if sh.name == "Status chip":
                chip = txt.strip()
                continue
            if sh.name.startswith("Slide number") or is_title or sh.name.startswith("Credit"):
                continue
            if sh.name.startswith("Legend chip") or (txt.strip() in STATUS_WORDS):
                continue
            counted += words
            # fit check
            if sh.has_text_frame and txt.strip():
                tf = sh.text_frame
                w_in = Emu(sh.width).inches - Emu(tf.margin_left).inches - Emu(tf.margin_right).inches
                h_in = Emu(sh.height).inches - Emu(tf.margin_top).inches - Emu(tf.margin_bottom).inches
                need = 0
                maxw = 0
                for para in tf.paragraphs:
                    if not para.runs:
                        continue
                    size = max(r.font.size.pt for r in para.runs)
                    bold = any(r.font.bold for r in para.runs)
                    t = "".join(r.text for r in para.runs)
                    if tf.word_wrap:
                        lines = wrap(t, size, bold, w_in)
                    else:
                        lines = [t]
                    need += len(lines) * size * 1.2 / 72
                    maxw = max(maxw, max(tw(l, size, bold) for l in lines))
                    if not tf.word_wrap and tw(t, size, bold) > w_in + 0.02:
                        problems.append((n, "no-wrap text wider than its box: %r (%.2f > %.2f)" % (t[:30], tw(t, size, bold), w_in)))
                if need > h_in + 0.03:
                    problems.append((n, "text taller than box: %r need %.2f have %.2f" % (txt[:30], need, h_in)))
                x0 = Emu(sh.left).inches + Emu(tf.margin_left).inches
                al = tf.paragraphs[0].alignment
                if al is not None and int(al) == 2:
                    tx1 = x0 + (w_in - maxw) / 2
                elif al is not None and int(al) == 3:
                    tx1 = x0 + w_in - maxw
                else:
                    tx1 = x0
                anchor = tf.vertical_anchor
                top = Emu(sh.top).inches + Emu(tf.margin_top).inches
                if anchor is not None and int(anchor) == 3:   # middle
                    ty1 = top + (h_in - need) / 2
                elif anchor is not None and int(anchor) == 4:  # bottom
                    ty1 = top + h_in - need
                else:
                    ty1 = top
                texts.append((tx1, ty1, tx1 + maxw, ty1 + need, txt[:25], sh))
                # contrast
                f = fill_of(sh) or BG
                for p, r in runs(sh):
                    try:
                        col = str(r.font.color.rgb)
                    except Exception:
                        continue
                    c = contrast(col, f)
                    if c < 4.5:
                        problems.append((n, "low contrast %.1f for %r on %s" % (c, r.text[:25], f)))
        # bounds
        for sh in leaves:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            if sh.left is None:
                continue
            l, t, w, h = [Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height)]
            if getattr(sh, "rotation", 0):
                continue
            if l < -0.01 or t < -0.01 or l + w > 13.35 or t + h > 7.51:
                problems.append((n, "shape outside slide: %s" % sh.name))
        # text/text overlaps
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                a, b = texts[i], texts[j]
                if a[0] < b[2] - 0.02 and b[0] < a[2] - 0.02 and a[1] < b[3] - 0.02 and b[1] < a[3] - 0.02:
                    problems.append((n, "text overlap: %r / %r" % (a[4], b[4])))
        # connector lines over text
        for sh in leaves:
            if sh.shape_type == MSO_SHAPE_TYPE.LINE or sh.name.startswith("Arrow") or sh.name.startswith("Leader"):
                try:
                    x1, y1, x2, y2 = [Emu(v).inches for v in (sh.begin_x, sh.begin_y, sh.end_x, sh.end_y)]
                except Exception:
                    continue
                for tx in texts:
                    if tx[5] is sh:
                        continue
                    if seg_hits_rect(x1, y1, x2, y2, tx[0], tx[1], tx[2], tx[3]):
                        problems.append((n, "line %r crosses text %r" % (sh.name, tx[4])))
        # alt text on groups and pictures
        for sh in leaves:
            el = sh._element
            cnv = None
            for tag in ("p:nvGrpSpPr", "p:nvPicPr"):
                nv = el.find(qn(tag))
                if nv is not None:
                    cnv = nv.find(qn("p:cNvPr"))
            if cnv is not None and not cnv.get("descr"):
                problems.append((n, "missing alt text on " + sh.name))
        # status chip rule
        if n in EXPECT_TAG:
            if chip is None or chip[:1] != {"A": "A", "P": "P", "V": "P", "U": "U"}[EXPECT_TAG[n]][:1] and False:
                pass
        want = {"A": "AGREED", "P": "PROPOSED", "V": "PROVISIONAL", "U": "UNDECIDED"}.get(EXPECT_TAG.get(n))
        if n in EXPECT_TAG and chip != want:
            problems.append((n, "status chip is %r, expected %r" % (chip, want)))
        if n not in EXPECT_TAG and chip:
            problems.append((n, "unexpected status chip %r" % chip))
        has_notes = sl.has_notes_slide
        if has_notes:
            problems.append((n, "has notes slide"))
        if sl._element.find(qn("p:transition")) is not None or sl._element.find(qn("p:timing")) is not None:
            problems.append((n, "has transition/animation"))
        for ph in sl.placeholders:
            if ph.has_text_frame and not ph.text_frame.text.strip():
                problems.append((n, "empty placeholder"))
        print("%5d | %6d %6d | %6s | %-11s | %s" % (n, counted, allw, minpt if minpt != 99 else "-", chip or "-",
                                                  "notes" if has_notes else "none"))
    print()
    if problems:
        print("PROBLEMS (%d):" % len(problems))
        for n, m in problems:
            print("  slide %d: %s" % (n, m))
    else:
        print("No problems found.")
    return problems


if __name__ == "__main__":
    audit(sys.argv[1])
