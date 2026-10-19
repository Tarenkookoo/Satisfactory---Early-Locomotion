"""Slides 1 to 8 of Deck 4."""
import math

from pptx.enum.dml import MSO_LINE

from lib import *  # noqa: F401,F403
from lib import _font  # noqa: F401

R02_ALT = ("High-detail render of the mod's two-rail track. Visual development only: not Unreal-tested, "
           "not production-ready.")
R02_CREDIT = "Early Locomotion project (visual development, not final art)."


# ------------------------------------------------------------------ slide 1
def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = I(MARGIN), I(1.7), I(6.7), I(2.0)
    t.name = "TITLE"
    tf = t.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = 0.95
    r = p.add_run()
    r.text = "The Monorail Transition"
    _font(r, 54, TEXT, True)
    sh = s.shapes
    # reading order: title, visual (picture and credit), then rule and caption
    picture(sh, "R02_track_hero_hd_v004.jpg", 7.75, 0.45, 4.98, 6.2, R02_ALT, focus=(0.43, 0.5), radius=0.08,
            name="PICTURE R02")
    credit(sh, 7.75, 6.74, 4.98, R02_CREDIT)
    rect(sh, MARGIN, 3.9, 2.4, 0.08, ORANGE, None, name="TITLE RULE")
    label(sh, MARGIN, 4.2, 6.4, 1.6,
          "What happens when the game's electric trains arrive, and how the mod plans to handle it.",
          20, TEXT, align="l", anchor="t", name="CAPTION", margins=(0, 0, 0, 0))


# ------------------------------------------------------------------ slide 2
def slide2(prs):
    s = new_slide(prs, "How to read this guide", None)
    sh = s.shapes
    rows = [("Agreed", "Decided."), ("Proposed", "Suggested, not confirmed."),
            ("Provisional", "A test number that may change."), ("Undecided", "Still an open question."),
            ("Background", "A real-world fact.")]
    g = group(sh, "LEGEND", "Legend: five tag chips, each with its plain meaning.")
    for i, (kind, meaning) in enumerate(rows):
        y = 1.6 + i * 0.9
        rect(g.shapes, MARGIN, y - 0.1, 7.6, 0.6, PANEL, None, radius=0.08)
        chip(g.shapes, kind, 0.8, y, 1.7, 0.4, 14, name="LEGEND CHIP")
        label(g.shapes, 2.8, y - 0.1, 5.4, 0.6, meaning, 20, TEXT, align="l", anchor="m")
    # decorative sketch: a slide with a tag in its top-right corner
    sk = group(sh, "SKETCH", "Sketch of a slide with a tag in its top-right corner.")
    rect(sk.shapes, 9.0, 1.7, 3.7, 2.4, PANEL, LINE, 1.5, radius=0.08)
    rect(sk.shapes, 9.25, 1.95, 1.4, 0.16, LINE, None, radius=0.04)
    rect(sk.shapes, 9.25, 2.4, 2.6, 0.12, LINE, None, radius=0.03)
    rect(sk.shapes, 9.25, 2.65, 2.2, 0.12, LINE, None, radius=0.03)
    rect(sk.shapes, 9.25, 3.0, 3.2, 0.7, PANEL_ALT, None, radius=0.05)
    rect(sk.shapes, 11.55, 1.9, 0.95, 0.26, GREEN, None, radius=0.13)
    finish(s, "Every idea carries a tag in the top-right corner.", 2)


# ------------------------------------------------------------------ slide 3
def slide3(prs):
    s = new_slide(prs, "Two kinds of railway", "Agreed")
    sh = s.shapes
    ground = "#5B6068"
    # --- steam cross-section (left)
    g = group(sh, "DIAGRAM steam track", "End-on view of conventional track: two steel rails fixed to a sleeper, which rests on ballast.")
    gs = g.shapes
    rect(gs, 0.6, 1.45, 5.9, 4.7, PANEL, None, radius=0.08)
    cx = 3.55
    poly(gs, [(cx - 2.1, 4.0), (cx + 2.1, 4.0), (cx + 2.75, 4.85), (cx - 2.75, 4.85)], ground, None)
    rect(gs, cx - 1.9, 3.68, 3.8, 0.32, SLEEPER, None)
    for rx in (cx - 1.2, cx + 1.2):
        rect(gs, rx - 0.2, 2.6, 0.4, 0.22, STEEL, None, radius=0.06)  # rail head
        rect(gs, rx - 0.07, 2.82, 0.14, 0.78, STEEL, None)  # web
        rect(gs, rx - 0.34, 3.6, 0.68, 0.1, STEEL, None)  # foot
    label(gs, 0.8, 5.0, 5.5, 0.9, "Steam: two rails", 24, ORANGE, True, anchor="m")
    # --- monorail cross-section (right)
    g = group(sh, "DIAGRAM monorail", "End-on view of an electric carriage wrapped around a single raised beam on a pillar.")
    gm = g.shapes
    rect(gm, 6.83, 1.45, 5.9, 4.7, PANEL, None, radius=0.08)
    cx = 9.78
    rect(gm, cx - 2.6, 4.75, 5.2, 0.1, ground, None)
    rect(gm, cx - 0.5, 4.65, 1.0, 0.12, ground, None)
    rect(gm, cx - 0.17, 3.2, 0.34, 1.5, STEEL, None)
    rect(gm, cx - 0.42, 2.8, 0.84, 0.42, LIGHT, None)
    poly(gm, [(cx - 1.3, 1.85), (cx + 1.3, 1.85), (cx + 1.3, 3.4), (cx + 0.58, 3.4), (cx + 0.58, 2.65),
              (cx - 0.58, 2.65), (cx - 0.58, 3.4), (cx - 1.3, 3.4)], ELEC, None)
    for i in range(3):
        rect(gm, cx - 1.1 + i * 0.75, 2.0, 0.6, 0.36, BG, None, radius=0.05)
    label(gm, 7.03, 4.92, 5.5, 1.1,
          [[("Monorail: one beam", {"size": 24, "bold": True, "color": ELEC})],
           [("(the game's electric trains)", {"size": 18, "color": TEXT})]], anchor="m")
    finish(s, "Steam runs on two rails. The game's later trains hug a single beam.", 3)


# ------------------------------------------------------------------ slide 4
def slide4(prs):
    s = new_slide(prs, "They cannot swap", "Agreed")
    sh = s.shapes
    ground = "#5B6068"
    # steam wheel on a rail
    g = group(sh, "DIAGRAM steam wheel", "A steam wheel resting on a rail, with a green tick: it works on its own track.")
    gs = g.shapes
    rect(gs, 0.6, 1.45, 4.8, 4.65, PANEL, None, radius=0.08)
    cx, cy, r = 3.0, 3.05, 1.05
    wheel(gs, cx, cy, r, ORANGE, ORANGE, lw=9, spokes=8)
    rect(gs, 0.95, cy + r, 4.1, 0.14, STEEL, None)
    for sx in (1.15, 2.15, 3.15, 4.15):
        rect(gs, sx, cy + r + 0.14, 0.4, 0.3, SLEEPER, None)
    rect(gs, 0.9, cy + r + 0.44, 4.2, 0.1, ground, None)
    tick(gs, 4.95, 1.95, 0.55)
    label(gs, 0.8, 5.0, 4.4, 0.8, "Steam wheel", 20, TEXT, True, anchor="m")
    # clamp wheel on a beam
    g = group(sh, "DIAGRAM clamp wheel", "Clamp wheels gripping a beam from above and both sides, with a green tick: it works on its own track.")
    gm = g.shapes
    rect(gm, 7.93, 1.45, 4.8, 4.65, PANEL, None, radius=0.08)
    cx = 10.33
    rect(gm, cx - 0.2, 3.9, 0.4, 0.9, STEEL, None)
    rect(gm, cx - 0.7, 4.75, 1.4, 0.1, ground, None)
    rect(gm, cx - 0.6, 2.6, 1.2, 1.3, LIGHT, None)
    poly(gm, [(cx - 1.18, 3.95), (cx - 1.18, 2.0), (cx + 1.18, 2.0), (cx + 1.18, 3.95)], None, ELEC, 4, closed=False)
    rect(gm, cx - 0.76, 2.2, 1.52, 0.34, ELEC, None, radius=0.1)
    rect(gm, cx - 1.0, 2.62, 0.34, 1.1, ELEC, None, radius=0.1)
    rect(gm, cx + 0.66, 2.62, 0.34, 1.1, ELEC, None, radius=0.1)
    tick(gm, 12.25, 1.95, 0.55)
    label(gm, 8.13, 5.0, 4.4, 0.8, "Clamp wheel", 20, TEXT, True, anchor="m")
    # crossed double arrow
    g = group(sh, "DIAGRAM cannot swap", "A double arrow between the two wheels, crossed out in red: they cannot swap.")
    gx = g.shapes
    darrow(gx, 5.6, 3.2, 7.73, 3.2, LIGHT, 3.5)
    cross(gx, 6.665, 3.2, 0.8, RED, 6)
    label(gx, 5.5, 3.85, 2.33, 0.5, "Cannot swap", 20, TEXT, True, anchor="m")
    finish(s, "Steam wheels sit on rails. Monorail trains grip a beam. Neither can use the other's track.", 4)


# ------------------------------------------------------------------ slide 5
def slide5(prs):
    s = new_slide(prs, "Why electric wins in the end", "Agreed")
    sh = s.shapes
    g = group(sh, "DIAGRAM paired bars",
              "Paired bars for steam and electric. Steam is longer for cheap to start; electric is longer for speed, "
              "throughput, traffic density and convenience. Shape only, not measured.")
    gs = g.shapes
    data = [("Cheap to start", 85, 45), ("Speed", 40, 90), ("Throughput", 35, 90), ("Traffic density", 35, 90),
            ("Convenience", 45, 90)]
    tx, tw = 3.9, 6.0
    for i, (name, a, b) in enumerate(data):
        y = 1.6 + i * 0.9
        label(gs, MARGIN, y, 3.2, 0.74, name, 20, TEXT, True, align="l", anchor="m")
        rect(gs, tx, y, tw, 0.34, PANEL_ALT, None)
        rect(gs, tx, y, tw * a / 100, 0.34, ORANGE, None)
        rect(gs, tx, y + 0.4, tw, 0.34, PANEL_ALT, None)
        rect(gs, tx, y + 0.4, tw * b / 100, 0.34, ELEC, None)
    # legend
    rect(gs, 10.45, 1.7, 2.28, 1.3, PANEL, None, radius=0.08)
    rect(gs, 10.65, 1.9, 0.3, 0.3, ORANGE, None)
    label(gs, 11.05, 1.85, 1.6, 0.4, "Steam", 18, TEXT, align="l")
    rect(gs, 10.65, 2.45, 0.3, 0.3, ELEC, None)
    label(gs, 11.05, 2.4, 1.6, 0.4, "Electric", 18, TEXT, align="l")
    label(gs, 10.45, 3.25, 2.28, 0.8, "Shape only. Not measured.", 16, MUTED, align="l", anchor="t")
    finish(s, "Electric rail wins on speed, throughput, traffic and convenience once a factory matures.", 5)


# ------------------------------------------------------------------ slide 6
def slide6(prs):
    s = new_slide(prs, "Upgrade in place", "Agreed")
    sh = s.shapes
    pw, ph, py = 3.7, 4.4, 1.55
    xs = [0.6, 4.82, 9.03]
    labels = ["Choose the electric track", "Point at your steam track", "It is replaced"]
    alts = ["Panel 1: a build menu card showing the electric beam, with a mouse pointer choosing it.",
            "Panel 2: the pointer aims at the steam track, which is highlighted.",
            "Panel 3: a blue beam now sits where the rails were."]
    gy = py + 2.55  # ground line
    for i, x in enumerate(xs):
        g = group(sh, "PANEL %d" % (i + 1), alts[i])
        gs = g.shapes
        rect(gs, x, py, pw, ph, PANEL, None, radius=0.08)
        num_badge(gs, x + 0.4, py + 0.4, i + 1)
        label(gs, x + 0.15, py + 3.1, pw - 0.3, 1.1, labels[i], 20, TEXT, True, anchor="m")
        if i == 0:
            rect(gs, x + 0.7, py + 0.95, 2.3, 1.4, PANEL_ALT, LINE, 1.5, radius=0.08)
            rect(gs, x + 0.95, py + 1.3, 1.8, 0.4, ELEC, None, radius=0.1)  # beam icon
            rect(gs, x + 1.7, py + 1.7, 0.3, 0.4, STEEL, None)
            rect(gs, x + 1.25, py + 2.05, 1.2, 0.08, ELEC, None)
            cursor(gs, x + 2.35, py + 1.7, 1.0)
        elif i == 1:
            rect(gs, x + 0.3, py + 1.15, 3.1, 1.25, ELEC, WATER, 2.5, radius=0.1, dash=MSO_LINE.DASH, alpha=0.18)
            rails_h(gs, x + 0.45, py + 1.55, 2.8, ORANGE, gap=0.34, step=0.4, rw=0.07)
            cursor(gs, x + 2.4, py + 1.95, 1.0)
        else:
            for px in (x + 0.8, x + 1.85, x + 2.9):
                rect(gs, px - 0.1, py + 1.75, 0.2, 0.7, STEEL, None)
            rect(gs, x + 0.35, py + 1.4, 3.0, 0.42, ELEC, None, radius=0.1)
            rect(gs, x + 0.3, py + 2.43, 3.1, 0.07, "#5B6068", None)
        if i < 2:
            arrow(sh, x + pw + 0.07, py + ph / 2, xs[i + 1] - 0.07, py + ph / 2, ORANGE, 3)
    finish(s, "Point the upgrade at your old track. It swaps in place, like upgrading a belt.", 6)


# ------------------------------------------------------------------ slide 7
def _path(x0, x1, yc, amp, n=80):
    pts = []
    L = x1 - x0
    for i in range(n + 1):
        x = x0 + L * i / n
        pts.append((x, yc + amp * math.sin(2 * math.pi * (x - x0) / L)))
    return pts


def _offset(pts, d):
    out = []
    for i, (x, y) in enumerate(pts):
        a = pts[max(i - 1, 0)]
        b = pts[min(i + 1, len(pts) - 1)]
        dx, dy = b[0] - a[0], b[1] - a[1]
        n = math.hypot(dx, dy)
        out.append((x - dy / n * d, y + dx / n * d))
    return out


def slide7(prs):
    s = new_slide(prs, "Your route stays", "Agreed")
    sh = s.shapes
    x0, x1, amp = 1.7, 11.6, 0.35
    top_c, bot_c = 2.4, 4.15
    # top copy: steam track
    g = group(sh, "DIAGRAM steam route", "A curved route between two stations drawn as orange steam track on sleepers.")
    gs = g.shapes
    pts = _path(x0, x1, top_c, amp)
    for i in range(0, len(pts), 3):
        (px, py), (qx, qy) = _offset(pts, -0.19)[i], _offset(pts, 0.19)[i]
        line(gs, px, py, qx, qy, SLEEPER, 3.5)
    poly(gs, _offset(pts, -0.09), None, ORANGE, 3, closed=False)
    poly(gs, _offset(pts, 0.09), None, ORANGE, 3, closed=False)
    for px, py in (pts[0], pts[-1]):
        oval(gs, px - 0.24, py - 0.24, 0.48, 0.48, LIGHT, DARK, 3)
    label(gs, 8.7, 1.4, 2.4, 0.4, "Steam track", 18, ORANGE, True, align="l")
    # bottom copy: monorail beam
    g = group(sh, "DIAGRAM electric route", "The same curved route between the same two stations, now drawn as a blue beam on supports.")
    gs = g.shapes
    pts2 = _path(x0, x1, bot_c, amp)
    for i in range(2, len(pts2) - 1, 8):
        px, py = pts2[i]
        rect(gs, px - 0.07, py, 0.14, 0.5, STEEL, None)
    poly(gs, pts2, None, ELEC, 13, closed=False)
    for px, py in (pts2[0], pts2[-1]):
        oval(gs, px - 0.24, py - 0.24, 0.48, 0.48, LIGHT, DARK, 3)
    label(gs, 8.7, 4.72, 2.6, 0.4, "Monorail beam", 18, ELEC, True, align="l")
    # same-path markers
    g = group(sh, "DIAGRAM same path", "Dashed lines join the two copies: the path is the same.")
    gs = g.shapes
    for x in (x0, x1):
        line(gs, x, top_c + 0.3, x, bot_c - 0.3, MUTED, 2, MSO_LINE.DASH)
    for x in (5.0, 8.3):
        ty = top_c + amp * math.sin(2 * math.pi * (x - x0) / (x1 - x0))
        arrow(gs, x, ty + 0.4, x, ty + (bot_c - top_c) - 0.4, MUTED, 2.25, MSO_LINE.DASH)
    label(gs, 5.5, 3.1, 2.3, 0.45, "Same path", 20, TEXT, True)
    # note with inline chip
    g = group(sh, "NOTE proposed", "Note: supports and junction shapes kept where possible. Proposed.")
    gs = g.shapes
    rect(gs, MARGIN, 5.35, 12.13, 0.65, PANEL, None, radius=0.08)
    chip(gs, "Proposed", 0.85, 5.5, None, 0.34, 14, name="ICHIP")
    label(gs, 0.85 + chip_w("Proposed") + 0.25, 5.35, 9.5, 0.65, "Supports and junction shapes kept where possible",
          20, TEXT, align="l")
    finish(s, "The route you planned is kept. You do not start again.", 7)


# ------------------------------------------------------------------ slide 8
def slide8(prs):
    s = new_slide(prs, "A warning first", "Proposed")
    sh = s.shapes
    py, ph = 1.5, 4.4
    g = group(sh, "DIAGRAM hologram", "A mock build preview: a ghost blue beam hologram laid over orange steam track.")
    gs = g.shapes
    rect(gs, MARGIN, py, 6.9, ph, PANEL, None, radius=0.08)
    gy = 5.35
    rect(gs, 0.9, gy, 6.3, 0.08, "#5B6068", None)
    for px in (1.5, 3.65, 5.8):
        rect(gs, px, 4.25, 0.24, gy - 4.25, None, ELEC, 2, dash=MSO_LINE.DASH)
    rails_h(gs, 1.0, 3.62, 6.1, ORANGE, gap=0.34, step=0.4, rw=0.08)
    rect(gs, 0.9, 3.3, 6.3, 0.95, ELEC, ELEC, 2.5, radius=0.1, dash=MSO_LINE.DASH, alpha=0.35)
    label(gs, 0.9, 1.8, 2.6, 0.45, "Hologram preview", 18, ELEC, True, align="l")
    arrow(gs, 2.2, 2.3, 2.2, 3.2, ELEC, 2.25)
    # dialog
    g = group(sh, "DIAGRAM warning dialog",
              "A warning dialog: Steam trains cannot use this section. Two buttons, Confirm and Cancel.")
    gs = g.shapes
    rect(gs, 7.85, py, 4.88, ph, PANEL_ALT, None, radius=0.08)
    rect(gs, 8.15, 1.85, 4.28, 2.3, AMBER, None, radius=0.08)
    warning_triangle(gs, 8.4, 2.68, 0.72, 0.64, fill=DARK, edge=DARK, bang=AMBER)
    label(gs, 9.3, 2.0, 3.0, 2.0, "Steam trains cannot use this section.", 20, DARK, True, align="l", anchor="m")
    rect(gs, 8.15, 4.75, 2.0, 0.7, ORANGE, None, radius=0.1, text="Confirm", size=20, color=DARK, bold=True)
    rect(gs, 10.43, 4.75, 2.0, 0.7, None, TEXT, 2, radius=0.1, text="Cancel", size=20, color=TEXT, bold=True)
    finish(s, "Before you confirm, the game warns you. Steam trains cannot use the new track.", 8)


SLIDES = [slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8]
