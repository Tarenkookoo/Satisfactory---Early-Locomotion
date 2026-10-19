"""Slides 16 to 22 of Deck 4."""
import math

from pptx.enum.dml import MSO_LINE

from lib import *  # noqa: F401,F403
from lib import _font  # noqa: F401

R01_ALT = ("Early test model of a tank engine (a yellow 'visual probe') on a short length of track. "
           "Not final art.")


# ------------------------------------------------------------------ slide 16
def slide16(prs):
    s = new_slide(prs, "What could go wrong?", "Undecided")
    sh = s.shapes
    cw, gap, y0, ch = 2.905, 0.17, 1.5, 3.85
    titles = ["Train on the track", "Steam train stranded", "Half-upgraded line", "Junction boundary"]
    alts = ["A steam train on a section of track where a blue beam preview overlaps it. Marked with a question mark.",
            "A steam train sitting on a section that is already a blue beam. Marked with a question mark.",
            "A line with orange steam track on the left and a blue beam on the right. Marked with a question mark.",
            "A junction where the line splits, with a dashed boundary and blue beams beyond it. Marked with a question mark."]
    for i, title in enumerate(titles):
        x = MARGIN + i * (cw + gap)
        g = group(sh, "SCENARIO %d" % (i + 1), alts[i])
        gs = g.shapes
        rect(gs, x, y0, cw, ch, PANEL, None, radius=0.08)
        rect(gs, x + 0.12, y0 + 0.65, cw - 0.24, 2.2, BG, None, radius=0.06)
        qbadge(gs, x + cw - 0.45, y0 + 0.38, 0.5)
        label(gs, x + 0.1, y0 + 2.95, cw - 0.2, 0.75, title, 20, TEXT, True)
        dx0, wd, yc = x + 0.3, cw - 0.6, y0 + 1.75
        if i == 0:  # train under a hologram
            rails_h(gs, dx0, yc - 0.17, wd, ORANGE, gap=0.34, step=0.3, rw=0.06)
            rect(gs, dx0 + 0.2, yc - 0.24, 0.9, 0.48, ORANGE, None, radius=0.08)
            oval(gs, dx0 + 0.95, yc - 0.09, 0.18, 0.18, COAL, None)
            rect(gs, dx0 + 1.2, yc - 0.2, 0.9, 0.4, WAGON, None, radius=0.05)
            rect(gs, dx0 - 0.05, yc - 0.45, wd + 0.1, 0.9, ELEC, ELEC, 2.5, radius=0.1, dash=MSO_LINE.DASH,
                 alpha=0.3)
        elif i == 1:  # train stranded on a beam
            rails_h(gs, dx0, yc - 0.17, 0.85, ORANGE, gap=0.34, step=0.3, rw=0.06)
            rect(gs, dx0 + 0.95, yc - 0.13, wd - 0.95, 0.26, ELEC, None, radius=0.06)
            rect(gs, dx0 + 1.05, yc - 0.3, 1.1, 0.5, ORANGE, None, radius=0.08)
            oval(gs, dx0 + 1.9, yc - 0.14, 0.18, 0.18, COAL, None)
        elif i == 2:  # half upgraded
            rails_h(gs, dx0, yc - 0.17, 1.1, ORANGE, gap=0.34, step=0.3, rw=0.06)
            rect(gs, dx0 + 1.2, yc - 0.13, wd - 1.2, 0.26, ELEC, None, radius=0.06)
            line(gs, dx0 + 1.15, yc - 0.6, dx0 + 1.15, yc + 0.6, MUTED, 2.5, MSO_LINE.DASH)
        else:  # junction boundary
            jx = dx0 + 0.95
            poly(gs, [(dx0, yc), (jx, yc)], None, ORANGE, 7, closed=False)
            poly(gs, [(jx, yc), (jx + 0.5, yc - 0.5), (dx0 + wd, yc - 0.5)], None, ELEC, 7, closed=False)
            poly(gs, [(jx, yc), (jx + 0.5, yc + 0.5), (dx0 + wd, yc + 0.5)], None, ELEC, 7, closed=False)
            line(gs, jx + 0.18, yc - 0.85, jx + 0.18, yc + 0.85, MUTED, 2.5, MSO_LINE.DASH)
    label(sh, 10.7, 5.55, 2.03, 0.3, "OQ12", 14, MUTED, align="r", name="FOOTER")
    finish(s, "Some cases still need answers: busy track, stranded trains, half-converted lines.", 16)


# ------------------------------------------------------------------ slide 17
def slide17(prs):
    s = new_slide(prs, "Costs, refunds and undo", "Undecided")
    sh = s.shapes
    cw, gap, y0, ch = 3.877, 0.25, 1.55, 3.9
    qs = ["Cost and refund?", "Undo or recovery?", "Longest piece?"]
    alts = ["Question card: cost and refund. Two small stacks of plates with arrows both ways.",
            "Question card: undo or recovery. A circular undo arrow.",
            "Question card: longest piece. A short piece of track with a measuring line."]
    for i, q in enumerate(qs):
        x = MARGIN + i * (cw + gap)
        cx = x + cw / 2
        g = group(sh, "QUESTION %d" % (i + 1), alts[i])
        gs = g.shapes
        rect(gs, x, y0, cw, ch, PANEL, None, radius=0.08)
        label(gs, x + 0.15, y0 + 2.7, cw - 0.3, 0.9, q, 26, TEXT, True)
        yc = y0 + 1.35
        if i == 0:
            for k in range(3):
                rect(gs, cx - 1.65, yc + 0.45 - k * 0.32, 1.1, 0.26, STEEL, None, radius=0.05)
                rect(gs, cx + 0.55, yc + 0.45 - k * 0.32, 1.1, 0.26, STEEL, None, radius=0.05)
            arrow(gs, cx - 0.4, yc - 0.1, cx + 0.4, yc - 0.1, ORANGE, 4)
            arrow(gs, cx + 0.4, yc + 0.4, cx - 0.4, yc + 0.4, ORANGE, 4)
        elif i == 1:
            arc(gs, cx, yc + 0.1, 0.8, 25, 300, ORANGE, 6, True)
        else:
            rails_h(gs, cx - 1.6, yc - 0.3, 3.2, STEEL, gap=0.36, step=0.4, rw=0.08)
            darrow(gs, cx - 1.6, yc + 0.6, cx + 1.6, yc + 0.6, ORANGE, 3.5)
            line(gs, cx - 1.6, yc + 0.42, cx - 1.6, yc + 0.78, ORANGE, 3.5)
            line(gs, cx + 1.6, yc + 0.42, cx + 1.6, yc + 0.78, ORANGE, 3.5)
    label(sh, 10.7, 5.6, 2.03, 0.3, "OQ12", 14, MUTED, align="r", name="FOOTER")
    finish(s, "Upgrade costs, refunds and an undo are not decided. Nor is the longest piece the game accepts.", 17)


# ------------------------------------------------------------------ slide 18
def slide18(prs):
    s = new_slide(prs, "The road to electric", "Provisional")
    sh = s.shapes
    cw, y0, ch = 2.75, 1.85, 2.1
    gap = (12.13 - 4 * cw) / 3
    tiers = [("Tier 3:", "Early Steam Railways"), ("Tier 3 to 4:", "Railway Expansion"),
             ("Tier 4 to 5:", "Heavy Steam"), ("Tier 6:", "Monorail")]
    g = group(sh, "DIAGRAM tier timeline",
              "A timeline of four tiers: Early Steam Railways, Railway Expansion, Heavy Steam, then the Monorail in blue.")
    gs = g.shapes
    for i, (t, name) in enumerate(tiers):
        x = MARGIN + i * (cw + gap)
        electric = i == 3
        accent = ELEC if electric else ORANGE
        rect(gs, x, y0, cw, ch, PANEL, ELEC if electric else None, 2.5, radius=0.08)
        rect(gs, x, y0, cw, 0.14, accent, None, radius=0.05)
        label(gs, x + 0.15, y0 + 0.35, cw - 0.3, 0.55, t, 24, accent, True, align="l")
        label(gs, x + 0.15, y0 + 0.95, cw - 0.3, 1.3, name, 20, TEXT, True, align="l", anchor="t")
        if i < 3:
            arrow(gs, x + cw + 0.04, y0 + ch / 2, x + cw + gap - 0.04, y0 + ch / 2, ORANGE, 3)
        oval(gs, x + cw / 2 - 0.13, 4.3, 0.26, 0.26, accent, None)
    line(gs, MARGIN + cw / 2, 4.43, MARGIN + 3 * (cw + gap) + cw / 2 - 0.15, 4.43, LINE, 3)
    note_strip(sh, "Undecided", "Tiers and prerequisites: provisional and undecided (OQ01)", 5.1)
    finish(s, "Steam starts early. The monorail arrives in a later tier.", 18)


# ------------------------------------------------------------------ slide 19
def _belt_icon(sh, x, y):
    rect(sh, x, y + 0.2, 1.7, 0.2, "#59616C", None, radius=0.1)
    for k in range(5):
        cx = x + 0.2 + k * 0.32
        poly(sh, [(cx, y + 0.23), (cx + 0.1, y + 0.3), (cx, y + 0.37)], None, LIGHT, 2, closed=False)
    oval(sh, x - 0.05, y + 0.15, 0.3, 0.3, STEEL, None)
    oval(sh, x + 1.45, y + 0.15, 0.3, 0.3, STEEL, None)
    rect(sh, x + 0.6, y - 0.12, 0.34, 0.3, LIGHT, DARK, 1.5)
    rect(sh, x + 1.05, y - 0.12, 0.34, 0.3, LIGHT, DARK, 1.5)


def _tractor_icon(sh, x, y):
    rect(sh, x + 0.1, y + 0.15, 1.1, 0.3, AMBER, None, radius=0.06)
    rect(sh, x + 0.75, y - 0.12, 0.5, 0.35, AMBER, None, radius=0.05)
    rect(sh, x + 0.85, y - 0.05, 0.3, 0.18, BG, None)
    line(sh, x + 0.3, y + 0.15, x + 0.3, y - 0.1, STEEL, 3)
    oval(sh, x + 0.85, y + 0.2, 0.62, 0.62, DARK, STEEL, 3)
    oval(sh, x + 0.05, y + 0.37, 0.4, 0.4, DARK, STEEL, 3)


def slide19(prs):
    s = new_slide(prs, "Every transport has a job", "Agreed")
    sh = s.shapes
    lanes = [("Belts:", " permanent short and medium links", TEXT, "belt"),
             ("Tractors:", " flexible, little infrastructure", AMBER, "tractor"),
             ("Steam:", " early long-distance bulk", ORANGE, "steam"),
             ("Electric:", " mature speed and throughput", ELEC, "electric")]
    alts = ["Lane: belts, with a conveyor belt icon.", "Lane: tractors, with a tractor icon.",
            "Lane: steam, with a small steam engine icon.", "Lane: electric, with a blue monorail carriage on a beam."]
    for i, (term, rest, col, icon) in enumerate(lanes):
        y = 1.5 + i * 1.14
        g = group(sh, "LANE %d" % (i + 1), alts[i])
        gs = g.shapes
        rect(gs, MARGIN, y, 12.13, 1.02, PANEL, None, radius=0.08)
        rect(gs, MARGIN, y + 0.08, 0.08, 0.86, col, None)
        ix, iy = 1.0, y + 0.04
        if icon == "belt":
            _belt_icon(gs, ix, iy + 0.28)
        elif icon == "tractor":
            _tractor_icon(gs, ix, iy + 0.16)
        elif icon == "steam":
            loco_side(gs, ix, iy, 1.7, ORANGE, "right")
        else:
            rect(gs, ix - 0.05, iy + 0.78, 1.8, 0.12, LIGHT, None)
            carriage_side(gs, ix, iy + 0.18, 1.7, ELEC)
        label(gs, 3.2, y, 9.3, 1.02, [[(term, {"size": 24, "bold": True, "color": col}),
                                       (rest, {"size": 20, "color": TEXT})]], 20, TEXT, align="l", anchor="m")
    finish(s, "Belts, tractors, steam and electric rail each have a job. Electric leads when the factory matures.", 19)


# ------------------------------------------------------------------ slide 20
def slide20(prs):
    s = new_slide(prs, "Where this stands", "Provisional")
    sh = s.shapes
    picture(sh, "R01_pannier_visual_probe_v004.jpg", MARGIN, 1.5, 2.4, 1.35, R01_ALT, radius=0.08,
            name="PICTURE R01")
    credit(sh, MARGIN, 2.88, 2.4, "Early test model, not final art.")
    note_strip_at(sh, "Provisional", "Plans can change.", 3.4, 1.85, 9.33, 0.65)
    steps = [("Test engine on the game's rails:", "Under way"), ("Upgrade in place:", "Planned"),
             ("Station replacement:", "Planned"), ("Mixed networks:", "Planned"), ("Interchange:", "Planned")]
    cw = 2.3
    gap = (12.13 - 5 * cw) / 4
    y0, ch = 3.4, 2.55
    g = group(sh, "ROADMAP", "Five steps. The test engine on the game's rails is under way. The other four steps are planned.")
    gs = g.shapes
    for i, (t, st) in enumerate(steps):
        x = MARGIN + i * (cw + gap)
        under = i == 0
        rect(gs, x, y0, cw, ch, PANEL, ORANGE if under else None, 2.5 if under else 0, radius=0.08)
        num_badge(gs, x + 0.4, y0 + 0.4, i + 1)
        label(gs, x + 0.12, y0 + 0.75, cw - 0.24, 1.0, t.rstrip(":"), 18, TEXT, True, align="l", anchor="t")
        if under:
            rect(gs, x + 0.15, y0 + 1.85, cw - 0.3, 0.45, ORANGE, None, radius=0.22, text="Under way", size=16,
                 color=DARK, bold=True)
        else:
            rect(gs, x + 0.15, y0 + 1.85, cw - 0.3, 0.45, None, MUTED, 2, radius=0.22, text="Planned", size=16,
                 color=TEXT, bold=True)
    finish(s, "Conversion is planned. It is not built yet.", 20)


# ------------------------------------------------------------------ slide 21
def slide21(prs):
    s = new_slide(prs, "Still undecided", "Undecided")
    sh = s.shapes
    # (question, open-question code): the code sits in the card's Undecided chip to keep the slide under 45 words
    qs = [("Exact copies of shapes", "OQ12"), ("Costs and refunds", "OQ12"), ("Undo and recovery", "OQ12"),
          ("Occupied and stranded trains", "OQ12"), ("Station and timetable carry-over", "OQ12"),
          ("Signal upgrades", "OQ03"), ("Transfer buildings in or out of Core?", "OQ13"),
          ("Which tier unlocks the upgrade?", "OQ01")]
    cw, ch = 2.9, 2.1
    gx = (12.13 - 4 * cw) / 3
    g = group(sh, "QUESTION GRID", "Eight open questions, each marked Undecided with its open-question code.")
    gs = g.shapes
    for i, (q, code) in enumerate(qs):
        r, c = divmod(i, 4)
        x = MARGIN + c * (cw + gx)
        y = 1.5 + r * (ch + 0.25)
        rect(gs, x, y, cw, ch, PANEL, None, radius=0.08)
        chip(gs, "Undecided", x + 0.15, y + 0.17, None, 0.34, 14, name="ICHIP", text="UNDECIDED \u00b7 " + code)
        label(gs, x + 0.15, y + 0.65, cw - 0.3, 1.38, q, 20, TEXT, align="l", anchor="t")
    finish(s, "These questions are open. Nothing here has been chosen yet.", 21)


# ------------------------------------------------------------------ slide 22
def slide22(prs):
    s = new_slide(prs, "Image credits", None)
    sh = s.shapes
    rows = [["Slot", "Slide", "Credit"],
            ["R01", "20", "Early Locomotion project (test model, not final art)."],
            ["R02", "1", "Early Locomotion project (visual development, not final art)."]]
    table(sh, MARGIN, 1.6, [1.3, 1.3, 9.53], rows, size=11, row_h=0.5,
          alt="Image credits table: R01 on slide 20 and R02 on slide 1, both Early Locomotion project renders, not final art.")
    finish(s, "This deck uses no photographs.", 22)


SLIDES = [slide16, slide17, slide18, slide19, slide20, slide21, slide22]
