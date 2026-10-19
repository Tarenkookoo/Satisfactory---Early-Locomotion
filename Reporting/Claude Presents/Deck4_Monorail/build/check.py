"""Audit the deck: word counts, smallest font, text-overflow estimate, alt text, notes, transitions.

usage: python3 check.py [first] [last]
Word count = caption + diagram labels + inline text. Not counted: title, status chips (top-right and
inline plain ones), number badges, slide number, 10 pt credit lines, and open-question codes (OQ12 and so on).
Both counts are printed ("words" excludes OQ codes, "+OQ" includes them).
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from lib import wrap_count  # noqa: E402

PPTX = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                                    "The_Monorail_Transition_Explained.pptx"))
SKIP_NAMES = ("TITLE", "CHIP", "NUM", "BADGE", "ICHIP", "LEGEND CHIP", "TICK")
EXEMPT = {2}
OQ = re.compile(r"^\(?OQ\d+\)?[.,]?$")


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)


def words_of(text):
    return [w for w in re.split(r"\s+", text.strip()) if re.search(r"[A-Za-z0-9]", w)]


def main():
    prs = Presentation(PPTX)
    n = len(prs.slides)
    first = int(sys.argv[1]) if len(sys.argv) > 1 else 1
    last = int(sys.argv[2]) if len(sys.argv) > 2 else n
    exempt = EXEMPT | {n}
    problems = 0
    print("slide  words  +OQ  minpt(body)  minpt(all)  notes  trans/anim  order")
    for i, slide in enumerate(prs.slides, 1):
        if i < first or i > last:
            continue
        w_count = 0
        w_oq = 0
        sizes_body, sizes_all = [], []
        warns = []
        for sh in walk(slide.shapes):
            name = sh.name or ""
            el = sh._element
            # alt text
            if sh.shape_type in (6, 13, 19):  # group, picture, table
                c = el.xpath("./p:nvGrpSpPr/p:cNvPr | ./p:nvPicPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr")[0]
                if not (c.get("descr") or "").strip():
                    warns.append("no alt text on %s" % name)
            tfs = []
            if sh.has_text_frame and sh.text_frame.text.strip():
                tfs.append((sh.text_frame, sh.width, sh.height))
            if sh.shape_type == 19:
                tbl = sh.table
                for r in tbl.rows:
                    for c in r.cells:
                        tfs.append((c.text_frame, None, None))
            for tf, wd, ht in tfs:
                txt = tf.text
                is_credit = name.startswith("CREDIT")
                is_table = sh.shape_type == 19
                psizes = []
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.size is not None:
                            psizes.append(r.font.size.pt)
                if not psizes:
                    continue
                sizes_all += psizes
                if not is_credit and not is_table:
                    sizes_body += psizes
                if not (name in SKIP_NAMES or name.startswith("CREDIT") or name.startswith("TICK")):
                    ws = words_of(txt)
                    w_oq += len(ws)
                    w_count += len([w for w in ws if not OQ.match(w)])
                # overflow estimate for fixed boxes
                if wd is not None and not is_table:
                    boldrun = any(r.font.bold for p in tf.paragraphs for r in p.runs)
                    size = max(psizes)
                    wi = Emu(wd).inches - 0.12 - (0.12 if sh.shape_type == 1 and "ROUND" in str(sh.auto_shape_type) else 0)
                    lines = sum(wrap_count(p.text, max([r.font.size.pt for r in p.runs if r.font.size] or [size]), wi,
                                           boldrun) for p in tf.paragraphs)
                    need = lines * size * 1.2 / 72.0 + 0.04
                    if need > Emu(ht).inches + 0.04 and name != "TITLE":
                        warns.append("overflow? %s: needs %.2f in, has %.2f in (%r)" % (name, need, Emu(ht).inches, txt[:40]))
        sld = slide._element
        has_trans = bool(sld.xpath(".//p:transition")) or bool(sld.xpath(".//p:timing"))
        order = [s.name for s in list(slide.shapes)[:3]]
        mn_b = min(sizes_body) if sizes_body else 0
        mn_a = min(sizes_all) if sizes_all else 0
        flag = ""
        if i not in exempt and w_count > 45:
            flag = "  OVER 45"
            problems += 1
        elif i not in exempt and w_count > 25:
            flag = "  >25"
        if mn_b and mn_b < 14:
            flag += "  SMALL BODY FONT"
            problems += 1
        print("%5d  %5d  %3d  %11.0f  %10.0f  %5s  %10s  %s%s" % (
            i, w_count, w_oq, mn_b, mn_a, slide.has_notes_slide, has_trans, order, flag))
        for w in warns:
            print("        WARN", w)
            problems += 1
    print("core props:", prs.core_properties.title, "|", prs.core_properties.author, "|", prs.core_properties.language)
    print("problems:", problems)


if __name__ == "__main__":
    main()
