"""Verify Deck 4 against DECK_BRIEF.md: titles, tags, exact captions, tag counts, contrast,
alt text, notes, transitions/animations, document properties, image set and sizes.

usage: python3 verify.py
"""
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

PPTX = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                                    "The_Monorail_Transition_Explained.pptx"))

# (title, tag, caption) exactly as in DECK_BRIEF.md
BRIEF = {
    1: ("The Monorail Transition", None,
        "What happens when the game's electric trains arrive, and how the mod plans to handle it."),
    2: ("How to read this guide", None, "Every idea carries a tag in the top-right corner."),
    3: ("Two kinds of railway", "Agreed", "Steam runs on two rails. The game's later trains hug a single beam."),
    4: ("They cannot swap", "Agreed",
        "Steam wheels sit on rails. Monorail trains grip a beam. Neither can use the other's track."),
    5: ("Why electric wins in the end", "Agreed",
        "Electric rail wins on speed, throughput, traffic and convenience once a factory matures."),
    6: ("Upgrade in place", "Agreed", "Point the upgrade at your old track. It swaps in place, like upgrading a belt."),
    7: ("Your route stays", "Agreed", "The route you planned is kept. You do not start again."),
    8: ("A warning first", "Proposed",
        "Before you confirm, the game warns you. Steam trains cannot use the new track."),
    9: ("Will every shape fit?", "Undecided", "Not every shape may convert. Exact copies are not promised yet."),
    10: ("Stations change too", "Proposed",
         "A big steam station is replaced by a compact electric one on the same line."),
    11: ("Names and timetables", "Proposed",
         "Your station name and timetable could carry over. Old coal and water gear is removed."),
    12: ("Signals change", "Proposed", "Mechanical signals may give way to electronic ones. The details are not set."),
    13: ("Mixed networks", "Proposed", "Steam branches could live beside an electric main line."),
    14: ("Cargo crosses, trains do not", "Proposed",
         "Freight moves between the two networks at an interchange. Steam trains cannot cross over."),
    15: ("A Transfer Station?", "Undecided", "A special transfer building is only an idea. It is not promised."),
    16: ("What could go wrong?", "Undecided",
         "Some cases still need answers: busy track, stranded trains, half-converted lines."),
    17: ("Costs, refunds and undo", "Undecided",
         "Upgrade costs, refunds and an undo are not decided. Nor is the longest piece the game accepts."),
    18: ("The road to electric", "Provisional", "Steam starts early. The monorail arrives in a later tier."),
    19: ("Every transport has a job", "Agreed",
         "Belts, tractors, steam and electric rail each have a job. Electric leads when the factory matures."),
    20: ("Where this stands", "Provisional", "Conversion is planned. It is not built yet."),
    21: ("Still undecided", "Undecided", "These questions are open. Nothing here has been chosen yet."),
    22: ("Image credits", None, "This deck uses no photographs."),
}
EXPECT_COUNTS = {"Agreed": 6, "Proposed": 6, "Undecided": 5, "Provisional": 2, None: 3}


def lum(hexc):
    h = hexc.lstrip("#")
    c = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    c = [x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4 for x in c]
    return 0.2126 * c[0] + 0.7152 * c[1] + 0.0722 * c[2]


def contrast(a, b):
    la, lb = lum(a), lum(b)
    hi, lo = max(la, lb), min(la, lb)
    return (hi + 0.05) / (lo + 0.05)


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def fill_hex(sh):
    try:
        if sh.fill.type == 1:  # solid
            el = sh._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
            clr = el.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            if clr is not None and clr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha") is None:
                return "#" + clr.get("val")
    except Exception:
        pass
    return None


def main():
    fails = 0
    prs = Presentation(PPTX)
    n = len(prs.slides)
    print("slides:", n, "(expected 22)")
    if n != 22:
        fails += 1
    counts = {}
    for i, slide in enumerate(prs.slides, 1):
        title, tag, caption = BRIEF[i]
        names = {sh.name: sh for sh in slide.shapes}
        t = slide.shapes.title.text_frame.text if slide.shapes.title is not None else None
        cap = names["CAPTION"].text_frame.text if "CAPTION" in names else None
        chip = names["CHIP"].text_frame.text if "CHIP" in names else None
        if t != title:
            print("slide %d TITLE differs: %r vs %r" % (i, t, title))
            fails += 1
        if cap != caption:
            print("slide %d CAPTION differs: %r vs %r" % (i, cap, caption))
            fails += 1
        if (chip.title() if chip else None) != tag:
            print("slide %d CHIP differs: %r vs %r" % (i, chip, tag))
            fails += 1
        counts[tag] = counts.get(tag, 0) + 1
        if slide.has_notes_slide:
            print("slide %d has a notes slide" % i)
            fails += 1
        x = slide._element
        if x.xpath(".//p:transition") or x.xpath(".//p:timing"):
            print("slide %d has transition or animation" % i)
            fails += 1
        # reading order: title first, chip second (if any), caption bar after the visuals
        order = [s.name for s in slide.shapes]
        if order[0] != "TITLE":
            print("slide %d: title is not first" % i)
            fails += 1
        if tag and order[1] != "CHIP":
            print("slide %d: chip is not second" % i)
            fails += 1
        if i not in (1, 2, 22) and "CAPTION BAR" in order and order.index("CAPTION BAR") < len(order) - 4:
            print("slide %d: caption is not near the end of the reading order" % i)
            fails += 1
        # language
        langs = set()
        for sh in walk(slide.shapes):
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        langs.add(r.font.language_id)
        # contrast
        shapes = list(walk(slide.shapes))
        for k, sh in enumerate(shapes):
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            bg = fill_hex(sh)
            if bg is None:
                cx = sh.left + sh.width // 2
                cy = sh.top + sh.height // 2
                bg = "#1E2126"
                for other in shapes[:k]:
                    if other.shape_type == 6 or other is sh:
                        continue
                    f = fill_hex(other)
                    if f and other.left <= cx <= other.left + other.width and other.top <= cy <= other.top + other.height:
                        bg = f
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    try:
                        col = "#" + str(r.font.color.rgb)
                    except Exception:
                        continue
                    cr = contrast(col, bg)
                    if cr < 4.5:
                        print("slide %d low contrast %.2f: %r %s on %s" % (i, cr, r.text[:30], col, bg))
                        fails += 1
        # alt text on pictures and groups
        for sh in shapes:
            if sh.shape_type in (6, 13, 19):
                c = sh._element.xpath("./p:nvGrpSpPr/p:cNvPr | ./p:nvPicPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr")[0]
                if not (c.get("descr") or "").strip():
                    print("slide %d: no alt text on %s" % (i, sh.name))
                    fails += 1
    print("tag counts:", {str(k): v for k, v in counts.items()}, "expected",
          {str(k): v for k, v in EXPECT_COUNTS.items()})
    if counts != EXPECT_COUNTS:
        fails += 1
    cp = prs.core_properties
    print("title/author/language:", cp.title, "|", cp.author, "|", cp.language)
    if cp.title != "The Monorail Transition, Explained" or cp.author != "Claude (Early Locomotion project)" \
            or cp.language != "en-GB":
        fails += 1
    # package checks
    with zipfile.ZipFile(PPTX) as z:
        bad = z.testzip()
        names = z.namelist()
        media = [m for m in names if m.startswith("ppt/media/")]
        print("zip ok:", bad is None, "| parts:", len(names), "| media:", media)
        if any("notesSlide" in m for m in names):
            print("notes parts present")
            fails += 1
        if len(media) != 2:
            print("expected exactly two images (R01, R02)")
            fails += 1
        from lxml import etree
        for m in names:
            if m.endswith(".xml") or m.endswith(".rels"):
                etree.fromstring(z.read(m))
        # language attribute: every run should be en-GB
        langs = set()
        for m in names:
            if m.startswith("ppt/slides/slide") and m.endswith(".xml"):
                for lang in re.findall(rb'lang="([^"]+)"', z.read(m)):
                    langs.add(lang.decode())
        print("run languages in slides:", langs)
        if langs - {"en-GB"}:
            fails += 1
    print("size: %.0f KB" % (os.path.getsize(PPTX) / 1024))
    print("FAILS:", fails)
    return fails


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
