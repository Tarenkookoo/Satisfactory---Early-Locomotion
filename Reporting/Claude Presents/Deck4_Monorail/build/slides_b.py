"""Slides 9 to 15 of Deck 4."""
import math

from pptx.enum.dml import MSO_LINE

from lib import *  # noqa: F401,F403
from lib import _font  # noqa: F401

HATCH_FG = "#6A727D"


# ------------------------------------------------------------------ slide 9
def slide9(prs):
    s = new_slide(prs, "Will every shape fit?", "Undecided")
    sh = s.shapes
    x0, x1 = 1.0, 12.3

    def yfn(x):
        return 3.4 - 1.1 * (x - x0) / (x1 - x0) + 0.38 * math.sin(2 * math.pi * (x - x0) / 5.6)

    g = group(sh, "DIAGRAM sloping track",
              "A curved, sloping piece of track on rising ground, with five markers pointing at the places that need checking.")
    gs = g.shapes
    rect(gs, MARGIN, 1.45, 12.13, 2.55, PANEL, None, radius=0.08)
    pts = path_pts(x0, x1, yfn)
    poly(gs, pts + [(x1, 3.9), (x0, 3.9)], PANEL_ALT, None)
    for i in range(0, len(pts), 2):
        a, b = offset_pts(pts, -0.2)[i], offset_pts(pts, 0.2)[i]
        line(gs, a[0], a[1], b[0], b[1], SLEEPER, 3.5)
    poly(gs, offset_pts(pts, -0.1), None, ORANGE, 3.5, closed=False)
    poly(gs, offset_pts(pts, 0.1), None, ORANGE, 3.5, closed=False)
    terms = ["Curves", "Slopes", "Clearances", "Piece lengths", "Connections"]
    cw, gap = 2.29, 0.17
    for i, term in enumerate(terms):
        x = MARGIN + i * (cw + gap)
        cx = x + cw / 2
        my = yfn(cx)
        line(gs, cx, my + 0.14, cx, 4.3, MUTED, 2, MSO_LINE.DASH)
        oval(gs, cx - 0.13, my - 0.13, 0.26, 0.26, LIGHT, ORANGE, 3)
    g = group(sh, "CARDS checks", "Five check cards, each marked with a question mark: curves, slopes, clearances, piece lengths, connections.")
    gc = g.shapes
    for i, term in enumerate(terms):
        x = MARGIN + i * (cw + gap)
        rect(gc, x, 4.3, cw, 1.45, PANEL, None, radius=0.08)
        label(gc, x + 0.05, 4.42, cw - 0.1, 0.5, term, 20, TEXT, True)
        qbadge(gc, x + cw / 2, 5.28, 0.5)
    label(sh, 10.7, 5.85, 2.03, 0.3, "OQ12", 14, MUTED, align="r", name="FOOTER")
    finish(s, "Not every shape may convert. Exact copies are not promised yet.", 9)


# ------------------------------------------------------------------ slide 10
def slide10(prs):
    s = new_slide(prs, "Stations change too", "Proposed")
    sh = s.shapes
    ty = 3.95  # track y
    g = group(sh, "DIAGRAM station replacement",
              "Plan view: a large dashed steam station with water tower, coal stage, signal box and platform blocks, "
              "replaced by a small solid electric station on the same track.")
    gs = g.shapes
    rails_h(gs, 0.8, ty, 11.9, STEEL, gap=0.26, step=0.4)
    # steam station (dashed)
    rect(gs, 0.8, 1.65, 6.4, 4.05, None, ORANGE, 2.5, radius=0.1, dash=MSO_LINE.DASH)
    label(gs, 1.0, 1.72, 3.0, 0.4, "Steam station", 18, ORANGE, True, align="l")
    rect(gs, 1.1, 2.45, 1.7, 0.8, LINE, None, radius=0.05, text="Coal stage", size=14, color=TEXT, bold=True)
    line(gs, 1.95, 3.25, 1.95, ty - 0.05, STEEL, 3)
    oval(gs, 3.95, 2.25, 1.15, 1.15, "#2F5A86", WATER, 2.5, text="Water tower", size=14, color=TEXT, bold=True)
    line(gs, 4.52, 3.4, 4.52, ty - 0.05, WATER, 3)
    for i in range(3):
        rect(gs, 1.9 + i * 1.3, 4.5, 1.22, 0.5, LINE, None, text="Platform" if i == 1 else "", size=14,
             color=TEXT, bold=True)
    rect(gs, 5.95, 4.7, 0.85, 0.5, LINE, None)
    poly(gs, [(5.85, 4.7), (6.375, 4.38), (6.9, 4.7)], "#59616C", None)
    label(gs, 5.7, 5.25, 1.4, 0.35, "Signal box", 14, TEXT, True)
    # arrow and label
    arrow(gs, 7.4, 2.7, 9.5, 2.7, ORANGE, 4)
    label(gs, 7.3, 2.82, 2.3, 0.7, "Aligned replacement, not a morph", 14, TEXT, True)
    # electric station (solid)
    label(gs, 10.0, 1.72, 2.7, 0.4, "Electric station", 18, ELEC, True, align="l")
    rect(gs, 10.35, 2.85, 1.5, 0.85, ELEC, None, radius=0.05)
    rect(gs, 10.55, 3.05, 0.45, 0.35, BG, None)
    rect(gs, 11.2, 3.05, 0.45, 0.35, BG, None)
    rect(gs, 10.1, 4.5, 2.2, 0.45, PANEL_ALT, ELEC, 2.5)
    finish(s, "A big steam station is replaced by a compact electric one on the same line.", 10)


# ------------------------------------------------------------------ slide 11
def slide11(prs):
    s = new_slide(prs, "Names and timetables", "Proposed")
    sh = s.shapes
    g = group(sh, "DIAGRAM station card",
              "A station card with a name bar and a timetable list, each ticked as kept where possible.")
    gs = g.shapes
    rect(gs, MARGIN, 1.5, 6.1, 4.25, PANEL, None, radius=0.08)
    rect(gs, 0.85, 1.75, 5.6, 0.75, PANEL_ALT, None, radius=0.06)
    label(gs, 1.05, 1.75, 3.8, 0.75, "Station name", 20, TEXT, True, align="l")
    tick(gs, 6.05, 2.125, 0.5)
    label(gs, 0.85, 2.65, 3.0, 0.35, "Timetable", 14, MUTED, True, align="l")
    for i, wbar in enumerate((2.6, 3.2, 2.0)):
        y = 3.05 + i * 0.58
        rect(gs, 0.85, y, 5.6, 0.48, PANEL_ALT, None, radius=0.06)
        oval(gs, 1.05, y + 0.11, 0.26, 0.26, None, MUTED, 2)
        rect(gs, 1.55, y + 0.16, wbar, 0.16, "#59616C", None, radius=0.04)
        tick(gs, 6.05, y + 0.24, 0.38)
    rect(gs, 1.5, 4.95, 4.0, 0.55, None, GREEN, 2.5, radius=0.27, text="Kept where possible", size=18, color=TEXT,
         bold=True)
    # removed gear (hatched, faded)
    g = group(sh, "DIAGRAM removed gear",
              "A coal stage and a water tower drawn hatched and faded, labelled Removed: the old coal and water gear goes.")
    gr = g.shapes
    label(gr, 7.4, 1.55, 5.3, 0.6, "Removed", 28, TEXT, True)
    gy = 4.55
    line(gr, 7.6, gy, 12.6, gy, "#59616C", 2, MSO_LINE.DASH)
    for px in (8.05, 8.85, 9.65):
        hatch(gr, px, 3.55, 0.12, gy - 3.55, HATCH_FG, BG, "#59616C", 1.5, MSO_LINE.DASH)
    hatch(gr, 7.85, 3.35, 2.0, 0.2, HATCH_FG, BG, "#59616C", 1.5, MSO_LINE.DASH)
    poly(gr, [(8.1, 3.35), (8.5, 2.8), (9.1, 2.65), (9.6, 3.35)], None, "#59616C", 2, dash=MSO_LINE.DASH)
    poly(gr, [(9.85, 3.4), (10.35, 4.1), (10.5, 4.1), (10.05, 3.4)], None, "#59616C", 2, dash=MSO_LINE.DASH)
    for px in (11.15, 12.0):
        hatch(gr, px, 3.7, 0.12, gy - 3.7, HATCH_FG, BG, "#59616C", 1.5, MSO_LINE.DASH)
    hatch(gr, 10.95, 2.55, 1.4, 1.1, HATCH_FG, BG, "#59616C", 1.5, MSO_LINE.DASH, radius=0.12)
    label(gr, 7.7, 4.7, 2.7, 0.35, "Coal stage", 14, MUTED, True)
    label(gr, 10.6, 4.7, 2.1, 0.35, "Water tower", 14, MUTED, True)
    finish(s, "Your station name and timetable could carry over. Old coal and water gear is removed.", 11)


# ------------------------------------------------------------------ slide 12
def slide12(prs):
    s = new_slide(prs, "Signals change", "Proposed")
    sh = s.shapes
    silver = "#C9CED4"
    g = group(sh, "DIAGRAM semaphore", "Silhouette of a mechanical semaphore signal with a flat arm.")
    gs = g.shapes
    rect(gs, MARGIN, 1.5, 4.4, 4.4, PANEL, None, radius=0.08)
    cx = 2.75
    rect(gs, cx - 0.09, 2.15, 0.18, 2.75, silver, None)
    rect(gs, cx - 0.35, 4.8, 0.7, 0.14, silver, None)
    rect(gs, cx - 0.05, 2.3, 1.25, 0.22, silver, None)
    oval(gs, cx - 0.5, 2.18, 0.34, 0.34, None, silver, 3)
    label(gs, 0.8, 5.05, 4.0, 0.7, "Semaphore", 20, TEXT, True)
    g = group(sh, "DIAGRAM electronic signal", "Silhouette of an electronic signal with three lamps on a post.")
    gs = g.shapes
    rect(gs, 8.33, 1.5, 4.4, 4.4, PANEL, None, radius=0.08)
    cx = 10.53
    rect(gs, cx - 0.09, 3.6, 0.18, 1.3, silver, None)
    rect(gs, cx - 0.35, 4.8, 0.7, 0.14, silver, None)
    rect(gs, cx - 0.42, 2.0, 0.84, 1.7, DARK, silver, 3, radius=0.12)
    for i, (f, o) in enumerate((("#5A2A2D", RED), ("#5A4C27", AMBER), ("#1F4A2A", GREEN))):
        oval(gs, cx - 0.21, 2.12 + i * 0.5, 0.42, 0.42, f, o, 2.5)
    label(gs, 8.53, 5.05, 4.0, 0.7, "Electronic signal", 20, TEXT, True)
    g = group(sh, "DIAGRAM arrow", "An arrow from the semaphore to the electronic signal, marked not yet specified.")
    gs = g.shapes
    arrow(gs, 5.3, 3.5, 8.1, 3.5, ORANGE, 4.5)
    label(gs, 5.2, 3.7, 3.0, 0.5, "Not yet specified", 20, TEXT, True)
    finish(s, "Mechanical signals may give way to electronic ones. The details are not set.", 12)


# ------------------------------------------------------------------ slide 13
def slide13(prs):
    s = new_slide(prs, "Mixed networks", "Proposed")
    sh = s.shapes
    ty = 3.65
    g = group(sh, "DIAGRAM mixed network",
              "A blue electric main line with orange steam branches joining it at three squares labelled Interchange.")
    gs = g.shapes
    rect(gs, 0.9, ty - 0.1, 11.5, 0.2, ELEC, None, radius=0.05)
    # branches: up, down, up
    poly(gs, [(3.4, ty), (3.4, 2.3), (1.7, 2.3)], None, ORANGE, 6, closed=False)
    poly(gs, [(6.8, ty), (6.8, 4.95), (8.5, 4.95)], None, ORANGE, 6, closed=False)
    poly(gs, [(10.2, ty), (10.2, 2.3), (11.9, 2.3)], None, ORANGE, 6, closed=False)
    for (tx, tyy) in ((1.55, 2.3), (8.65, 4.95), (12.05, 2.3)):
        oval(gs, tx - 0.2, tyy - 0.2, 0.4, 0.4, ORANGE, LIGHT, 2.5)
    for nx, up in ((3.4, True), (6.8, False), (10.2, True)):
        rect(gs, nx - 0.2, ty - 0.2, 0.4, 0.4, LIGHT, DARK, 2.5)
        label(gs, nx - 0.9, ty + 0.3 if up else ty - 0.68, 1.8, 0.38, "Interchange", 16, TEXT, True)
    label(gs, 0.9, ty - 0.7, 2.4, 0.4, "Electric main line", 18, ELEC, True, align="l")
    label(gs, 1.7, 1.75, 2.8, 0.4, "Steam branches", 18, ORANGE, True, align="l")
    finish(s, "Steam branches could live beside an electric main line.", 13)


# ------------------------------------------------------------------ slide 14
def slide14(prs):
    s = new_slide(prs, "Cargo crosses, trains do not", "Proposed")
    sh = s.shapes
    ty = 4.2
    d = -0.15  # vertical offset of the yard block
    g = group(sh, "DIAGRAM steam side", "A steam train stopped at a buffer at the end of its own track.")
    gs = g.shapes
    rect(gs, 0.7, ty, 3.75, 0.08, STEEL, None)
    wagon_side(gs, 0.9, ty - 0.6, 1.3, WAGON)
    loco_side(gs, 2.35, ty - 0.96, 1.7, ORANGE, "right")
    rect(gs, 4.3, ty - 0.55, 0.16, 0.6, RED, None)
    label(gs, 0.7, ty + 0.2, 3.7, 0.35, "Steam network", 16, ORANGE, True)
    g = group(sh, "DIAGRAM transfer yard",
              "A grey transfer yard called Interchange: a crane lifts crates and arrows carry them across.")
    gs = g.shapes
    rect(gs, 4.75, 1.7 + d, 3.9, 3.3, LINE, None, radius=0.1)
    label(gs, 4.75, 1.78 + d, 3.9, 0.45, "Interchange", 18, TEXT, True)
    # gantry crane
    rect(gs, 5.05, 2.95 + d, 0.14, 1.95, AMBER, None)
    rect(gs, 8.2, 2.95 + d, 0.14, 1.95, AMBER, None)
    rect(gs, 4.95, 2.75 + d, 3.49, 0.2, AMBER, None)
    rect(gs, 6.47, 2.95 + d, 0.5, 0.28, ORANGE, None)
    line(gs, 6.72, 3.23 + d, 6.72, 3.55 + d, DARK, 2.5)
    rect(gs, 6.52, 3.55 + d, 0.4, 0.4, LIGHT, DARK, 1.5)
    arrow(gs, 5.6, 2.42 + d, 7.9, 2.42 + d, ORANGE, 3.5)
    # crates and arrows
    for cx in (5.35, 5.8, 7.3, 7.75):
        rect(gs, cx, 4.45 + d, 0.4, 0.4, LIGHT, DARK, 1.5)
    arrow(gs, 6.4, 4.65 + d, 7.15, 4.65 + d, ORANGE, 3.5)
    g = group(sh, "DIAGRAM electric side", "An electric train on a raised beam, on the right.")
    gs = g.shapes
    rect(gs, 9.0, ty - 0.02, 3.73, 0.2, LIGHT, None)
    for px in (9.5, 11.0, 12.3):
        rect(gs, px, ty + 0.18, 0.16, 0.5, STEEL, None)
    carriage_side(gs, 9.2, ty - 0.62 - 0.02, 1.9, ELEC)
    carriage_side(gs, 11.2, ty - 0.62 - 0.02, 1.5, ELEC)
    label(gs, 9.0, ty + 0.75, 3.7, 0.35, "Electric network", 16, ELEC, True)
    note_strip(sh, "Agreed", "Steam trains cannot cross over", 5.5)
    finish(s, "Freight moves between the two networks at an interchange. Steam trains cannot cross over.", 14)


# ------------------------------------------------------------------ slide 15
def slide15(prs):
    s = new_slide(prs, "A Transfer Station?", "Undecided")
    sh = s.shapes
    g = group(sh, "DIAGRAM transfer station idea",
              "A dashed outline card showing a building silhouette with a large question mark, labelled Transfer Station.")
    gs = g.shapes
    rect(gs, 3.67, 1.5, 6.0, 4.4, PANEL, MUTED, 3, radius=0.1, dash=MSO_LINE.DASH)
    poly(gs, [(4.9, 2.95), (6.67, 1.85), (8.45, 2.95)], "#3A4049", None)
    rect(gs, 5.0, 2.95, 3.35, 1.85, "#3A4049", None)
    label(gs, 5.0, 2.9, 3.35, 1.9, "?", 96, TEXT, True)
    label(gs, 3.67, 5.0, 6.0, 0.7, "Transfer Station", 28, TEXT, True)
    finish(s, "A special transfer building is only an idea. It is not promised.", 15)


SLIDES = [slide9, slide10, slide11, slide12, slide13, slide14, slide15]
