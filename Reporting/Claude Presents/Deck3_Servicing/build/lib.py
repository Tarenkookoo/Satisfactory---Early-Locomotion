"""Shape library for Deck 3 ("Coal, Water and Servicing, Explained").

All helpers take `sh`, a shapes collection (slide.shapes or group.shapes), so the
same call works on a slide or inside a group. Sizes are in inches.
"""
import os
import re

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.abspath(os.path.join(HERE, "..", "..", "images"))

# --- style tokens (OUTLINE.md) -------------------------------------------
BG = "#1E2126"
PANEL = "#2B2F36"
PANEL_ALT = "#343A43"
LINE = "#4A515B"
TEXT = "#F2F2F2"
MUTED = "#A9B0B8"
ORANGE = "#F28C28"
RED = "#E5484D"
GREEN = "#3FB950"
AMBER = "#F2C14E"
WATER = "#4DA3FF"
ELEC = "#6CA0DC"
COAL = "#1B1D21"
COAL_LINE = "#6B7078"
STEEL = "#9AA0A6"
SLEEPER = "#6B5B4B"
PURPLE = "#B084F5"
DARK = "#101418"  # text on bright fills
WAGON = "#6C7480"

CHIPS = {
    "Agreed": (GREEN, DARK, None),
    "Proposed": (PURPLE, DARK, None),
    "Provisional": (WATER, DARK, None),
    "Undecided": (None, TEXT, MUTED),
    "Background": (LINE, TEXT, None),
}

FONT = "Arial"
FONT_R = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
_fcache = {}


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def I(v):
    return Inches(v)


# --- text measuring (Liberation Sans has Arial metrics) --------------------
def text_w(s, size, bold=False):
    key = bold
    if key not in _fcache:
        _fcache[key] = ImageFont.truetype(FONT_B if bold else FONT_R, 200)
    return _fcache[key].getlength(s) * size / 200.0 / 72.0


def wrap_count(s, size, width, bold=False):
    """Number of lines `s` takes at `size` pt in `width` inches."""
    lines = 0
    for para in s.split("\n"):
        words = para.split(" ")
        cur = ""
        n = 1
        for w in words:
            t = (cur + " " + w).strip()
            if text_w(t, size, bold) <= width or not cur:
                cur = t
            else:
                n += 1
                cur = w
        lines += n
    return lines


# --- deck ------------------------------------------------------------------
def new_deck(title, subject=""):
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    cp = prs.core_properties
    cp.title = title
    cp.author = "Claude (Early Locomotion project)"
    cp.last_modified_by = "Claude (Early Locomotion project)"
    cp.language = "en-GB"
    cp.subject = subject
    cp.keywords = "Satisfactory; Early Locomotion; steam; coal; water; servicing"
    # Theme: Arial fonts, readable hyperlink colours on a dark background.
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    xml = theme_part.blob.decode("utf8")
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*(")', r"\1Arial\2", xml)
    xml = re.sub(r'(<a:hlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}(")', r"\g<1>8FC7FF\2", xml)
    xml = re.sub(r'(<a:folHlink>\s*<a:srgbClr val=")[0-9A-Fa-f]{6}(")', r"\g<1>C9B2F5\2", xml)
    theme_part._blob = xml.encode("utf8")
    return prs


def new_slide(prs, title, chip_kind=None, title_size=36, title_w=10.2, title_box=None):
    """Title Only layout: real title placeholder (good for navigation), then the chip."""
    layout = prs.slide_layouts[5]
    s = prs.slides.add_slide(layout)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(BG)
    t = s.shapes.title
    bx = title_box or (0.6, 0.4, title_w, 0.8)
    t.left, t.top, t.width, t.height = I(bx[0]), I(bx[1]), I(bx[2]), I(bx[3])
    t.name = "TITLE"
    tf = t.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    _font(r, title_size, TEXT, True)
    if chip_kind:
        chip(s.shapes, chip_kind, 13.333 - 0.6 - 1.7, 0.58, name="CHIP")
    return s


def finish(slide, caption, n, caption_lines=None):
    """Caption bar (20 pt) and slide number, added last (reading order)."""
    sh = slide.shapes
    bar_y, bar_h = 6.3, 0.8
    bar = rect(sh, 0.6, bar_y, 12.13, bar_h, PANEL, None, radius=0.08, name="CAPTION BAR")
    rect(sh, 0.6, bar_y + 0.06, 0.08, bar_h - 0.12, ORANGE, None, name="CAPTION RULE")
    tb = label(sh, 0.85, bar_y, 11.75, bar_h, caption, 20, TEXT, align="l", anchor="m", name="CAPTION")
    if n is not None:
        label(sh, 11.73, 7.14, 1.0, 0.3, str(n), 14, MUTED, align="r", name="NUM")
    return tb


# --- low-level styling ------------------------------------------------------
def _font(run, size, color, bold=False, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.language_id = MSO_LANGUAGE_ID.ENGLISH_UK


def _strip_style(shape):
    el = shape._element
    for st in el.findall(qn("p:style")):
        el.remove(st)


def _set_alpha(fill_parent_elm, alpha):
    clr = fill_parent_elm.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))


def style(shape, fill=None, line=None, lw=1.5, dash=None, alpha=None):
    _strip_style(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _set_alpha(shape._element.spPr.find(qn("a:solidFill")), alpha)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
        if dash:
            shape.line.dash_style = dash


def set_alt(shape, descr, name=None):
    el = shape._element
    found = el.xpath(
        "./p:nvSpPr/p:cNvPr | ./p:nvGrpSpPr/p:cNvPr | ./p:nvPicPr/p:cNvPr | "
        "./p:nvCxnSpPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr"
    )
    c = found[0]
    c.set("descr", descr)
    if name:
        c.set("name", name)


def _fill_tf(tf, text, size, color, bold, align, anchor, margins, spacing=None):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [I(m) for m in margins]
    paras = text if isinstance(text, list) else [text]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        runs = para if isinstance(para, list) else [para]
        for rn in runs:
            if isinstance(rn, tuple):
                t, o = rn
            else:
                t, o = rn, {}
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))
            if o.get("link"):
                r.hyperlink.address = o["link"]


# --- primitives --------------------------------------------------------------
def rect(sh, x, y, w, h, fill=None, line=None, lw=1.5, radius=None, dash=None, alpha=None, name=None,
         text=None, size=14, color=TEXT, bold=False, align="c", anchor="m", margins=(0.06, 0.03, 0.06, 0.03)):
    shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    if text is not None:
        _fill_tf(shp.text_frame, text, size, color, bold, align, anchor, margins)
    return shp


def oval(sh, x, y, w, h, fill=None, line=None, lw=1.5, alpha=None, name=None, text=None, size=14, color=TEXT,
         bold=False, dash=None):
    shp = sh.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(w), I(h))
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    if text is not None:
        _fill_tf(shp.text_frame, text, size, color, bold, "c", "m", (0, 0, 0, 0))
    return shp


def label(sh, x, y, w, h, text, size=14, color=TEXT, bold=False, align="c", anchor="m", name=None,
          margins=(0.02, 0.01, 0.02, 0.01), spacing=None):
    tb = sh.add_textbox(I(x), I(y), I(w), I(h))
    _fill_tf(tb.text_frame, text, size, color, bold, align, anchor, margins, spacing)
    if name:
        tb.name = name
    return tb


def poly(sh, pts, fill=None, line=None, lw=1.5, closed=True, alpha=None, name=None, dash=None):
    """Freeform polygon / polyline through `pts` (inches)."""
    e = [(int(Inches(px)), int(Inches(py))) for px, py in pts]
    fb = sh.build_freeform(e[0][0], e[0][1], scale=1.0)
    fb.add_line_segments(e[1:], close=closed)
    shp = fb.convert_to_shape()
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    return shp


def _ends(conn, head, tail):
    ln = conn.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "med")
            e.set("len", "med")


def line(sh, x1, y1, x2, y2, color=STEEL, w=2.25, dash=None, head=False, tail=False, name=None):
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    _ends(c, head, tail)
    if name:
        c.name = name
    return c


def arrow(sh, x1, y1, x2, y2, color=ORANGE, w=2.25, dash=None, name=None):
    return line(sh, x1, y1, x2, y2, color, w, dash, False, True, name)


def curve(sh, p0, p1, p2, color=ORANGE, w=2.25, n=28, head=True, dash=None, name=None):
    """Quadratic Bezier polyline p0 -> p2 bending towards p1, optional arrow head at the end."""
    pts = []
    for i in range(n + 1):
        t = i / n
        pts.append(((1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * p1[0] + t * t * p2[0],
                    (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * p1[1] + t * t * p2[1]))
    shp = poly(sh, pts, None, color, w, closed=False, name=name, dash=dash)
    if head:
        _ends(shp, False, True)
    return shp


def group(sh, name, alt):
    g = sh.add_group_shape()
    g.name = name
    set_alt(g, alt, name)
    return g


# --- chips --------------------------------------------------------------------
def chip(sh, kind, x, y, w=1.7, h=0.4, size=14, text=None, name="CHIP", align="c"):
    fill, tcol, outline = CHIPS[kind]
    t = text if text is not None else kind.upper()
    shp = rect(sh, x, y, w, h, fill, outline, 2.0, radius=min(0.12, h / 2), name=name, text=t, size=size,
               color=tcol, bold=True, align=align, margins=(0.06, 0.0, 0.06, 0.0))
    return shp


def ichip(sh, kind, x, y, w=None, text=None, h=0.32):
    """Inline chip (14 pt, per MULTI_DECK_RULES: nothing below 14 pt)."""
    t = text if text is not None else kind.upper()
    if w is None:
        w = text_w(t, 14, True) + 0.3
    return chip(sh, kind, x, y, w, h, 14, t, name="ICHIP")


# --- pictures ---------------------------------------------------------------
def add_picture_fit(sh, fname, x, y, w, h, alt, mode="cover", focus=(0.5, 0.5), radius=0.08, name=None):
    path = os.path.join(IMG_DIR, fname)
    with Image.open(path) as im:
        iw, ih = im.size
    ar_box, ar_img = w / h, iw / ih
    if mode == "contain":
        if ar_img > ar_box:
            nw, nh = w, w / ar_img
        else:
            nw, nh = h * ar_img, h
        x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
    pic = sh.add_picture(path, I(x), I(y), I(w), I(h))
    if mode == "cover":
        if ar_img > ar_box:
            vis = ar_box / ar_img
            left = (1 - vis) * focus[0]
            pic.crop_left, pic.crop_right = left, 1 - vis - left
        else:
            vis = ar_img / ar_box
            top = (1 - vis) * focus[1]
            pic.crop_top, pic.crop_bottom = top, 1 - vis - top
    if radius:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        for g in list(av):
            av.remove(g)
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val %d" % int(min(0.5, radius / min(w, h)) * 100000))
    set_alt(pic, alt, name or os.path.splitext(fname)[0])
    return pic


def credit(sh, x, y, w, text, align="l", h=0.24, color=MUTED):
    return label(sh, x, y, w, h, text, 10, color, align=align, anchor="t", name="CREDIT", margins=(0, 0, 0, 0))


# --- tables ----------------------------------------------------------------------
def cell_border(cell, color, w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    # borders must precede the fill element inside tcPr
    first = tcPr[0] if len(tcPr) else None
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(w_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", color.lstrip("#"))
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        if first is not None:
            first.addprevious(ln)
        else:
            tcPr.append(ln)


def set_cell(cell, text, size, color=TEXT, bold=False, fill=None, align="l", border=LINE, margins=(0.08, 0.04, 0.08, 0.04)):
    cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = [I(m) for m in margins]
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    paras = text if isinstance(text, list) else [text]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        runs = para if isinstance(para, list) else [para]
        for rn in runs:
            t, o = (rn if isinstance(rn, tuple) else (rn, {}))
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold))
            if o.get("link"):
                r.hyperlink.address = o["link"]
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    else:
        cell.fill.background()
    if border:
        cell_border(cell, border)


def add_table(sh, x, y, w, h, rows, cols, col_w, alt, name="TABLE"):
    gf = sh.add_table(rows, cols, I(x), I(y), I(w), I(h))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = I(cw)
    for r in tbl.rows:
        r.height = I(h / rows)
    set_alt(gf, alt, name)
    return gf, tbl


# --- reusable objects (RESEARCH_NOTES "Reusable objects" and style tokens) ---------
def rails(sh, x, y, length, color=STEEL):
    rect(sh, x, y, length, 0.06, color, None, name="rail")
    rect(sh, x, y + 0.22, length, 0.06, color, None, name="rail")


def sleepers(sh, x, y, length, pitch=0.35, h=0.4):
    n = int(length / pitch)
    for i in range(n + 1):
        rect(sh, x + i * pitch, y - 0.06, 0.06, h, SLEEPER, None, name="sleeper")


def wheel(sh, cx, cy, d, fill=COAL, rim=STEEL, lw=1.5):
    oval(sh, cx - d / 2, cy - d / 2, d, d, fill, rim, lw, name="wheel")
    oval(sh, cx - d * 0.12, cy - d * 0.12, d * 0.24, d * 0.24, rim, None, name="hub")


def tank_gauge(sh, x, y, w, h, fill_frac, fill_color, outline=STEEL, interior=PANEL, name="tank"):
    rect(sh, x, y, w, h, interior, outline, 2.25, radius=0.1, name=name)
    fh = (h - 0.12) * fill_frac
    if fill_frac > 0:
        f = rect(sh, x + 0.06, y + h - 0.06 - fh, w - 0.12, fh, fill_color, "#8A9099" if fill_color == COAL else None,
                 1.25, name=name + " fill")
    return None


def arc_poly(cx, cy, r_out, r_in, a0, a1, n=24):
    """Points of a ring segment from angle a0 to a1 (degrees, 0 = right, counter-clockwise up)."""
    import math
    pts = []
    for i in range(n + 1):
        a = math.radians(a0 + (a1 - a0) * i / n)
        pts.append((cx + r_out * math.cos(a), cy - r_out * math.sin(a)))
    for i in range(n, -1, -1):
        a = math.radians(a0 + (a1 - a0) * i / n)
        pts.append((cx + r_in * math.cos(a), cy - r_in * math.sin(a)))
    return pts


def loco_side(sh, x, y, scale=1.0, tender=False, color=ORANGE, wheels=3, facing="right", name="loco"):
    """Small side-view steam locomotive drawn from shapes. (x, y) is the top-left of the cab.
    Natural size about 2.0 x 1.0 in (plus 1.0 for a tender) at scale 1."""
    s = scale
    out = []
    flip = facing == "left"
    W = 2.0 * s
    def X(dx, w=0):
        return x + (W - dx - w) if flip else x + dx

    # boiler (long, rounded)
    out.append(rect(sh, X(0.55 * s, 1.35 * s), y + 0.22 * s, 1.35 * s, 0.5 * s, color, None, radius=0.2 * s, name="boiler"))
    # cab
    out.append(rect(sh, X(0.0, 0.62 * s), y, 0.62 * s, 0.72 * s, color, None, radius=0.04 * s, name="cab"))
    out.append(rect(sh, X(0.08 * s, 0.4 * s), y + 0.12 * s, 0.4 * s, 0.2 * s, COAL, None, name="cab window"))
    # chimney
    out.append(rect(sh, X(1.7 * s, 0.14 * s), y + 0.02 * s, 0.14 * s, 0.26 * s, COAL, COAL_LINE, 1.0, name="chimney"))
    # dome
    out.append(oval(sh, X(1.0 * s, 0.2 * s), y + 0.12 * s, 0.2 * s, 0.16 * s, color, None, name="dome"))
    # footplate / frame
    out.append(rect(sh, X(0.0, 1.95 * s), y + 0.72 * s, 1.95 * s, 0.08 * s, COAL, COAL_LINE, 1.0, name="frame"))
    # wheels
    d = 0.42 * s
    n = wheels
    span = 1.2 * s
    start = 0.45 * s
    for i in range(n):
        cx = start + (span * i / max(1, n - 1)) if n > 1 else start
        wcx = x + (W - cx) if flip else x + cx
        out.append(oval(sh, wcx - d / 2, y + 0.8 * s - 0.02 * s, d, d, COAL, STEEL, 1.5, name="wheel"))
    if tender:
        tx = X(-1.0 * s) if not flip else x + W + 0.08 * s
        tx = x - 1.0 * s - 0.08 * s if not flip else x + W + 0.08 * s
        out.append(rect(sh, tx, y + 0.12 * s, 1.0 * s, 0.6 * s, WAGON, None, radius=0.03 * s, name="tender"))
        out.append(rect(sh, tx + 0.1 * s, y + 0.06 * s, 0.8 * s, 0.1 * s, COAL, COAL_LINE, 1.0, name="coal"))
        out.append(rect(sh, tx, y + 0.72 * s, 1.0 * s, 0.08 * s, COAL, COAL_LINE, 1.0, name="tender frame"))
        for k in (0.25, 0.75):
            out.append(oval(sh, tx + k * s - d / 2, y + 0.8 * s - 0.02 * s, d, d, COAL, STEEL, 1.5, name="wheel"))
    return out
