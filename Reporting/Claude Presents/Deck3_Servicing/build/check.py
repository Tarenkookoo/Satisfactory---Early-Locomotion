"""Per-slide checks: word counts, smallest font, overflow estimate, alt text, notes, chips."""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402
from lib import wrap_count, text_w  # noqa: E402

PATH = os.path.abspath(os.path.join(HERE, "..", "Coal_Water_and_Servicing_Explained.pptx"))
SKIP_COUNT = ("TITLE", "CHIP", "ICHIP", "NUM", "CREDIT", "CHIP LEGEND")


def walk(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:  # group
            yield from walk(s.shapes)


def words(t):
    """Whitespace tokens that contain a letter (bare numbers such as badge digits are not words)."""
    return len([w for w in t.split() if any(c.isalpha() for c in w)])


def main():
    prs = Presentation(PATH)
    print("slides:", len(prs.slides), " size KB:", os.path.getsize(PATH) // 1024)
    hdr = "%3s %-34s %-12s %5s %5s %6s %s" % ("#", "title", "chip", "body", "total", "minpt", "issues")
    print(hdr)
    for i, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title.text_frame.text if sl.shapes.title is not None else "(none)"
        chip, body, total, minpt = "-", 0, 0, 99
        issues = []
        if sl.has_notes_slide:
            issues.append("HAS NOTES")
        for s in walk(sl.shapes):
            nm = s.name or ""
            if s.shape_type == 13 or s.shape_type == 6 and nm:
                c = s._element.xpath("./p:nvPicPr/p:cNvPr | ./p:nvGrpSpPr/p:cNvPr")
                if c and not c[0].get("descr"):
                    issues.append("no alt: " + nm)
            texts = []
            if s.has_text_frame and s.text_frame.text.strip():
                texts.append((s.text_frame, s.width, s.height, s.text_frame.margin_left, s.text_frame.margin_right,
                              s.text_frame.margin_top, s.text_frame.margin_bottom))
            if getattr(s, "has_table", False) and s.has_table:
                for r in s.table.rows:
                    for c in r.cells:
                        if c.text_frame.text.strip():
                            total += words(c.text_frame.text)
                            body += words(c.text_frame.text)
                            for p in c.text_frame.paragraphs:
                                for rn in p.runs:
                                    if rn.font.size:
                                        minpt = min(minpt, rn.font.size.pt) if nm != "CREDIT" else minpt
                continue
            for tf, w, h, ml, mr, mt, mb in texts:
                t = tf.text
                if nm == "CHIP" and chip == "-":
                    chip = t
                if not nm.startswith("NUM") and nm != "CREDIT":
                    total += words(t)
                if nm not in SKIP_COUNT:
                    body += words(t)
                sizes = [rn.font.size.pt for p in tf.paragraphs for rn in p.runs if rn.font.size]
                if sizes and nm != "CREDIT":
                    minpt = min(minpt, min(sizes))
                # overflow estimate
                if sizes:
                    size = max(sizes)
                    bold = any(rn.font.bold for p in tf.paragraphs for rn in p.runs)
                    inner_w = Emu(w - ml - mr).inches
                    inner_h = Emu(h - mt - mb).inches
                    nlines = sum(wrap_count(p.text if p.text else " ", size, inner_w, bold) for p in tf.paragraphs)
                    need = nlines * size * 1.2 / 72
                    longest = max((text_w(wd, size, bold) for p in tf.paragraphs for wd in p.text.split()), default=0)
                    if need > inner_h + 0.02:
                        issues.append("overflow? %s '%s' need %.2f have %.2f" % (nm, t[:24], need, inner_h))
                    if longest > inner_w + 0.01:
                        issues.append("word too wide: %s '%s'" % (nm, t[:24]))
        flag = ""
        if body > 45 and i not in (2, 24):
            flag = " >45!"
        elif body > 25 and i not in (2, 24):
            flag = " >25"
        print("%3d %-34s %-12s %5d %5d %6s %s%s" % (i, title[:34], chip[:12], body, total,
                                                    ("%g" % minpt) if minpt < 99 else "-", "; ".join(issues), flag))


if __name__ == "__main__":
    main()
