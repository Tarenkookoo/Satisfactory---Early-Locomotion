"""Verify the built deck against DECK_BRIEF.md and MULTI_DECK_RULES.md. Usage: python3 verify_brief.py"""
import os
import re
import sys
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "..", "Signals_and_Signal_Boxes_Explained.pptx")
BRIEF = os.path.join(HERE, "..", "DECK_BRIEF.md")
IMG = os.path.join(HERE, "..", "..", "images")
EXPECT_TAGS = {"A": [15, 16, 17, 19, 20, 22, 23, 24], "P": [21], "V": [18, 25], "U": [26],
               "B": list(range(3, 15)), "-": [1, 2, 27]}
TAGNAME = {"AGREED": "A", "PROPOSED": "P", "PROVISIONAL": "V", "UNDECIDED": "U", "BACKGROUND": "B"}


def brief_rows():
    rows = {}
    for ln in open(BRIEF, encoding="utf-8"):
        m = re.match(r"\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*\"(.+?)\"\s*\|", ln)
        if m:
            tag = ln.split("|")[5].strip()
            rows[int(m.group(1))] = (m.group(2), m.group(3), tag)
    return rows


def main():
    prs = Presentation(DECK)
    ok = True
    rows = brief_rows()
    print("slides:", len(prs.slides), "(expected 27)")
    ok &= len(prs.slides) == 27
    tagcount = {k: [] for k in EXPECT_TAGS}
    for i, sl in enumerate(prs.slides, 1):
        title = sl.shapes.title.text_frame.text
        cap = None
        main_chip = None
        for sh in sl.shapes:
            if sh.name == "Caption":
                cap = sh.text_frame.text
            if sh.name.startswith("Chip:") and sh.left > 10.5 * 914400 and sh.top < 914400:
                main_chip = TAGNAME[sh.text_frame.text.strip().upper()]
            if i == 1 and sh.name == "Caption":
                cap = sh.text_frame.text
        tagcount[main_chip or "-"].append(i)
        b = rows.get(i)
        if b:
            if title != b[0]:
                print(f"  title differs slide {i}: {title!r} vs {b[0]!r}"); ok = False
            if cap != b[1]:
                print(f"  caption differs slide {i}: {cap!r} vs {b[1]!r}"); ok = False
        else:
            print(f"  slide {i}: no brief row parsed")
            ok = False
        # pictures: aspect ratio preserved after crop
        for sh in sl.shapes:
            if sh.shape_type == 13:
                iw, ih = Image.open(os.path.join(IMG, os.path.basename(sh.image.filename or "")) if False else
                                    __import__("io").BytesIO(sh.image.blob)).size
                vis_w = iw * (1 - sh.crop_left - sh.crop_right)
                vis_h = ih * (1 - sh.crop_top - sh.crop_bottom)
                a_img, a_box = vis_w / vis_h, sh.width / sh.height
                if abs(a_img / a_box - 1) > 0.01:
                    print(f"  slide {i}: picture {sh.name} stretched ({a_img:.3f} vs {a_box:.3f})"); ok = False
    for k, exp in EXPECT_TAGS.items():
        got = tagcount[k]
        flag = "ok" if got == exp else "MISMATCH"
        print(f"tag {k}: {len(got)} {got} {flag}")
        ok &= got == exp
    # package-level checks
    z = zipfile.ZipFile(DECK)
    names = z.namelist()
    print("notes slides:", [n for n in names if "notesSlide" in n] or "none")
    ok &= not [n for n in names if "notesSlide" in n]
    bad_lang = 0
    sizes = set()
    for n in names:
        if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
            xml = z.read(n).decode("utf-8")
            for m in re.finditer(r"<a:(rPr|endParaRPr)\b([^>]*)>", xml):
                if 'lang="en-GB"' not in m.group(2):
                    bad_lang += 1
            sizes.update(int(s) for s in re.findall(r'<a:rPr[^>]*\bsz="(\d+)"', xml))
            if "<p:transition" in xml or "<p:timing" in xml:
                print("  transition/animation found in", n); ok = False
            if re.search(r"<p:ph [^>]*type=\"(body|dt|ftr|sldNum)\"", xml):
                print("  leftover placeholder in", n); ok = False
    print("runs without lang en-GB:", bad_lang, "| font sizes used (pt):", sorted(s / 100 for s in sizes))
    ok &= bad_lang == 0
    print("document title:", prs.core_properties.title, "| author:", prs.core_properties.author,
          "| language:", prs.core_properties.language)
    print("ALL CHECKS PASSED" if ok else "SOME CHECKS FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
