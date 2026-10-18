"""Shape library for the Deck 2 build (python-pptx, native shapes only).

Everything is positioned in inches on a 13.333 x 7.5 in canvas. Style tokens come from
Reporting/Claude Presents/OUTLINE.md ("Style tokens"). Diagram objects follow RESEARCH_NOTES.md
("Reusable objects").
"""
import math
import os
import re

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.normpath(os.path.join(HERE, "..", "..", "images"))

FONT = "Arial"
LANG = "en-GB"
AUTHOR = "Claude (Early Locomotion project)"

C = dict(
    bg="#1E2126", panel="#2B2F36", panel_alt="#343A43", line="#4A515B", text="#F2F2F2",
    muted="#A9B0B8", orange="#F28C28", red="#E5484D", green="#3FB950", amber="#F2C14E",
    blue="#4DA3FF", elec="#6CA0DC", coal="#1B1D21", steel="#9AA0A6", sleeper="#6B5B4B",
    wagon="#6C7480", post="#E8E8E8", ink="#101418", coal_line="#6B7078",
    red_text="#FF7A7E",  # lighter red for text on dark (4.5:1 needs more than #E5484D)
    purple="#B084F5",
)

# tag letter -> (label, fill, text colour, outline)
TAGS = {
    "A": ("Agreed", C["green"], C["ink"], None),
    "P": ("Proposed", C["purple"], C["ink"], None),
    "V": ("Provisional", C["blue"], C["ink"], None),
    "U": ("Undecided", None, C["text"], C["muted"]),
    "B": ("Background", C["line"], C["text"], None),
}
TAG_WORDS = {v[0].upper() for v in TAGS.values()}


def E(v):
    return Emu(int(round(v * 914400)))


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


# ----------------------------------------------------------------------------------------------
# low-level helpers
# ----------------------------------------------------------------------------------------------
def _strip_style(shape):
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)


def set_alt(shape, text, name=None):
    el = shape._element.xpath(".//p:cNvPr")[0]
    el.set("descr", text)
    if name:
        el.set("name", name)


def set_name(shape, name):
    shape._element.xpath(".//p:cNvPr")[0].set("name", name)


def set_alpha(shape, opacity_pct):
    sf = shape._element.spPr.find(qn("a:solidFill"))
    clr = sf.find(qn("a:srgbClr"))
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(opacity_pct * 1000)))


def style_shape(shape, fill=None, line=None, line_w=1.5, dash=None, alpha=None):
    _strip_style(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            set_alpha(shape, alpha)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
        if dash:
            shape.line.dash_style = dash
    return shape


def add_arrowheads(shape, head=False, tail=True, size="med"):
    ln = shape.line._get_or_add_ln()
    if head:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle"); e.set("w", size); e.set("len", size)
    if tail:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle"); e.set("w", size); e.set("len", size)


def set_text(shape, text, size=14, color=None, bold=False, align="l", anchor="m",
             margins=(0.05, 0.03, 0.05, 0.03), wrap=True, spacing=None, italic=False):
    """text: str (use \\n for paragraphs) or list of paragraphs; a paragraph is a str or a list of
    (text, {bold, color, size}) runs."""
    color = color or C["text"]
    tf = shape.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [E(m) for m in margins]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = text.split("\n") if isinstance(text, str) else text
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for rt, opt in runs:
            r = p.add_run()
            r.text = rt
            f = r.font
            f.name = FONT
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = italic
            f.color.rgb = rgb(opt.get("color", color))
    return shape


def textbox(shapes, x, y, w, h, text, size=14, color=None, bold=False, align="l", anchor="m",
            margins=(0.02, 0.02, 0.02, 0.02), name=None, wrap=True, spacing=None, italic=False):
    tb = shapes.add_textbox(E(x), E(y), E(w), E(h))
    _strip_style(tb)
    set_text(tb, text, size, color, bold, align, anchor, margins, wrap, spacing, italic)
    if name:
        set_name(tb, name)
    return tb


def rect(shapes, x, y, w, h, fill=None, line=None, line_w=1.5, dash=None, alpha=None, radius=None,
         name=None, shape=None):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = shapes.add_shape(st, E(x), E(y), E(w), E(h))
    style_shape(s, fill, line, line_w, dash, alpha)
    if radius and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if name:
        set_name(s, name)
    return s


def ellipse(shapes, cx, cy, d, fill=None, line=None, line_w=1.5, dash=None, alpha=None, name=None, dh=None):
    dh = dh if dh is not None else d
    s = shapes.add_shape(MSO_SHAPE.OVAL, E(cx - d / 2), E(cy - dh / 2), E(d), E(dh))
    style_shape(s, fill, line, line_w, dash, alpha)
    if name:
        set_name(s, name)
    return s


def line(shapes, x1, y1, x2, y2, color=None, w=1.5, dash=None, head=False, tail=False, name=None):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color or C["steel"])
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    if head or tail:
        add_arrowheads(c, head, tail)
    if name:
        set_name(c, name)
    return c


def poly(shapes, pts, fill=None, line_c=None, line_w=1.5, close=True, dash=None, alpha=None,
         head=False, tail=False, name=None):
    """Freeform through absolute inch points. close=False gives an open polyline."""
    fb = shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=close)
    s = fb.convert_to_shape()
    style_shape(s, fill if close else None, line_c, line_w, dash, alpha)
    if head or tail:
        add_arrowheads(s, head, tail)
    if name:
        set_name(s, name)
    return s


def group(shapes, name, alt=None):
    g = shapes.add_group_shape()
    set_name(g, name)
    if alt:
        set_alt(g, alt)
    return g


def arrow(shapes, x1, y1, x2, y2, color=None, w=2.25, dash=None, both=False, name=None):
    return line(shapes, x1, y1, x2, y2, color or C["orange"], w, dash, head=both, tail=True, name=name)


def chip(shapes, tag, x=11.05, y=0.5, w=1.65, h=0.4, size=14, text=None, name=None):
    """Status chip. tag is one of A, P, V, U, B. text overrides the label (inline chips with words)."""
    label, fill, tcol, outline = TAGS[tag]
    s = rect(shapes, x, y, w, h, fill=fill, line=outline, line_w=2, radius=0.08)
    set_text(s, text if text is not None else label.upper(), size, tcol, bold=True, align="c",
             margins=(0.03, 0.01, 0.03, 0.01), wrap=text is not None)
    set_name(s, name or f"Chip: {label}")
    set_alt(s, f"Status tag: {label}")
    return s


# ----------------------------------------------------------------------------------------------
# pictures
# ----------------------------------------------------------------------------------------------
def picture(shapes, fname, x, y, w, h, alt, focus=(0.5, 0.5), radius=0.08, name=None):
    """Place a bundle image, cropped (never stretched) to the w x h box."""
    path = os.path.join(IMG, fname)
    iw, ih = Image.open(path).size
    pic = shapes.add_picture(path, E(x), E(y), E(w), E(h))
    box_ar, img_ar = w / h, iw / ih
    if img_ar > box_ar:
        excess = 1 - box_ar / img_ar
        pic.crop_left = excess * focus[0]
        pic.crop_right = excess - excess * focus[0]
    elif img_ar < box_ar:
        excess = 1 - img_ar / box_ar
        pic.crop_top = excess * focus[1]
        pic.crop_bottom = excess - excess * focus[1]
    if radius:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        pg = pic._element.spPr.find(qn("a:prstGeom"))
        av = pg.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(pg, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", "val %d" % int(min(0.5, radius / min(w, h)) * 100000))
    set_alt(pic, alt, name or fname)
    return pic


def credit(shapes, x, y, w, text, h=0.2, align="l", name="Credit"):
    return textbox(shapes, x, y, w, h, text, size=10, color=C["muted"], anchor="t", align=align,
                   margins=(0, 0, 0, 0), name=name)


# ----------------------------------------------------------------------------------------------
# railway objects (RESEARCH_NOTES.md "Reusable objects")
# ----------------------------------------------------------------------------------------------
def rails(shapes, x, y, length, sleepers=True, pitch=0.35, name="Track"):
    """Two rails 0.06 in thick, 0.28 in apart (top rail at y). Sleepers 0.06 x 0.4 in."""
    g = group(shapes, name, "Two rails on sleepers.")
    gs = g.shapes
    if sleepers:
        n = int(length / pitch)
        off = (length - (n - 1) * pitch) / 2
        for i in range(n):
            rect(gs, x + off + i * pitch - 0.03, y - 0.06, 0.06, 0.40, fill=C["sleeper"])
    rect(gs, x, y, length, 0.06, fill=C["steel"])
    rect(gs, x, y + 0.28, length, 0.06, fill=C["steel"])
    return g


def _wheel(gs, cx, cy, d):
    ellipse(gs, cx, cy, d, fill=C["coal"], line=C["steel"], line_w=1)


def train(shapes, x, rail_y, facing="right", wagons=2, scale=1.0, name="Train", alt=None):
    """Locomotive plus wagons standing on rails whose top rail line is rail_y. x = left edge.
    Loco 1.0 x 0.55 in orange; wagons 0.8 x 0.45 in; gaps 0.08 in."""
    s = scale
    g = group(shapes, name, alt or "A steam locomotive with wagons.")
    gs = g.shapes
    lw, lh, ww, wh, gap = 1.0 * s, 0.55 * s, 0.8 * s, 0.45 * s, 0.08 * s
    body_bottom = rail_y - 0.07 * s
    wd = 0.2 * s
    total = lw + wagons * (ww + gap)
    if facing == "right":
        loco_x = x + wagons * (ww + gap)
        wagon_xs = [x + i * (ww + gap) for i in range(wagons)]
    else:
        loco_x = x
        wagon_xs = [x + lw + gap + i * (ww + gap) for i in range(wagons)]
    for wx in wagon_xs:
        rect(gs, wx, body_bottom - wh, ww, wh, fill=C["wagon"], line=C["steel"], line_w=1, radius=0.03)
        for k in (0.22, 0.78):
            _wheel(gs, wx + ww * k, rail_y, wd)
    # locomotive body
    rect(gs, loco_x, body_bottom - lh, lw, lh, fill=C["orange"], radius=0.08)
    if facing == "right":
        chim_x = loco_x + lw - 0.30 * s
        bunk_x = loco_x + 0.05 * s
    else:
        chim_x = loco_x + 0.15 * s
        bunk_x = loco_x + lw - 0.05 * s - 0.28 * s
    rect(gs, chim_x, body_bottom - lh - 0.28 * s, 0.15 * s, 0.30 * s, fill=C["coal"], line=C["coal_line"], line_w=1)
    rect(gs, bunk_x, body_bottom - lh + 0.02 * s, 0.28 * s, 0.30 * s, fill=C["coal"], line=C["coal_line"], line_w=1)
    for k in (0.2, 0.5, 0.8):
        _wheel(gs, loco_x + lw * k, rail_y, wd * 1.1)
    return g, total


def _rot(px, py, lx, ly, theta_deg):
    """Local (lx along arm, ly downward) -> global, rotating the arm up by theta around (px, py)."""
    t = math.radians(theta_deg)
    return (px + lx * math.cos(t) + ly * math.sin(t), py - lx * math.sin(t) + ly * math.cos(t))


def semaphore(shapes, x, base_y, state="stop", scale=1.0, kind="stop", name=None, alt=None,
              lamp=True, arm_dir=1):
    """Signal post (0.12 x 1.1 in) with an arm pivoted 0.1 in below the top.
    state 'stop' = arm flat; 'clear' = arm raised 45 degrees. kind 'stop' (red, white stripe) or
    'distant' (amber, fishtail notch). x = post centre; base_y = bottom of the post. arm_dir 1 = arm
    points right, -1 = left. Returns (group, pivot)."""
    s = scale
    top = base_y - 1.1 * s
    px, py = x, top + 0.1 * s
    g = group(shapes, name or f"Signal ({kind}, {state})",
              alt or f"A semaphore signal post with its arm {'flat (stop)' if state == 'stop' else 'raised (clear)'}.")
    gs = g.shapes
    rect(gs, x - 0.15 * s, base_y - 0.08 * s, 0.30 * s, 0.08 * s, fill=C["steel"])
    rect(gs, x - 0.06 * s, top, 0.12 * s, 1.1 * s - 0.08 * s, fill=C["post"])
    L, t = 0.7 * s, 0.14 * s
    theta = 0 if state == "stop" else 45
    d = arm_dir

    def gp(lx, ly):
        gx, gy = _rot(px, py, lx, ly, theta)
        return (px + d * (gx - px), gy)

    if kind == "stop":
        pts = [gp(0, -t / 2), gp(L, -t / 2), gp(L, t / 2), gp(0, t / 2)]
        poly(gs, pts, fill=C["red"], line_c=None)
        a, b = 0.50 * L, 0.50 * L + 0.11 * s
        poly(gs, [gp(a, -t / 2), gp(b, -t / 2), gp(b, t / 2), gp(a, t / 2)], fill="#FFFFFF", line_c=None)
    else:
        n = 0.18 * s
        pts = [gp(0, -t / 2), gp(L, -t / 2), gp(L - n, 0), gp(L, t / 2), gp(0, t / 2)]
        poly(gs, pts, fill=C["amber"], line_c=None)
        a, b = 0.50 * L, 0.50 * L + 0.08 * s
        poly(gs, [gp(a, -t / 2), gp(b, -t / 2), gp(b, t / 2), gp(a, t / 2)], fill=C["coal"], line_c=None)
    ellipse(gs, px, py, 0.07 * s + 0.03, fill=C["coal"], line="#FFFFFF", line_w=0.75)
    if lamp:
        col = (C["red"] if kind == "stop" else C["amber"]) if state == "stop" else C["green"]
        ellipse(gs, px - d * 0.20 * s, py, 0.17 * s + 0.02, fill=col, line="#FFFFFF", line_w=1)
    return g, (px, py)


def block_band(shapes, x, y, w, state, label, size=18, h=0.4, name=None):
    col = C["red"] if state == "occupied" else C["green"]
    r = rect(shapes, x, y, w, h, fill=col, alpha=25, line=col, line_w=1)
    set_text(r, label, size, C["text"], bold=True, align="c", margins=(0.03, 0, 0.03, 0))
    set_name(r, name or f"Block band ({state})")
    return r


def person_icon(shapes, cx, base_y, scale=1.0, color=None, name="Signaller icon"):
    col = color or C["post"]
    g = group(shapes, name, "A simple icon of a person: the signaller.")
    gs = g.shapes
    hr = 0.13 * scale
    ellipse(gs, cx, base_y - 0.55 * scale + hr, hr * 2, fill=col)
    b = gs.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, E(cx - 0.24 * scale), E(base_y - 0.38 * scale),
                     E(0.48 * scale), E(0.38 * scale))
    style_shape(b, col)
    b.adjustments[0] = 0.45
    return g


def lever(shapes, x, base_y, color, pulled=False, h=0.9, w=0.12, scale=1.0):
    """A frame lever: vertical bar on a pivot. pulled = leaning back 25 degrees."""
    h, w = h * scale, w * scale
    bar = rect(shapes, x - w / 2, base_y - h, w, h, fill=color, line=C["ink"], line_w=0.75)
    if pulled:
        # rotate about the base: shift centre so the base stays put
        ang = 25
        cx, cy = x, base_y - h / 2
        t = math.radians(ang)
        ncx = x + (h / 2) * math.sin(t)
        ncy = base_y - (h / 2) * math.cos(t)
        bar.left, bar.top = E(ncx - w / 2), E(ncy - h / 2)
        bar.rotation = ang
    ellipse(shapes, x, base_y, 0.2 * scale, fill=C["steel"], line=C["ink"], line_w=0.75)
    return bar


def box_building(shapes, x, y, w, h, color=None, line_c=None, windows=True, roof=True, name="Signal box",
                 alt=None, door=True):
    """Small signal-box building: walls, pitched roof, window strip."""
    g = group(shapes, name, alt or "A small signal box building.")
    gs = g.shapes
    fill = color or C["panel_alt"]
    lc = line_c or C["orange"]
    rect(gs, x, y + h * 0.22, w, h * 0.78, fill=fill, line=lc, line_w=1.5)
    if roof:
        poly(gs, [(x - 0.08, y + h * 0.24), (x + w / 2, y), (x + w + 0.08, y + h * 0.24)],
             fill=C["line"], line_c=lc, line_w=1.5)
    if windows:
        rect(gs, x + w * 0.1, y + h * 0.36, w * 0.8, h * 0.22, fill=C["blue"], alpha=55, line=C["post"], line_w=1)
        for k in range(1, 4):
            line(gs, x + w * 0.1 + w * 0.8 * k / 4, y + h * 0.36, x + w * 0.1 + w * 0.8 * k / 4, y + h * 0.58,
                 C["post"], 1)
    if door:
        rect(gs, x + w * 0.42, y + h * 0.70, w * 0.16, h * 0.30, fill=C["coal"], line=C["steel"], line_w=1)
    return g


def dimension(shapes, x1, x2, y, label, size=16, color=None, label_above=True):
    col = color or C["steel"]
    g = group(shapes, "Dimension line", f"A dimension line labelled {label}.")
    gs = g.shapes
    line(gs, x1, y, x2, y, col, 1.5, head=True, tail=True)
    line(gs, x1, y - 0.12, x1, y + 0.12, col, 1.5)
    line(gs, x2, y - 0.12, x2, y + 0.12, col, 1.5)
    ty = y - 0.42 if label_above else y + 0.1
    textbox(gs, x1, ty, x2 - x1, 0.32, label, size, C["text"], bold=True, align="c")
    return g


# ----------------------------------------------------------------------------------------------
# deck and slide scaffolding
# ----------------------------------------------------------------------------------------------
class SlideCtx:
    def __init__(self, deck, n, title, tag, title_box=None, size=36):
        self.deck, self.n, self.tag = deck, n, tag
        self.slide = deck.prs.slides.add_slide(deck.layout)
        s = self.slide
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(C["bg"])
        t = s.shapes.title
        bx = title_box or (0.6, 0.45, 10.4, 0.8)
        t.left, t.top, t.width, t.height = [E(v) for v in bx]
        tf = t.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = E(0)
        tf.margin_top = tf.margin_bottom = E(0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = title
        r.font.name, r.font.size, r.font.bold = FONT, Pt(size), True
        r.font.color.rgb = rgb(C["text"])
        set_name(t, "Title")
        if tag:
            chip(s.shapes, tag)

    @property
    def shapes(self):
        return self.slide.shapes

    def finish(self, caption, number=True):
        sh = self.shapes
        cap = rect(sh, 0.6, 6.35, 12.13, 0.7, fill=C["panel"], name="Caption")
        set_text(cap, caption, 20, C["text"], anchor="m", margins=(0.28, 0.02, 0.15, 0.02))
        rule = rect(sh, 0.6, 6.35, 0.08, 0.7, fill=C["orange"], name="Caption rule")
        set_alt(rule, "")
        if number:
            textbox(sh, 11.7, 7.1, 1.03, 0.3, str(self.n), 14, C["muted"], align="r", name="Slide number")


class Deck:
    def __init__(self):
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        self.prs = prs
        master = prs.slide_master
        master.background.fill.solid()
        master.background.fill.fore_color.rgb = rgb(C["bg"])
        tp = master.part.part_related_by(RT.THEME)
        xml = tp.blob.decode("utf-8")
        xml = re.sub(r'(<a:latin typeface=")[^"]*(")', r"\1Arial\2", xml)
        tp._blob = xml.encode("utf-8")
        keep = None
        for lay in list(prs.slide_layouts):
            if lay.name == "Title Only":
                keep = lay
        for lay in list(prs.slide_layouts):
            if lay is not keep:
                prs.slide_layouts.remove(lay)
        self.layout = keep
        cp = prs.core_properties
        cp.title = "Signals and Signal Boxes, Explained"
        cp.author = AUTHOR
        cp.last_modified_by = AUTHOR
        cp.language = LANG
        cp.subject = "Early Locomotion (a Satisfactory mod): explainer deck 2"

    def slide(self, n, title, tag, **kw):
        return SlideCtx(self, n, title, tag, **kw)

    def finalize(self, path):
        for sl in self.prs.slides:
            for el in sl._element.iter(qn("a:rPr"), qn("a:endParaRPr")):
                el.set("lang", LANG)
        self.prs.save(path)


def callout(shapes, text, box, from_pt, to_pt, align="l", size=18, color=None, dot=True):
    """Label with a thin leader line ending in a dot on the part it names. box = (x, y, w, h)."""
    x, y, w, h = box
    g = group(shapes, f"Callout: {text}", f"A label, {text}, with a leader line to the part it names.")
    gs = g.shapes
    textbox(gs, x, y, w, h, text, size, color or C["text"], bold=True, align=align)
    line(gs, from_pt[0], from_pt[1], to_pt[0], to_pt[1], C["muted"], 1.5)
    if dot:
        ellipse(gs, to_pt[0], to_pt[1], 0.12, fill=C["text"], line=C["bg"], line_w=0.75)
    return g
