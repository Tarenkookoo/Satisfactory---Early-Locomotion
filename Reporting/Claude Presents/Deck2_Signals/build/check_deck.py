"""Structural checks for the Deck 2 .pptx. Usage: python3 check_deck.py [deck.pptx] [first last]

Prints per slide: chip, word counts (content = caption + labels, excluding title, status chips,
slide number and credit lines), smallest font size, pictures, notes slide, alt-text gaps, and an
estimate of text that may overflow its box (PIL metrics with Liberation Sans, same metrics as Arial).
"""
import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
PATH = sys.argv[1] if len(sys.argv) > 1 and sys.argv[1].endswith(".pptx") else os.path.join(
    HERE, "..", "Signals_and_Signal_Boxes_Explained.pptx")
RANGE = [int(a) for a in sys.argv if a.isdigit()]
TAGW = {"AGREED", "PROPOSED", "PROVISIONAL", "UNDECIDED", "BACKGROUND"}
FONTS = {}


def font(size, bold):
    key = (round(size * 4), bold)
    if key not in FONTS:
        f = "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"
        FONTS[key] = ImageFont.truetype("/usr/share/fonts/truetype/liberation/" + f, int(size * 4))
    return FONTS[key]


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from walk(sh.shapes)


def est_lines(par_runs, width_pt):
    """Greedy word wrap using PIL metrics. par_runs: list of (text, size, bold)."""
    words = []
    for text, size, bold in par_runs:
        for i, w in enumerate(text.split(" ")):
            words.append((w, size, bold))
    lines, cur = 1, 0.0
    for w, size, bold in words:
        f = font(size, bold)
        wl = f.getlength(w + " ") / 4
        if cur + f.getlength(w) / 4 > width_pt and cur > 0:
            lines += 1
            cur = 0
        cur += wl
    return lines


def main():
    prs = Presentation(PATH)
    print(f"{os.path.basename(PATH)}: {len(prs.slides)} slides, {os.path.getsize(PATH) / 1e6:.2f} MB")
    print("title:", prs.core_properties.title, "| author:", prs.core_properties.author)
    problems = []
    for i, sl in enumerate(prs.slides, 1):
        if RANGE and not (RANGE[0] <= i <= RANGE[-1]):
            continue
        chip_txt, content, allw, sizes, pics, noalt, over = [], 0, 0, [], 0, [], []
        title = sl.shapes.title.text_frame.text if sl.shapes.title is not None else "(no title)"
        for sh in walk(sl.shapes):
            if sh.shape_type == 13:
                pics += 1
                if not sh._element.xpath(".//p:cNvPr")[0].get("descr"):
                    noalt.append(sh.name)
            if sh.shape_type == 6 and not sh._element.xpath(".//p:cNvPr")[0].get("descr"):
                noalt.append("group:" + sh.name)
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            txt = sh.text_frame.text.strip()
            isTitle = sh == sl.shapes.title
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size is not None:
                        sizes.append(r.font.size.pt)
            n = len(txt.split())
            if sh.name.startswith("Chip:") and txt.upper().split()[0] in TAGW:
                chip_txt.append(txt)
                continue
            if sh.name.startswith("Chip:"):
                chip_txt.append("inline:" + txt[:24])
            allw += n
            if isTitle or sh.name in ("Slide number",) or sh.name.startswith("Credit"):
                continue
            content += n
            # overflow estimate
            tf = sh.text_frame
            w_pt = (sh.width - tf.margin_left - tf.margin_right) / 12700
            h_pt = (sh.height - tf.margin_top - tf.margin_bottom) / 12700
            if not tf.word_wrap:
                continue
            need = 0
            for p in tf.paragraphs:
                runs = [(r.text, r.font.size.pt if r.font.size else 18, bool(r.font.bold)) for r in p.runs]
                if not runs:
                    continue
                sz = max(r[1] for r in runs)
                need += est_lines(runs, w_pt) * sz * 1.17
            if need > h_pt + 1.5:
                over.append(f"{sh.name!r}: needs {need:.0f}pt, has {h_pt:.0f}pt")
        notes = sl.has_notes_slide
        status = "OK"
        if min(sizes or [99]) < 10 or notes or noalt or over:
            status = "CHECK"
        print(f"{i:2d} {title[:34]:34s} chips={chip_txt} content={content:2d} all={allw:2d} "
              f"minpt={min(sizes or [0]):g} pics={pics} notes={notes} {status}")
        for m in noalt:
            print("     no alt text:", m)
        for m in over:
            print("     overflow?:", m)
        if sizes:
            small = sorted({s for s in sizes if s < 14})
            if small:
                print("     sizes below 14 pt:", small)
        if content > 25:
            print(f"     NOTE: {content} content words (target 25, hard cap 45)")
        if content > 45:
            problems.append(i)
    if problems:
        print("OVER HARD CAP:", problems)


if __name__ == "__main__":
    main()
