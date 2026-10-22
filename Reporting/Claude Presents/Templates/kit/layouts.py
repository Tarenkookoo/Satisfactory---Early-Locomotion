"""layouts: builds the template package (theme, master, 13 custom layouts) for el_kit / build_template.

Every layout has real placeholders (title, status chip, caption bar, picture, table, body, slide number).
Panels, rules and other fixed decoration are ordinary shapes drawn on the layout behind the placeholders,
so a slide only holds the content you type. No layout carries a transition, animation or speaker note.

Placeholder idx map (stable across layouts, so changing a slide's layout keeps its text):
   0 title / centre title      1 subtitle        12 slide number
  13 status chip               14 caption bar    15 credit line      16 picture
  17 body / notes / next steps 18 note line      19 chapter label    20 table
  21..28 cards, stops or tiles (numbered)
"""
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.parts.slide import SlideLayoutPart

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import el_kit as k  # noqa: E402
from el_kit import E, _lvl, sldnum_para  # noqa: E402

LAYOUT_NAMES = [
    "Cover",
    "Chapter divider",
    "Title + visual + caption",
    "Picture + caption + credit line",
    "Two-panel compare",
    "Table",
    "Card grid (4)",
    "Card grid (8)",
    "Horizontal timeline",
    "Key-facts tiles",
    "Status legend",
    "Credits / sources table",
    "Closing / what next",
]


# ------------------------------------------------------------------ raw XML helpers
def _geom_round(radius, w, h):
    return ('<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val %d"/></a:avLst></a:prstGeom>'
            % int(min(0.5, radius / min(w, h)) * 100000))


_GEOM_RECT = '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'


def _scheme_fill(name):
    return '<a:solidFill><a:schemeClr val="%s"/></a:solidFill>' % name


def _scheme_ln(name, w_pt=1.5, dash=None):
    d = '<a:prstDash val="%s"/>' % dash if dash else ""
    return ('<a:ln w="%d"><a:solidFill><a:schemeClr val="%s"/></a:solidFill>%s</a:ln>'
            % (int(w_pt * 12700), name, d))


def _prompt(paras):
    """paras: list of (lvl, text)."""
    out = ""
    for lvl, text in paras:
        ppr = '<a:pPr lvl="%d"/>' % lvl if lvl else ""
        out += '<a:p>%s<a:r><a:rPr lang="%s"/><a:t>%s</a:t></a:r></a:p>' % (ppr, k.LANG, text)
    return out


def ph_sp(sid, name, ph_type, idx=None, *, box=None, anchor="t", ins=(0, 0, 0, 0), lst="", prompt=(),
          geom="", fill="", ln="", quarter=True):
    """One placeholder shape as XML. `box` (x, y, w, h in inches) is written so the layout is explicit."""
    attrs = 'type="%s"' % ph_type
    if quarter and ph_type not in ("title", "ctrTitle", "subTitle"):
        attrs += ' sz="quarter"'
    if idx is not None:
        attrs += ' idx="%d"' % idx
    sppr = ""
    if box:
        sppr += '<a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>' % tuple(E(v) for v in box)
    sppr += geom + fill + ln
    body = _prompt(prompt) if prompt else "<a:p><a:endParaRPr lang=\"%s\"/></a:p>" % k.LANG
    return (
        '<p:sp %s><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        '<p:nvPr><p:ph %s/></p:nvPr></p:nvSpPr><p:spPr>%s</p:spPr><p:txBody>'
        '<a:bodyPr vert="horz" wrap="square" lIns="%d" tIns="%d" rIns="%d" bIns="%d" rtlCol="0" anchor="%s">'
        '<a:noAutofit/></a:bodyPr><a:lstStyle>%s</a:lstStyle>%s</p:txBody></p:sp>'
        % (k._NS, sid, name, attrs, sppr, E(ins[0]), E(ins[1]), E(ins[2]), E(ins[3]), anchor, lst, body)
    )


def lv(n, sz, **kw):
    return _lvl(n, sz, **kw)


# ------------------------------------------------------------------ package assembly
class Builder:
    """Adds layouts to a presentation whose master has been installed."""

    def __init__(self, prs):
        self.prs = prs
        self.master = prs.slide_master
        self.next_lid = 2147483649

    def new_layout(self, name):
        pkg = self.prs.part.package
        partname = pkg.next_partname("/ppt/slideLayouts/slideLayout%d.xml")
        xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<p:sldLayout %s preserve="1" userDrawn="1"><p:cSld name="%s"><p:spTree><p:nvGrpSpPr>'
            '<p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="0" y="0"/>'
            '<a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree>'
            '</p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>' % (k._NS, name)
        )
        part = SlideLayoutPart.load(partname, CT.PML_SLIDE_LAYOUT, pkg, xml.encode("utf8"))
        part.relate_to(self.master.part, RT.SLIDE_MASTER)
        rid = self.master.part.relate_to(part, RT.SLIDE_LAYOUT)
        lst = self.master._element.find(qn("p:sldLayoutIdLst"))
        el = etree.SubElement(lst, qn("p:sldLayoutId"))
        el.set("id", str(self.next_lid))
        el.set(qn("r:id"), rid)
        self.next_lid += 1
        return Layout(part.slide_layout)


class Layout:
    def __init__(self, layout):
        self.layout = layout
        self.tree = layout._element.cSld.spTree

    @property
    def sh(self):
        return k.layout_shapes(self.layout)

    def _next_id(self):
        return max(int(i) for i in self.tree.xpath("//p:cNvPr/@id")) + 1

    def add(self, xml):
        self.tree.append(parse_xml(xml))

    # --- standard placeholders
    def title(self, box=k.TITLE_BOX, size=36, kind="title", anchor="ctr", prompt="Click to add the title",
              bold=True, color="tx1"):
        lst = lv(1, size, bold=bold, color=color, font="+mj-lt", line_pct=100 if size <= 40 else 90)
        self.add(ph_sp(self._next_id(), "Title", kind, None, box=box, anchor=anchor, lst=lst,
                       prompt=[(0, prompt)]))

    def chip(self):
        # No geometry or fill here on purpose: a layout placeholder that carries its own shape geometry gets its
        # prompt text drawn onto slides by LibreOffice. set_status() writes the rounded shape on the slide chip.
        lst = lv(1, 14, bold=True, color="tx2", algn="ctr")
        self.add(ph_sp(self._next_id(), "Status chip", "body", 13, box=k.CHIP_BOX, anchor="ctr",
                       ins=(0.06, 0, 0.06, 0), lst=lst, prompt=[(0, "STATUS")]))

    def caption(self):
        lst = lv(1, 20, color="tx1")
        self.add(ph_sp(self._next_id(), "Caption", "body", 14, box=k.CAPTION_BAR, anchor="ctr",
                       ins=(0.3, 0.04, 0.2, 0.04), lst=lst, prompt=[(0, "Caption: one idea, one or two lines")]))

    def credit(self, box):
        lst = lv(1, 10, color="tx2")
        self.add(ph_sp(self._next_id(), "Credit line", "body", 15, box=box, anchor="t", lst=lst,
                       prompt=[(0, "Credit line: author, licence, source")]))

    def number(self):
        self.add(ph_sp(self._next_id(), "Slide number", "sldNum", 12, box=None, anchor="ctr",
                       prompt=(), lst="").replace(
            "<a:p><a:endParaRPr lang=\"%s\"/></a:p>" % k.LANG, sldnum_para()))

    def body(self, name, idx, box, *, levels, prompt, anchor="t", ins=(0, 0, 0, 0), ph_type="body"):
        lst = "".join(levels)
        self.add(ph_sp(self._next_id(), name, ph_type, idx, box=box, anchor=anchor, ins=ins, lst=lst,
                       prompt=prompt))

    def caption_decor(self):
        k.caption_bar_decor(self.sh)


# ------------------------------------------------------------------ the 13 layouts
def build_cover(b):
    L = b.new_layout("Cover")
    sh = L.sh
    k.track_motif(sh, 10.35, 0, 2.3, k.SLIDE_H, vertical=True)
    k.rect(sh, 0.6, 4.5, 1.8, 0.08, k.ORANGE, None, name="Cover rule")
    L.title(box=(0.6, 1.1, 8.9, 3.15), size=54, kind="ctrTitle", anchor="b", prompt="Click to add the deck title")
    L.body("Subtitle", 1, (0.6, 4.8, 8.9, 1.1), levels=[lv(1, 24, color="tx1")], ph_type="subTitle",
           prompt=[(0, "Click to add the subtitle")])
    L.body("Cover note", 18, (0.6, 6.5, 8.9, 0.4), levels=[lv(1, 14, color="tx2")],
           prompt=[(0, "Version, date and author line")], anchor="ctr")


def build_chapter(b):
    L = b.new_layout("Chapter divider")
    sh = L.sh
    k.track_motif(sh, 0, 6.45, k.SLIDE_W, 0.62, vertical=False)
    k.rect(sh, 0.6, 2.05, 1.8, 0.08, k.ORANGE, None, name="Chapter rule")
    L.title(box=(0.6, 2.85, 11.5, 1.5), size=48, anchor="t", prompt="Click to add the chapter title")
    L.body("Chapter label", 19, (0.6, 2.3, 10.5, 0.5), levels=[lv(1, 24, bold=True, color="accent1", caps=True)],
           prompt=[(0, "Part number or label")], anchor="ctr")
    L.body("Summary", 17, (0.6, 4.5, 10.5, 1.2), levels=[lv(1, 20, color="tx2")],
           prompt=[(0, "One line saying what this part covers")])
    L.number()


def build_workhorse(b):
    L = b.new_layout("Title + visual + caption")
    L.caption_decor()
    L.title()
    L.chip()
    L.caption()
    L.number()


def build_picture(b):
    L = b.new_layout("Picture + caption + credit line")
    sh = L.sh
    L.caption_decor()
    k.rect(sh, 8.9, 1.45, 3.83, 4.5, k.PANEL, None, radius=0.08, name="Notes panel")
    k.rect(sh, 8.96, 1.57, 0.08, 4.26, k.ORANGE, None, name="Notes rule")
    L.title()
    L.chip()
    pic_box = (0.6, 1.45, 8.0, 4.5)
    L.add(ph_sp(L._next_id(), "Picture", "pic", 16, box=pic_box, anchor="ctr", quarter=True,
                lst=lv(1, 16, color="tx2", algn="ctr"), prompt=[(0, "Click the icon to add a picture")],
                geom=_geom_round(0.08, 8.0, 4.5), fill=_scheme_fill("bg2")))
    L.credit((0.6, 5.98, 8.0, 0.24))
    L.body("Notes", 17, (9.3, 1.65, 3.25, 4.1),
           levels=[lv(1, 18, marL=0.3, indent=-0.3, bullet="dot", spc_bef=6),
                   lv(2, 16, marL=0.6, indent=-0.3, bullet="dot", color="tx2", spc_bef=3)],
           prompt=[(0, "Short point to notice"), (0, "Another short point")])
    L.caption()
    L.number()


def build_compare(b):
    L = b.new_layout("Two-panel compare")
    sh = L.sh
    L.caption_decor()
    for i, x in enumerate((0.6, 6.76)):
        k.rect(sh, x, 1.45, 5.97, 4.7, k.PANEL, None, radius=0.08, name="Panel %d" % (i + 1))
        k.rect(sh, x + 0.25, 2.3, 1.2, 0.06, k.ORANGE, None, name="Panel %d rule" % (i + 1))
    L.title()
    L.chip()
    head = [lv(1, 20, bold=True, color="tx1")]
    bodyl = [lv(1, 18, marL=0.3, indent=-0.3, bullet="dot", spc_bef=6),
             lv(2, 16, marL=0.6, indent=-0.3, bullet="dot", color="tx2", spc_bef=3)]
    L.body("Left heading", 21, (0.85, 1.6, 5.47, 0.6), levels=head, prompt=[(0, "Left heading")], anchor="ctr")
    L.body("Left body", 22, (0.85, 2.55, 5.47, 3.4), levels=bodyl, prompt=[(0, "Short point"), (0, "Short point")])
    L.body("Right heading", 23, (7.01, 1.6, 5.47, 0.6), levels=head, prompt=[(0, "Right heading")], anchor="ctr")
    L.body("Right body", 24, (7.01, 2.55, 5.47, 3.4), levels=bodyl, prompt=[(0, "Short point"), (0, "Short point")])
    L.caption()
    L.number()


def build_table(b):
    L = b.new_layout("Table")
    L.caption_decor()
    L.title()
    L.chip()
    L.add(ph_sp(L._next_id(), "Table", "tbl", 20, box=(0.6, 1.45, 12.13, 4.6), anchor="ctr",
                lst=lv(1, 16, color="tx2", algn="ctr"), prompt=[(0, "Click the icon to add a table")],
                fill=_scheme_fill("bg2")))
    L.caption()
    L.number()


def _cards(b, name, cols, rows, hsize, bsize):
    L = b.new_layout(name)
    sh = L.sh
    L.caption_decor()
    gap = 0.2
    w = (12.13 - gap * (cols - 1)) / cols
    h = (4.7 - gap * (rows - 1)) / rows
    cells = []
    for r in range(rows):
        for c in range(cols):
            x, y = 0.6 + c * (w + gap), 1.45 + r * (h + gap)
            n = r * cols + c + 1
            k.rect(sh, x, y, w, h, k.PANEL, None, radius=0.08, name="Card %d panel" % n)
            k.rect(sh, x + 0.06, y + 0.12, 0.08, h - 0.24, k.ORANGE, None, name="Card %d rule" % n)
            cells.append((n, x, y))
    L.title()
    L.chip()
    levels = [lv(1, hsize, bold=True, color="tx1"), lv(2, bsize, color="tx2", spc_bef=6)]
    for n, x, y in cells:
        L.body("Card %d" % n, 20 + n, (x + 0.32, y + 0.15, w - 0.55, h - 0.3), levels=levels,
               prompt=[(0, "Card %d heading" % n), (1, "Detail line")])
    L.caption()
    L.number()


def build_timeline(b):
    L = b.new_layout("Horizontal timeline")
    sh = L.sh
    L.caption_decor()
    n, ax_y, x0, w = 5, 3.75, 0.6, 12.13
    k.rect(sh, x0, ax_y - 0.03, w, 0.06, k.STEEL, None, name="Timeline axis")
    centres = [x0 + w * (2 * i + 1) / (2 * n) for i in range(n)]
    for i, cx in enumerate(centres):
        above = i % 2 == 0
        if above:
            k.line(sh, cx, ax_y - 0.15, cx, ax_y - 0.45, k.LINE, 1.5, name="Stem %d" % (i + 1))
        else:
            k.line(sh, cx, ax_y + 0.15, cx, ax_y + 0.45, k.LINE, 1.5, name="Stem %d" % (i + 1))
        k.oval(sh, cx - 0.15, ax_y - 0.15, 0.3, 0.3, k.ORANGE, k.BG, 3.0, name="Node %d" % (i + 1))
    L.title()
    L.chip()
    for i, cx in enumerate(centres):
        above = i % 2 == 0
        levels = [lv(1, 18, bold=True, color="tx1", algn="ctr"), lv(2, 14, color="tx2", algn="ctr", spc_bef=3)]
        if above:
            box = (cx - 1.2, ax_y - 0.45 - 1.75, 2.4, 1.75)
        else:
            box = (cx - 1.2, ax_y + 0.45, 2.4, 1.75)
        L.body("Stop %d" % (i + 1), 21 + i, box, levels=levels, anchor="b" if above else "t",
               prompt=[(0, "Stop %d" % (i + 1)), (1, "Short detail")])
    L.caption()
    L.number()


def build_tiles(b):
    L = b.new_layout("Key-facts tiles")
    sh = L.sh
    L.caption_decor()
    n, gap = 4, 0.2
    w = (12.13 - gap * (n - 1)) / n
    y, h = 1.75, 3.9
    for i in range(n):
        x = 0.6 + i * (w + gap)
        k.rect(sh, x, y, w, h, k.PANEL, None, radius=0.08, name="Tile %d panel" % (i + 1))
        k.rect(sh, x + 0.25, y, w - 0.5, 0.08, k.ORANGE, None, name="Tile %d rule" % (i + 1))
    L.title()
    L.chip()
    levels = [lv(1, 48, bold=True, color="accent1", algn="ctr"), lv(2, 18, color="tx1", algn="ctr", spc_bef=6)]
    for i in range(n):
        x = 0.6 + i * (w + gap)
        L.body("Tile %d" % (i + 1), 21 + i, (x + 0.15, y + 0.35, w - 0.3, h - 0.6), levels=levels, anchor="ctr",
               prompt=[(0, "00"), (1, "Short label")])
    L.caption()
    L.number()


def build_legend(b):
    L = b.new_layout("Status legend")
    sh = L.sh
    L.caption_decor()
    for i, kind in enumerate(k.CHIPS):
        y = 1.45 + i * 0.98
        k.rect(sh, 0.6, y, 8.0, 0.78, k.PANEL, None, radius=0.08, name="Legend row %d" % (i + 1))
        k.chip(sh, kind, 0.85, y + 0.19, name="Chip: " + kind)
        k.label(sh, 2.85, y, 5.6, 0.78, k.CHIP_MEANING[kind], 20, k.TEXT, align="l", anchor="m",
                name="Legend meaning: " + kind, margins=(0, 0, 0, 0))
    # corner sketch: where the chip lives on every content slide
    g = k.group(sh, "Corner sketch", "Small drawing of a slide with a tag chip in its top-right corner")
    k.rect(g.shapes, 9.2, 1.7, 3.53, 1.99, k.BG, k.STEEL, 2.0, radius=0.08, name="sketch slide")
    k.rect(g.shapes, 9.4, 1.95, 1.7, 0.16, k.LINE, None, radius=0.05, name="sketch title")
    k.rect(g.shapes, 11.35, 1.9, 1.2, 0.26, k.ORANGE, None, radius=0.1, name="sketch chip")
    k.rect(g.shapes, 9.4, 3.3, 3.13, 0.24, k.PANEL, None, radius=0.05, name="sketch caption")
    k.label(sh, 9.2, 3.82, 3.53, 0.35, "Top-right corner of every content slide", 14, k.MUTED, align="l",
            anchor="t", name="Legend note 1", margins=(0, 0, 0, 0))
    k.label(sh, 9.2, 4.55, 3.53, 0.9, "A small chip beside a statement tags that statement.", 14, k.MUTED,
            align="l", anchor="t", name="Legend note 2", margins=(0, 0, 0, 0))
    L.title(box=k.TITLE_BOX_WIDE)
    L.caption()
    L.number()


def build_credits(b):
    L = b.new_layout("Credits / sources table")
    L.caption_decor()
    L.title(box=k.TITLE_BOX_WIDE)
    L.add(ph_sp(L._next_id(), "Table", "tbl", 20, box=(0.6, 1.45, 12.13, 3.85), anchor="ctr",
                lst=lv(1, 16, color="tx2", algn="ctr"), prompt=[(0, "Click the icon to add the credits table")],
                fill=_scheme_fill("bg2")))
    L.body("Note", 17, (0.6, 5.45, 12.13, 0.65), levels=[lv(1, 14, color="tx2")],
           prompt=[(0, "Licence note, for example how images were changed")])
    L.caption()
    L.number()


def build_closing(b):
    L = b.new_layout("Closing / what next")
    sh = L.sh
    k.track_motif(sh, 10.35, 0, 2.3, k.SLIDE_H, vertical=True)
    k.rect(sh, 0.6, 2.2, 1.8, 0.08, k.ORANGE, None, name="Closing rule")
    L.title(box=(0.6, 0.8, 9.0, 1.3), size=44, anchor="b", prompt="Click to add the closing title")
    L.body("Next steps", 17, (0.6, 2.6, 9.0, 3.4),
           levels=[lv(1, 24, marL=0.5, indent=-0.5, bullet="num", spc_bef=10),
                   lv(2, 18, marL=0.5, indent=0, color="tx2", spc_bef=3)],
           prompt=[(0, "First next step"), (0, "Second next step"), (0, "Third next step")])
    L.body("Closing note", 18, (0.6, 6.5, 9.0, 0.4), levels=[lv(1, 14, color="tx2")],
           prompt=[(0, "Contact, date or version line")], anchor="ctr")


def build_all(b):
    build_cover(b)
    build_chapter(b)
    build_workhorse(b)
    build_picture(b)
    build_compare(b)
    build_table(b)
    _cards(b, "Card grid (4)", 2, 2, 20, 16)
    _cards(b, "Card grid (8)", 4, 2, 18, 14)
    build_timeline(b)
    build_tiles(b)
    build_legend(b)
    build_credits(b)
    build_closing(b)


def new_template_presentation():
    """A presentation holding the theme, master and the 13 layouts, and no slides."""
    prs = Presentation()
    prs.slide_width = k.Emu(12192000)
    prs.slide_height = k.Emu(6858000)
    for lay in list(prs.slide_layouts):
        prs.slide_layouts.remove(lay)
    master = prs.slide_master
    theme_part = master.part.part_related_by(RT.THEME)
    theme_part._blob = k.theme_xml().encode("utf8")
    m = master._element
    new = parse_xml(k.master_xml())
    for ch in list(m):
        m.remove(ch)
    for ch in list(new):
        m.append(ch)
    pres = prs.part._element
    old = pres.find(qn("p:defaultTextStyle"))
    old.getparent().replace(old, parse_xml(k.default_text_style_xml()))
    build_all(Builder(prs))
    k.tidy_package(prs)
    return prs
