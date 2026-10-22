"""el_kit: one shape and template library for the Early Locomotion house style (python-pptx).

Merges the best of the Deck 3, Deck 5 and Deck 7 libraries (lib.py, lib.py, el_lib.py) and adds the
template-aware slide API. Native shapes only. Arial. British English.

Contents (search for the banner comments)
  1. Tokens          colours, type sizes, fixed geometry, chip and pill styles
  2. Text measuring  text_w, wrap_lines, wrap_count, balanced_width (Liberation Sans = Arial metrics)
  3. Styling         rgb, shade, style, set_alt, text fill
  4. Primitives      rect, oval, label, poly, line, arrow, path, curve, group, box
  5. Chips, pills    chip, ichip, pill (report workflow states), legend helpers
  6. Pictures        add_picture_fit, credit, picture_box (grey "Picture goes here" box)
  7. Tables          cell_border, set_cell, add_table
  8. Composites      caption_bar, track_motif, card, tile, timeline, flow_row, callout
  9. Theme           theme_xml, master_xml, default_text_style_xml (used by build_template.py)
 10. Template API    open_template, new_deck, add_slide, set_status, set_caption, fill_text,
                     fill_picture, fill_table, finish, clone_shape, clear_timing
 11. Package tidy    tidy_package, set_core_properties

All coordinates are inches on a 13.333 x 7.5 in canvas. Helpers take `sh`, a shapes collection
(slide.shapes, a group's .shapes, or a layout wrapped by layout_shapes()), so the same call
works on a slide, inside a group, or on a layout.

Type floor: nothing below 14 pt, except credit lines (10 pt) and credits tables (11 pt).
"""
import copy
import math
import os
import re
import uuid
import zipfile

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.shapes.shapetree import SlideShapes
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_DIR = os.path.abspath(os.path.join(HERE, ".."))
TEMPLATE_PATH = os.path.join(TEMPLATE_DIR, "EL_Claude_Template_v001.pptx")

# ============================================================================ 1. Tokens
BG = "#1E2126"          # master background; theme dk1
PANEL = "#2B2F36"       # panels, caption bar; theme dk2
PANEL_ALT = "#343A43"
LINE = "#4A515B"
TEXT = "#F2F2F2"       # theme lt1
MUTED = "#A9B0B8"      # theme lt2
ORANGE = "#F28C28"     # theme accent1
ORANGE_DK = "#C96F1B"
GREEN = "#3FB950"      # theme accent2
WATER = "#4DA3FF"      # theme accent3
PURPLE = "#B084F5"     # theme accent4
AMBER = "#F2C14E"      # theme accent5
RED = "#E5484D"        # theme accent6
ELEC = "#6CA0DC"
COAL = "#1B1D21"
COAL_LINE = "#6B7078"
STEEL = "#9AA0A6"
SLEEPER = "#6B5B4B"
DARKMETAL = "#3A3F47"
WAGON = "#6C7480"
DARK = "#101418"       # text on bright fills
HLINK = "#8FC7FF"      # theme hlink (light, readable on charcoal)
FOLHLINK = "#C9B2F5"   # theme folHlink

FONT = "Arial"
LANG = "en-GB"
AUTHOR = "Claude (Early Locomotion project)"

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
RIGHT_EDGE = SLIDE_W - MARGIN                  # 12.733
TITLE_BOX = (0.6, 0.4, 10.2, 0.8)              # 36 pt bold, anchored middle
TITLE_BOX_WIDE = (0.6, 0.4, 12.13, 0.8)        # slides without a status chip
CHIP_BOX = (11.03, 0.58, 1.7, 0.4)             # top-right status chip, 14 pt bold
CAPTION_BAR = (0.6, 6.3, 12.13, 0.8)           # 20 pt caption inside, orange rule on the left
CAPTION_RULE = (0.6, 6.36, 0.08, 0.68)
CAPTION_TEXT_X = 0.9
NUMBER_BOX = (11.73, 7.14, 1.0, 0.3)           # 14 pt muted, right aligned
CONTENT_TOP, CONTENT_BOTTOM = 1.45, 6.15

SZ_TITLE, SZ_CAPTION, SZ_LABEL, SZ_CREDIT, SZ_CREDITS_TABLE = 36, 20, 14, 10, 11
SZ_FLOOR = 14

# Design status chips (fill, text, outline). Colour is never the only signal: the word is always shown.
CHIPS = {
    "Agreed": (GREEN, DARK, None),
    "Proposed": (PURPLE, DARK, None),
    "Provisional": (WATER, DARK, None),
    "Undecided": (None, TEXT, MUTED),
    "Background": (LINE, TEXT, None),
}
CHIP_MEANING = {
    "Agreed": "Decided.",
    "Proposed": "Suggested, not confirmed.",
    "Provisional": "A test number that may change.",
    "Undecided": "Still an open question.",
    "Background": "A real-world fact.",
}

# Report workflow pills: NOT design statuses. Outline pill with a marker shape, so they cannot be mistaken for chips.
PILLS = {
    "Complete": GREEN,
    "In progress": AMBER,
    "Blocked": RED,
    "Planned": STEEL,
}


# ============================================================================ 2. Text measuring
_FONT_R = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
_FONT_B = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
_fcache = {}


def text_w(s, size, bold=False):
    """Width in inches of `s` at `size` pt (Liberation Sans has Arial's metrics)."""
    if bold not in _fcache:
        _fcache[bold] = ImageFont.truetype(_FONT_B if bold else _FONT_R, 200)
    return _fcache[bold].getlength(s) * size / 200.0 / 72.0


def wrap_lines(s, size, width, bold=False):
    """Greedy word wrap, as PowerPoint does it. Returns the list of lines."""
    out = []
    for para in s.split("\n"):
        cur = ""
        for w in para.split(" "):
            t = (cur + " " + w).strip()
            if text_w(t, size, bold) <= width or not cur:
                cur = t
            else:
                out.append(cur)
                cur = w
        out.append(cur)
    return out


def wrap_count(s, size, width, bold=False):
    return len(wrap_lines(s, size, width, bold))


def balanced_width(text, size, max_w, bold=False, min_w=6.0):
    """Widest box (<= max_w) that wraps `text` into two lines whose second line is not an orphan."""
    first = wrap_lines(text, size, max_w, bold)
    if len(first) != 2 or text_w(first[1], size, bold) >= 0.45 * text_w(first[0], size, bold):
        return max_w
    w = max_w
    while w > min_w:
        w = round(w - 0.1, 2)
        ls = wrap_lines(text, size, w, bold)
        if len(ls) != 2:
            break
        if text_w(ls[1], size, bold) >= 0.45 * text_w(ls[0], size, bold):
            return min(max_w, w + 0.3)
    return max_w


def word_count(t):
    """Whitespace tokens that contain a letter (bare numerals are not words)."""
    return len([w for w in t.split() if any(c.isalpha() for c in w)])


# ============================================================================ 3. Styling
def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def I(v):
    return Inches(v)


def shade(h, f):
    """Lighten (f > 0) or darken (f < 0) a hex colour."""
    h = h.lstrip("#")
    c = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
    c = [int(v + (255 - v) * f) for v in c] if f >= 0 else [int(v * (1 + f)) for v in c]
    return "#%02X%02X%02X" % tuple(c)


def _font(run, size=None, color=None, bold=None, italic=None, name=FONT):
    f = run.font
    if name:
        f.name = name
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    f.language_id = MSO_LANGUAGE_ID.ENGLISH_UK


def _strip_style(shape):
    el = shape._element
    for st in el.findall(qn("p:style")):
        el.remove(st)


def _set_alpha(fill_parent_elm, alpha):
    clr = fill_parent_elm.find(".//" + qn("a:srgbClr"))
    if clr is not None:
        for a in clr.findall(qn("a:alpha")):
            clr.remove(a)
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))


def style(shape, fill=None, line=None, lw=1.5, dash=None, alpha=None):
    """Flat fill and outline; removes the theme style reference so nothing is inherited by surprise."""
    _strip_style(shape)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _set_alpha(shape._element.spPr.find(qn("a:solidFill")), alpha)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(lw)
        if dash:
            shape.line.dash_style = dash


_CNVPR_XPATH = ("./p:nvSpPr/p:cNvPr | ./p:nvGrpSpPr/p:cNvPr | ./p:nvPicPr/p:cNvPr | "
                "./p:nvCxnSpPr/p:cNvPr | ./p:nvGraphicFramePr/p:cNvPr")


def set_alt(shape, descr, name=None):
    """Alt text (and optionally the name used in the Selection Pane) on any shape."""
    c = shape._element.xpath(_CNVPR_XPATH)[0]
    c.set("descr", descr)
    if name:
        c.set("name", name)


def _anchor(a):
    return {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[a]


def _align(a):
    return {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[a]


def fill_text(tf, paras, size=None, color=None, bold=None, align=None, anchor=None, margins=None, spacing=None,
              wrap=True, font=FONT):
    """Write text into a text frame.

    paras: str | list of paragraphs. A paragraph is str, a dict {"text"|"runs", "lvl", "align", + run options},
    or a list of runs. A run is str or (str, {"size","color","bold","italic","link"}).
    Any of size/color/bold/align/anchor/margins left as None is NOT written, so placeholder text keeps
    inheriting from the layout.
    """
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if anchor:
        tf.vertical_anchor = _anchor(anchor)
    if margins:
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [I(m) for m in margins]
    if isinstance(paras, (str, dict)):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = {}
        if isinstance(para, dict):
            opts = para
            runs = para.get("runs") or [para.get("text", "")]
        elif isinstance(para, list):
            runs = para
        else:
            runs = [para]
        al = opts.get("align", align)
        if al:
            p.alignment = _align(al)
        if opts.get("lvl"):
            p.level = opts["lvl"]
        if spacing:
            p.line_spacing = spacing
        for rn in runs:
            t, o = rn if isinstance(rn, tuple) else (rn, {})
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", opts.get("size", size)), o.get("color", opts.get("color", color)),
                  o.get("bold", opts.get("bold", bold)), o.get("italic", opts.get("italic")), name=font)
            if o.get("link"):
                r.hyperlink.address = o["link"]


# ============================================================================ 4. Primitives
def rect(sh, x, y, w, h, fill=None, line=None, lw=1.5, radius=None, dash=None, alpha=None, name=None,
         text=None, size=14, color=TEXT, bold=False, align="c", anchor="m", margins=(0.06, 0.03, 0.06, 0.03)):
    """Rectangle (rounded when `radius` in inches is given), optionally holding text."""
    shp = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    if radius:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    if text is not None:
        fill_text(shp.text_frame, text, size, color, bold, align, anchor, margins)
    return shp


def box(sh, x, y, w, h, text, size=14, fill=PANEL, line=LINE, lw=1.5, color=TEXT, bold=False, radius=0.08,
        dash=None, align="c", anchor="m", name=None):
    """Labelled box: the basic diagram node."""
    return rect(sh, x, y, w, h, fill, line, lw, radius, dash, name=name, text=(text or None), size=size,
                color=color, bold=bold, align=align, anchor=anchor, margins=(0.1, 0.05, 0.1, 0.05))


def oval(sh, x, y, w, h, fill=None, line=None, lw=1.5, alpha=None, name=None, text=None, size=14, color=TEXT,
         bold=False, dash=None):
    shp = sh.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(w), I(h))
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    if text is not None:
        fill_text(shp.text_frame, text, size, color, bold, "c", "m", (0, 0, 0, 0))
    return shp


def label(sh, x, y, w, h, text, size=14, color=TEXT, bold=False, align="c", anchor="m", name=None,
          margins=(0.02, 0.01, 0.02, 0.01), spacing=None):
    tb = sh.add_textbox(I(x), I(y), I(w), I(h))
    fill_text(tb.text_frame, text, size, color, bold, align, anchor, margins, spacing)
    if name:
        tb.name = name
    return tb


def poly(sh, pts, fill=None, line=None, lw=1.5, closed=True, alpha=None, name=None, dash=None):
    """Freeform polygon / polyline through `pts` (inches)."""
    e = [(int(Inches(px)), int(Inches(py))) for px, py in pts]
    fb = sh.build_freeform(e[0][0], e[0][1], scale=1.0)
    fb.add_line_segments(e[1:], close=closed)
    shp = fb.convert_to_shape()
    style(shp, fill, line, lw, dash, alpha)
    if name:
        shp.name = name
    return shp


def _ends(conn, head, tail):
    ln = conn.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head), ("a:tailEnd", tail)):
        if on:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "med")
            e.set("len", "med")


def line(sh, x1, y1, x2, y2, color=STEEL, w=2.25, dash=None, head=False, tail=False, name=None):
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = dash
    _ends(c, head, tail)
    if name:
        c.name = name
    return c


def arrow(sh, x1, y1, x2, y2, color=ORANGE, w=2.25, dash=None, name=None):
    """Straight arrow from (x1, y1) to (x2, y2); the head is at the end."""
    return line(sh, x1, y1, x2, y2, color, w, dash, False, True, name)


def path(sh, pts, color=ORANGE, w=2.25, head=True, dash=None, name=None):
    """Open polyline (elbow arrows)."""
    shp = poly(sh, pts, None, color, w, closed=False, name=name, dash=dash)
    if head:
        _ends(shp, False, True)
    return shp


def curve(sh, p0, p1, p2, color=ORANGE, w=2.25, n=28, head=True, dash=None, name=None):
    """Quadratic Bezier drawn as a polyline from p0 to p2, bending towards p1."""
    pts = []
    for i in range(n + 1):
        t = i / n
        pts.append(((1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * p1[0] + t * t * p2[0],
                    (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * p1[1] + t * t * p2[1]))
    return path(sh, pts, color, w, head, dash, name)


def group(sh, name, alt):
    """Group with alt text. Every diagram is one group (or a few), and every group needs alt text."""
    g = sh.add_group_shape()
    g.name = name
    set_alt(g, alt, name)
    return g


# ============================================================================ 5. Chips and pills
def chip(sh, kind, x, y, w=1.7, h=0.4, size=14, text=None, name="Chip", align="c"):
    """Design-status chip: a filled rounded rectangle (Undecided is outline only). The word is always shown."""
    fill, tcol, outline = CHIPS[kind]
    t = text if text is not None else kind.upper()
    return rect(sh, x, y, w, h, fill, outline, 2.0, radius=min(0.12, h / 2), name=name, text=t, size=size,
                color=tcol, bold=True, align=align, margins=(0.06, 0.0, 0.06, 0.0))


def ichip(sh, kind, x, y, text=None, w=None, h=0.34):
    """Inline chip beside a single statement (14 pt). Gives its own words; width follows the text."""
    t = text if text is not None else kind.upper()
    if w is None:
        w = text_w(t, 14, True) + 0.34
    return chip(sh, kind, x, y, w, h, 14, t, name="Inline chip: " + kind)


def pill(sh, state, x, y, w=1.9, h=0.4):
    """Report workflow pill (Complete / In progress / Blocked / Planned). Not a design status.

    Drawn as an outline pill with a marker shape and the state in words, so it cannot be mistaken for a
    design chip: chips are filled blocks with upper-case words; pills are open outlines with a marker.
    Returns a group (with alt text).
    """
    col = PILLS[state]
    g = group(sh, "Workflow pill: " + state, "Report workflow state: " + state + ". Not a design status.")
    gs = g.shapes
    dash = MSO_LINE.DASH if state == "Planned" else None
    rect(gs, x, y, w, h, None, col, 2.0, radius=h / 2, dash=dash, name="pill outline")
    d = h * 0.5
    mx, my = x + h * 0.5, y + h / 2
    if state == "Complete":
        oval(gs, mx - d / 2, my - d / 2, d, d, col, None, name="marker")
        poly(gs, [(mx - d * 0.25, my), (mx - d * 0.05, my + d * 0.2), (mx + d * 0.27, my - d * 0.2)], None, DARK,
             1.75, closed=False, name="marker tick")
    elif state == "In progress":
        oval(gs, mx - d / 2, my - d / 2, d, d, None, col, 2.0, name="marker")
        for k in (-1, 0, 1):
            oval(gs, mx + k * d * 0.26 - d * 0.07, my - d * 0.07, d * 0.14, d * 0.14, col, None, name="marker dot")
    elif state == "Blocked":
        oval(gs, mx - d / 2, my - d / 2, d, d, col, None, name="marker")
        rect(gs, mx - d * 0.3, my - d * 0.08, d * 0.6, d * 0.16, DARK, None, name="marker bar")
    else:  # Planned
        oval(gs, mx - d / 2, my - d / 2, d, d, None, col, 2.0, dash=MSO_LINE.DASH, name="marker")
    label(gs, x + h, y, w - h - 0.12, h, state, 14, TEXT, False, "l", "m", name="pill text")
    return g


# ============================================================================ 6. Pictures
def add_picture_fit(sh, file_path, x, y, w, h, alt, focus=(0.5, 0.5), radius=0.08, name=None):
    """Cover-fit picture: crops (never stretches) to the box, rounded corners, alt text."""
    with Image.open(file_path) as im:
        iw, ih = im.size
    ar_box, ar_img = w / h, iw / ih
    pic = sh.add_picture(file_path, I(x), I(y), I(w), I(h))
    if ar_img > ar_box:
        vis = ar_box / ar_img
        left = (1 - vis) * focus[0]
        pic.crop_left, pic.crop_right = left, 1 - vis - left
    else:
        vis = ar_img / ar_box
        top = (1 - vis) * focus[1]
        pic.crop_top, pic.crop_bottom = top, 1 - vis - top
    round_corners(pic, w, h, radius)
    set_alt(pic, alt, name or os.path.splitext(os.path.basename(file_path))[0])
    return pic


def round_corners(pic, w, h, radius=0.08):
    """Give a picture rounded-rectangle geometry (radius in inches)."""
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = etree.SubElement(pic._element.spPr, qn("a:prstGeom"))
        xfrm = pic._element.spPr.find(qn("a:xfrm"))
        if xfrm is not None:
            xfrm.addnext(geom)
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for g in list(av):
        av.remove(g)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", "val %d" % int(min(0.5, radius / min(w, h)) * 100000))


def credit(sh, x, y, w, text, align="l", h=0.24, color=MUTED):
    """Credit line: 10 pt, directly under its picture. The only text on a content slide below 14 pt."""
    return label(sh, x, y, w, h, text, SZ_CREDIT, color, align=align, anchor="t", name="Credit line",
                 margins=(0, 0, 0, 0))


def picture_box(sh, x, y, w, h, text="Picture goes here", alt="Grey box standing in for a picture",
                name="Picture box"):
    """Labelled grey stand-in for a picture (used in the Patterns deck, which has no photographs)."""
    g = group(sh, name, alt)
    rect(g.shapes, x, y, w, h, PANEL_ALT, LINE, 1.5, radius=0.08, dash=MSO_LINE.DASH, name="picture box outline")
    # simple landscape glyph: sun and hills, so the box reads as a picture
    cx, cy = x + w / 2, y + h / 2 - 0.25
    oval(g.shapes, cx + 0.45, cy - 0.45, 0.3, 0.3, LINE, None, name="glyph sun")
    poly(g.shapes, [(cx - 0.9, cy + 0.5), (cx - 0.35, cy - 0.15), (cx + 0.05, cy + 0.3), (cx + 0.35, cy + 0.0),
                    (cx + 0.9, cy + 0.5)], LINE, None, name="glyph hills")
    label(g.shapes, x, cy + 0.65, w, 0.4, text, 18, MUTED, True, "c", "m", name="picture box label")
    return g


# ============================================================================ 7. Tables
def cell_border(cell, color, w_pt=0.75):
    """All four cell borders in one colour. Border elements must precede the cell fill in a:tcPr."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    first = tcPr[0] if len(tcPr) else None
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(w_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color.lstrip("#"))
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        if first is not None:
            first.addprevious(ln)
        else:
            tcPr.append(ln)


def set_cell(cell, text, size, color=TEXT, bold=False, fill=None, align="l", border=LINE,
             margins=(0.08, 0.04, 0.08, 0.04)):
    """Format one table cell. `text` follows fill_text(); a run may carry {"link": url}."""
    cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom = [I(m) for m in margins]
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text(cell.text_frame, text, size, color, bold, align)
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    else:
        cell.fill.background()
    if border:
        cell_border(cell, border)


def _format_table(tbl, data, col_w, size, header, header_fill, aligns, row_fill, border, total_h):
    rows, cols = len(data), len(data[0])
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1" if header else "0")
    tblPr.set("bandRow", "0")
    for sid in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(sid)
    tot = sum(col_w)
    full = sum(c.width for c in tbl.columns)
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Emu(int(full * cw / tot))
    for r in tbl.rows:
        r.height = Emu(int(I(total_h) / rows))
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            is_head = header and ri == 0
            set_cell(tbl.cell(ri, ci), val, size, TEXT, bold=is_head, fill=header_fill if is_head else row_fill,
                     align=(aligns[ci] if aligns else "l"), border=border)


def add_table(sh, x, y, w, h, data, col_w, alt, size=14, header=True, header_fill=PANEL_ALT, aligns=None,
              row_fill=None, border=LINE, name="Table"):
    """Free-standing table (no placeholder). data = list of rows of cell texts."""
    gf = sh.add_table(len(data), len(data[0]), I(x), I(y), I(w), I(h))
    _format_table(gf.table, data, col_w, size, header, header_fill, aligns, row_fill, border, h)
    set_alt(gf, alt, name)
    return gf


# ============================================================================ 8. Composites
def caption_bar_decor(sh, name="Caption bar"):
    """The panel and orange rule behind a caption (the layouts draw these; the caption text is a placeholder)."""
    x, y, w, h = CAPTION_BAR
    bar = rect(sh, x, y, w, h, PANEL, None, radius=0.08, name=name)
    rx, ry, rw, rh = CAPTION_RULE
    rule = rect(sh, rx, ry, rw, rh, ORANGE, None, name="Caption rule")
    return bar, rule


def caption_bar(sh, text, y=None, h=None, name="Caption"):
    """Free-standing caption bar with text (for slides not built on a layout). 20 pt."""
    x, by, w, bh = CAPTION_BAR
    by = by if y is None else y
    bh = bh if h is None else h
    rect(sh, x, by, w, bh, PANEL, None, radius=0.08, name="Caption bar")
    rect(sh, x, by + 0.06, 0.08, bh - 0.12, ORANGE, None, name="Caption rule")
    cw = balanced_width(text, SZ_CAPTION, 11.4)
    return label(sh, CAPTION_TEXT_X, by, cw, bh, text, SZ_CAPTION, TEXT, align="l", anchor="m", name=name)


def track_motif(sh, x, y, w, h, vertical=True, pitch=0.42):
    """Generic top-down track: two rails and sleepers. A brand motif, not a real location or vehicle."""
    g = group(sh, "Track motif", "Decorative drawing of two rails and sleepers seen from above")
    gs = g.shapes
    if vertical:
        gauge = w * 0.56
        for k in range(int(h / pitch) + 1):
            if y + k * pitch + 0.14 <= y + h:
                rect(gs, x, y + k * pitch, w, 0.14, SLEEPER, None, name="sleeper")
        for cx in (x + w / 2 - gauge / 2, x + w / 2 + gauge / 2):
            rect(gs, cx - 0.045, y, 0.09, h, STEEL, None, name="rail")
    else:
        gauge = h * 0.56
        for k in range(int(w / pitch) + 1):
            if x + k * pitch + 0.14 <= x + w:
                rect(gs, x + k * pitch, y, 0.14, h, SLEEPER, None, name="sleeper")
        for cy in (y + h / 2 - gauge / 2, y + h / 2 + gauge / 2):
            rect(gs, x, cy - 0.045, w, 0.09, STEEL, None, name="rail")
    return g


def card(sh, x, y, w, h, heading, body=None, accent=ORANGE, hsize=18, bsize=14, name="Card"):
    """Panel with an orange left rule, a bold heading and an optional detail line. Returns a group."""
    g = group(sh, name, "Card: " + heading)
    gs = g.shapes
    rect(gs, x, y, w, h, PANEL, None, radius=0.08, name="card panel")
    rect(gs, x + 0.06, y + 0.12, 0.08, h - 0.24, accent, None, name="card rule")
    paras = [{"text": heading, "bold": True, "size": hsize, "color": TEXT}]
    if body:
        paras.append({"text": body, "size": bsize, "color": MUTED})
    label(gs, x + 0.3, y + 0.12, w - 0.5, h - 0.24, paras, bsize, TEXT, False, "l", "t", name="card text",
          margins=(0, 0, 0, 0))
    return g


def tile(sh, x, y, w, h, value, caption_text, accent=ORANGE, vsize=48, csize=18, name="Tile"):
    """Key-fact tile: one big value over a short label. Returns a group."""
    g = group(sh, name, "Key fact: %s, %s" % (value, caption_text))
    gs = g.shapes
    rect(gs, x, y, w, h, PANEL, None, radius=0.08, name="tile panel")
    rect(gs, x + 0.25, y, w - 0.5, 0.08, accent, None, name="tile rule")
    label(gs, x + 0.15, y + 0.3, w - 0.3, h * 0.5, value, vsize, accent, True, "c", "m", name="tile value")
    label(gs, x + 0.2, y + 0.3 + h * 0.5, w - 0.4, h * 0.5 - 0.5, caption_text, csize, TEXT, False, "c", "t",
          name="tile label")
    return g


def timeline(sh, x, y, w, stops, node_d=0.3, label_w=2.4, up_h=1.6, down_h=1.6, hsize=18, dsize=14,
             name="Timeline"):
    """Horizontal timeline. stops = [(heading, detail), ...]; labels alternate above and below the axis.
    (x, y) is the left end of the axis; the group spans up_h above and down_h below. Returns a group."""
    g = group(sh, name, "Timeline of %d stops: %s" % (len(stops), "; ".join(s[0] for s in stops)))
    gs = g.shapes
    rect(gs, x, y - 0.03, w, 0.06, STEEL, None, name="axis")
    n = len(stops)
    for k, (head, detail) in enumerate(stops):
        cx = x + w * (2 * k + 1) / (2 * n)
        above = k % 2 == 0
        oval(gs, cx - node_d / 2, y - node_d / 2, node_d, node_d, ORANGE, BG, 3.0, name="node")
        if above:
            line(gs, cx, y - node_d / 2, cx, y - 0.45, LINE, 1.5, name="stem")
            label(gs, cx - label_w / 2, y - 0.45 - up_h, label_w, up_h,
                  [{"text": head, "bold": True, "size": hsize}, {"text": detail, "size": dsize, "color": MUTED}],
                  dsize, TEXT, False, "c", "b", name="stop text", margins=(0, 0, 0, 0.03))
        else:
            line(gs, cx, y + node_d / 2, cx, y + 0.45, LINE, 1.5, name="stem")
            label(gs, cx - label_w / 2, y + 0.45, label_w, down_h,
                  [{"text": head, "bold": True, "size": hsize}, {"text": detail, "size": dsize, "color": MUTED}],
                  dsize, TEXT, False, "c", "t", name="stop text", margins=(0, 0.03, 0, 0))
    return g


def flow_row(sh, x, y, w, h, labels, gap=0.55, fills=None, color=ORANGE, size=18):
    """Row of boxes joined by arrows. Returns the list of boxes (arrows start and end in the gaps)."""
    n = len(labels)
    bw = (w - gap * (n - 1)) / n
    out = []
    for i, t in enumerate(labels):
        bx = x + i * (bw + gap)
        out.append(box(sh, bx, y, bw, h, t, size, fill=(fills[i] if fills else PANEL), line=LINE, bold=True,
                       name="step"))
        if i < n - 1:
            arrow(sh, bx + bw + 0.06, y + h / 2, bx + bw + gap - 0.06, y + h / 2, color, 2.25, name="flow arrow")
    return out


def callout(sh, bx, by, bw, bh, text, target, side="b", size=14, color=ORANGE):
    """Label box plus a leader that starts at the box edge and points away from it (never across the text).
    side: which edge the leader leaves from: 't', 'b', 'l' or 'r'. target = (x, y) in inches."""
    b = box(sh, bx, by, bw, bh, text, size, fill=PANEL_ALT, line=color, name="callout label")
    if side == "b":
        p = (bx + bw / 2, by + bh)
    elif side == "t":
        p = (bx + bw / 2, by)
    elif side == "l":
        p = (bx, by + bh / 2)
    else:
        p = (bx + bw, by + bh / 2)
    arrow(sh, p[0], p[1], target[0], target[1], color, 2.0, name="leader")
    return b


def wheel(sh, cx, cy, d, fill=COAL, rim=STEEL, lw=1.5):
    oval(sh, cx - d / 2, cy - d / 2, d, d, fill, rim, lw, name="wheel")
    oval(sh, cx - d * 0.13, cy - d * 0.13, d * 0.26, d * 0.26, rim, None, name="hub")


def generic_wagon(sh, x, gy, u=0.2, length=4.4, color=WAGON):
    """Plain box-wagon silhouette (neutral, not any real vehicle). (x, gy) = left edge and rail-top line; u = unit."""
    def rc(a, h0, w, h1, fill, line_=None, radius=None, nm="part"):
        return rect(sh, x + a * u, gy - h1 * u, w * u, (h1 - h0) * u, fill, line_, 0.75, radius=radius, name=nm)
    rc(0, 1.0, length, 1.4, DARKMETAL, STEEL, nm="wagon frame")
    rc(0, 1.4, length, 3.5, color, None, radius=0.05 * u, nm="wagon body")
    rc(-0.1, 3.5, length + 0.2, 3.75, shade(color, -0.3), None, nm="wagon roof")
    for a in (0.9, length - 0.9):
        wheel(sh, x + a * u, gy - 0.6 * u, 1.2 * u, COAL, STEEL, 1.25)


def rails_strip(sh, x, y, length, sleepers=True):
    """Side-on rail line (a thin steel bar), y = top of rail."""
    rect(sh, x, y, length, 0.07, STEEL, None, name="rail")
    if sleepers:
        k = 0
        while k * 0.4 < length:
            rect(sh, x + k * 0.4, y + 0.07, 0.14, 0.09, SLEEPER, None, name="sleeper")
            k += 1


# ============================================================================ 9. Theme
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS = 'xmlns:a="%s" xmlns:r="%s" xmlns:p="%s"' % (_A, _R, _P)


def E(v):
    """Inches to EMU."""
    return int(round(v * 914400))


def theme_xml():
    """Charcoal-and-orange theme: Arial for headings and body, light hyperlink colours, flat format scheme."""
    def c(tag, hexv):
        return '<a:%s><a:srgbClr val="%s"/></a:%s>' % (tag, hexv.lstrip("#"), tag)
    fills = "".join('<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>' for _ in range(3))
    lns = "".join('<a:ln w="%d" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/>'
                  '</a:solidFill><a:prstDash val="solid"/></a:ln>' % w for w in (12700, 19050, 28575))
    effs = "".join("<a:effectStyle><a:effectLst/></a:effectStyle>" for _ in range(3))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<a:theme xmlns:a="%s" name="Early Locomotion Charcoal and Orange"><a:themeElements>'
        '<a:clrScheme name="Early Locomotion Charcoal and Orange">'
        % _A
        + c("dk1", BG) + c("lt1", TEXT) + c("dk2", PANEL) + c("lt2", MUTED)
        + c("accent1", ORANGE) + c("accent2", GREEN) + c("accent3", WATER) + c("accent4", PURPLE)
        + c("accent5", AMBER) + c("accent6", RED) + c("hlink", HLINK) + c("folHlink", FOLHLINK)
        + '</a:clrScheme><a:fontScheme name="Arial">'
        '<a:majorFont><a:latin typeface="Arial"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>'
        '<a:minorFont><a:latin typeface="Arial"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont>'
        '</a:fontScheme><a:fmtScheme name="Flat">'
        "<a:fillStyleLst>" + fills + "</a:fillStyleLst><a:lnStyleLst>" + lns + "</a:lnStyleLst>"
        "<a:effectStyleLst>" + effs + "</a:effectStyleLst><a:bgFillStyleLst>" + fills + "</a:bgFillStyleLst>"
        "</a:fmtScheme></a:themeElements><a:objectDefaults/><a:extraClrSchemeLst/></a:theme>"
    )


def _lvl(n, sz, *, marL=0, indent=0, algn="l", color="tx1", bold=False, bullet=None, caps=False, spc_bef=0,
         line_pct=None, font="+mn-lt", tag=None):
    """One a:lvlNpPr (or `tag`) with the child order PowerPoint requires: lnSpc, spcBef, bullet, defRPr."""
    tag = tag or "a:lvl%dpPr" % n
    lnspc = '<a:lnSpc><a:spcPct val="%d"/></a:lnSpc>' % (line_pct * 1000) if line_pct else ""
    if bullet == "dot":
        bu = ('<a:buClr><a:schemeClr val="accent1"/></a:buClr><a:buFont typeface="Arial"/>'
              '<a:buChar char="&#8226;"/>')
    elif bullet == "num":
        bu = ('<a:buClr><a:schemeClr val="accent1"/></a:buClr><a:buFont typeface="+mj-lt"/>'
              '<a:buAutoNum type="arabicPeriod"/>')
    else:
        bu = "<a:buNone/>"
    return (
        '<%s marL="%d" indent="%d" algn="%s" defTabSz="914400" rtl="0" eaLnBrk="1" latinLnBrk="0" '
        'hangingPunct="1">%s<a:spcBef><a:spcPts val="%d"/></a:spcBef>%s'
        '<a:defRPr lang="%s" sz="%d" b="%d"%s kern="1200"><a:solidFill><a:schemeClr val="%s"/></a:solidFill>'
        '<a:latin typeface="%s"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/></a:defRPr></%s>'
        % (tag.replace("a:", "a:"), E(marL), E(indent), algn, lnspc, spc_bef * 100, bu, LANG, int(sz * 100),
           1 if bold else 0, ' cap="all"' if caps else "", color, font, tag)
    )


def default_text_style_xml():
    """presentation.xml defaultTextStyle: 14 pt Arial, tx1, en-GB. New text boxes in PowerPoint start at the floor."""
    out = '<p:defaultTextStyle %s><a:defPPr><a:defRPr lang="%s"/></a:defPPr>' % (_NS, LANG)
    for n in range(1, 10):
        out += _lvl(n, 14, marL=(n - 1) * 0.5)
    return out + "</p:defaultTextStyle>"


def _master_ph(sid, name, ph, box_in, anchor, lst, prompt, wrap_extra=""):
    x, y, w, h = box_in
    return (
        '<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
        '<p:nvPr><p:ph %s/></p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/>'
        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody>'
        '<a:bodyPr vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="%s">'
        '<a:noAutofit/></a:bodyPr><a:lstStyle>%s</a:lstStyle>%s</p:txBody></p:sp>'
        % (sid, name, ph, E(x), E(y), E(w), E(h), anchor, lst, prompt)
    )


def _prompt(*texts):
    out = ""
    for i, t in enumerate(texts):
        ppr = '<a:pPr lvl="%d"/>' % i if i else ""
        out += '<a:p>%s<a:r><a:rPr lang="%s"/><a:t>%s</a:t></a:r></a:p>' % (ppr, LANG, t)
    return out


SLDNUM_FLD_ID = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"


def sldnum_para(lang=LANG, text="&#8249;#&#8250;"):
    return ('<a:p><a:fld id="%s" type="slidenum"><a:rPr lang="%s"/><a:t>%s</a:t></a:fld>'
            '<a:endParaRPr lang="%s"/></a:p>' % (SLDNUM_FLD_ID, lang, text, lang))


def master_xml():
    """Slide master: charcoal background, title / body / slide-number placeholders, Arial text styles.
    The layout list is filled in by build_template.py."""
    title = _master_ph(2, "Title Placeholder 1", 'type="title"', TITLE_BOX, "ctr", "",
                       _prompt("Click to edit the title"))
    body = _master_ph(3, "Text Placeholder 2", 'type="body" idx="1"', (0.6, CONTENT_TOP, 12.13, 4.7), "t", "",
                      _prompt("Click to edit the text", "Second level", "Third level", "Fourth level", "Fifth level"))
    num = _master_ph(4, "Slide Number Placeholder 3", 'type="sldNum" sz="quarter" idx="4"', NUMBER_BOX, "ctr",
                     _lvl(1, 14, algn="r", color="tx2"), sldnum_para())
    title_style = ('<p:titleStyle>' + _lvl(1, 36, bold=True, color="tx1", font="+mj-lt") + "</p:titleStyle>")
    body_style = "<p:bodyStyle>"
    sizes = {1: 20, 2: 18, 3: 16}
    for n in range(1, 10):
        body_style += _lvl(n, sizes.get(n, 14), marL=0.3 * n, indent=-0.3, bullet="dot", spc_bef=6 if n == 1 else 3)
    body_style += "</p:bodyStyle>"
    other_style = "<p:otherStyle>" + '<a:defPPr><a:defRPr lang="%s"/></a:defPPr>' % LANG
    for n in range(1, 10):
        other_style += _lvl(n, 14, marL=(n - 1) * 0.5)
    other_style += "</p:otherStyle>"
    return (
        '<p:sldMaster %s><p:cSld><p:bg><p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill><a:effectLst/>'
        '</p:bgPr></p:bg><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/>'
        '</a:xfrm></p:grpSpPr>%s%s%s</p:spTree></p:cSld>'
        '<p:clrMap bg1="dk1" tx1="lt1" bg2="dk2" tx2="lt2" accent1="accent1" accent2="accent2" accent3="accent3" '
        'accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
        '<p:sldLayoutIdLst/><p:txStyles>%s%s%s</p:txStyles></p:sldMaster>'
        % (_NS, BG.lstrip("#"), title, body, num, title_style, body_style, other_style)
    )


# ============================================================================ 10. Template API
def layout_shapes(layout):
    """A shapes collection that can draw on a slide layout (python-pptx does not offer one)."""
    return SlideShapes(layout._element.cSld.spTree, layout)


def open_template(path=None):
    """Open the template as a Presentation (no slides)."""
    return Presentation(path or TEMPLATE_PATH)


def new_deck(title, subject="", keywords="", template=None, author=AUTHOR):
    """Start a deck from the template. Sets title, author, language and revision."""
    prs = open_template(template)
    set_core_properties(prs, title, subject, keywords, author)
    return prs


def get_layout(prs, name):
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    raise KeyError("no layout named %r; have %s" % (name, [l.name for l in prs.slide_layouts]))


def ph(slide, key):
    """Placeholder on a slide by name (as named in the layout) or by idx."""
    for s in slide.placeholders:
        if s.name == key or s.placeholder_format.idx == key:
            return s
    raise KeyError("slide has no placeholder %r; have %s" % (key, [s.name for s in slide.placeholders]))


def has_ph(slide, key):
    try:
        ph(slide, key)
        return True
    except KeyError:
        return False


def add_slide(prs, layout, title=None, status=None, caption=None):
    """New slide from a layout (by name or object). Placeholders take the layout's names."""
    lay = get_layout(prs, layout) if isinstance(layout, str) else layout
    s = prs.slides.add_slide(lay)
    by_idx = {p.placeholder_format.idx: p for p in lay.placeholders}
    for shp in s.placeholders:
        src = by_idx.get(shp.placeholder_format.idx)
        if src is not None:
            shp.name = src.name
    if title is not None:
        fill_text(s.shapes.title.text_frame, title)
    if status:
        set_status(s, status)
    if caption is not None:
        set_caption(s, caption)
    return s


def set_status(slide, kind, text=None):
    """Fill the top-right status chip placeholder and colour it for `kind` (Agreed, Proposed, ...)."""
    shp = ph(slide, "Status chip")
    fill, tcol, outline = CHIPS[kind]
    shp.text_frame.clear()
    fill_text(shp.text_frame, [[(text or kind.upper(), {"color": tcol, "bold": True})]])
    _strip_style(shp)
    _round_geom(shp, 0.12)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if outline is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(outline)
        shp.line.width = Pt(2.0)
        shp.line.dash_style = MSO_LINE.SOLID
    set_alt(shp, "Status: " + kind)
    return shp


def _round_geom(shp, radius):
    """Explicit position (all four values together) and rounded-rectangle geometry on a placeholder."""
    x, y, w, h = shp.left, shp.top, shp.width, shp.height
    shp.left, shp.top, shp.width, shp.height = x, y, w, h
    spPr = shp._element.spPr
    for g in spPr.findall(qn("a:prstGeom")):
        spPr.remove(g)
    geom = etree.Element(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    gd = etree.SubElement(etree.SubElement(geom, qn("a:avLst")), qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", "val %d" % int(min(0.5, radius / min(Emu(w).inches, Emu(h).inches)) * 100000))
    spPr.find(qn("a:xfrm")).addnext(geom)


def set_caption(slide, text):
    """Caption text (20 pt, set by the layout)."""
    shp = ph(slide, "Caption")
    shp.text_frame.clear()
    fill_text(shp.text_frame, text)
    return shp


def put(slide, key, paras, **kw):
    """Fill any placeholder by name or idx, keeping the layout's formatting. Returns the shape."""
    shp = ph(slide, key)
    shp.text_frame.clear()
    fill_text(shp.text_frame, paras, **kw)
    return shp


def fill_picture(slide, file_path, alt, key="Picture", credit_text=None, radius=0.08):
    """Insert a picture into the layout's picture placeholder (cropped to fit, never stretched), with alt text,
    rounded corners and an optional credit line in the layout's credit placeholder."""
    shp = ph(slide, key)
    x, y, w, h = shp.left, shp.top, shp.width, shp.height
    pic = shp.insert_picture(file_path)
    pic.left, pic.top, pic.width, pic.height = x, y, w, h       # set all four together
    round_corners(pic, Emu(w).inches, Emu(h).inches, radius)
    set_alt(pic, alt, "Picture")
    if credit_text is not None:
        put(slide, "Credit line", credit_text)
    return pic


def fill_table(slide, data, col_w, key="Table", alt="Table", size=14, header=True, aligns=None,
               header_fill=PANEL_ALT, row_fill=None, row_h=None):
    """Insert a table into the layout's table placeholder and format every cell (fills, borders, size)."""
    shp = ph(slide, key)
    h_in = Emu(shp.height).inches
    gf = shp.insert_table(len(data), len(data[0]))
    gf.name = "Table"
    total_h = row_h * len(data) if row_h else h_in
    _format_table(gf.table, data, col_w, size, header, header_fill, aligns, row_fill, LINE, total_h)
    set_alt(gf, alt, "Table")
    return gf


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def _add_slide_number(slide):
    """Slide-number placeholder with a live field, cloned from the layout's (python-pptx skips it)."""
    for shp in slide.shapes:
        if shp.is_placeholder and shp.placeholder_format.type is not None and "SLIDE_NUMBER" in str(
                shp.placeholder_format.type):
            return shp
    lay = slide.slide_layout
    src = None
    for p in lay.placeholders:
        if "SLIDE_NUMBER" in str(p.placeholder_format.type):
            src = p
    if src is None:
        return None
    ids = [int(i) for i in slide.shapes._spTree.xpath("//p:cNvPr/@id")]
    el = copy.deepcopy(src._element)
    for x in el.xpath("./p:spPr/*"):
        x.getparent().remove(x)                       # inherit position from the layout
    c = el.xpath(_CNVPR_XPATH)[0]
    c.set("id", str(max(ids) + 1))
    c.set("name", "Slide number")
    idx = list(slide.part.package.presentation_part.presentation.slides).index(slide) + 1
    for t in el.iter(qn("a:t")):
        t.text = str(idx)
    slide.shapes._spTree.append(el)
    return el


def finish(slide, prune=True, number=True):
    """Final tidy for one slide: drop unfilled placeholders, add the slide number, put shapes in reading order
    (title, status chip, visuals, caption, slide number)."""
    tree = slide.shapes._spTree
    if prune:
        for shp in list(slide.placeholders):
            if shp.has_text_frame and not shp.text_frame.text.strip():
                remove_shape(shp)
    if number and has_layout_number(slide.slide_layout):
        _add_slide_number(slide)
    order_reading(slide)
    return slide


def has_layout_number(layout):
    return any("SLIDE_NUMBER" in str(p.placeholder_format.type) for p in layout.placeholders)


def _role(shp):
    n = (shp.name or "").lower()
    if shp.is_placeholder:
        t = str(shp.placeholder_format.type)
        if "TITLE" in t and "SUB" not in t:
            return 0
        if n == "status chip":
            return 1
        if "SLIDE_NUMBER" in t:
            return 4
        if n == "caption":
            return 3
        return 2
    if n.startswith("status chip"):
        return 1
    if n == "slide number":
        return 4
    if n == "caption":
        return 3
    return 2


def order_reading(slide):
    """Reorder top-level shapes: title, chip, everything else (in current order), caption, slide number."""
    tree = slide.shapes._spTree
    shapes = list(slide.shapes)
    keyed = sorted(enumerate(shapes), key=lambda t: (_role(t[1]), t[0]))
    anchor = tree.find(qn("p:grpSpPr"))
    for _, shp in keyed:
        tree.remove(shp._element)
    cur = anchor
    for _, shp in keyed:
        cur.addnext(shp._element)
        cur = shp._element


def move_before(shape, other):
    """Put `shape` immediately before `other` in the z-order and reading order of the same slide."""
    el = shape._element
    el.getparent().remove(el)
    other._element.addprevious(el)


def clone_shape(shape, slide):
    """Copy a shape onto another slide the safe way: strips p:custDataLst and creationId extensions that
    otherwise carry over and confuse PowerPoint. Pictures need their own relationship, so copy those with
    add_picture instead."""
    el = copy.deepcopy(shape._element)
    for bad in el.xpath(".//p:custDataLst"):
        bad.getparent().remove(bad)
    for ext in el.xpath(".//p:extLst/p:ext[a:*[local-name()='creationId']] | .//a:extLst/a:ext[*[local-name()='creationId']]"):
        ext.getparent().remove(ext)
    ids = [int(i) for i in slide.shapes._spTree.xpath("//p:cNvPr/@id")]
    for k, c in enumerate(el.iter(qn("p:cNvPr"))):
        c.set("id", str(max(ids) + 1 + k))
    slide.shapes._spTree.append(el)
    return el


def clear_timing(slide):
    """Remove any transition or animation block (a cloned slide can carry a copied p:timing)."""
    sld = slide._element
    for tag in ("p:transition", "p:timing"):
        for el in sld.findall(qn(tag)):
            sld.remove(el)
    for el in sld.findall("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"):
        sld.remove(el)


# ============================================================================ 11. Package tidy
def set_core_properties(prs, title, subject="", keywords="", author=AUTHOR):
    import datetime
    cp = prs.core_properties
    cp.title = title
    cp.author = author
    cp.last_modified_by = author
    cp.language = LANG
    cp.subject = subject
    cp.keywords = keywords
    cp.comments = ""
    cp.revision = 1
    now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None, microsecond=0)
    cp.created = now
    cp.modified = now


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<Application>Early Locomotion kit (python-pptx)</Application>'
           '<PresentationFormat>Widescreen</PresentationFormat><Company></Company></Properties>')


def tidy_package(prs):
    """Drop the stale parts python-pptx's default template carries (Mac printer settings, a blank thumbnail,
    4:3 slide-size tag, old app.xml)."""
    pkg = prs.part.package
    for rId, rel in list(pkg._rels.items()):
        if rel.reltype.endswith("/thumbnail"):
            pkg._rels.pop(rId)
    for rId, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/printerSettings"):
            prs.part.rels.pop(rId)
    sz = prs.part._element.find(qn("p:sldSz"))
    if sz is not None and "type" in sz.attrib:
        del sz.attrib["type"]
    for part in pkg.iter_parts():
        if str(part.partname) == "/docProps/app.xml":
            part._blob = APP_XML.encode("utf8")


def write_potx(src_pptx, dst_potx):
    """Copy a .pptx as a .potx: only the main part's content type changes."""
    old = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
    new = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
    with zipfile.ZipFile(src_pptx) as zin, zipfile.ZipFile(dst_potx, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                s = data.decode("utf8")
                assert old in s, "main part content type not found"
                data = s.replace(old, new).encode("utf8")
            zout.writestr(item, data)
